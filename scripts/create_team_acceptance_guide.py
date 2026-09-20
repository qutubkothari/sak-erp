from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
from datetime import date

OUT = Path('output/docx/SAK_ERP_Team_Acceptance_Guide.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = '3F2B1D'
BROWN = '80613B'
GOLD = 'B18B4B'
BEIGE = 'F7F2E9'
LIGHT = 'FBF9F5'
GREEN = '2F7D59'
RED = 'A33232'
GREY = '666666'

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.67)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.72)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.32)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}'); tcMar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); node = OxmlElement('w:tblHeader'); node.set(qn('w:val'), 'true'); trPr.append(node)

def set_table_widths(table, widths):
    tbl = table._tbl; tblPr = tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), '9360'); tblW.set(qn('w:type'), 'dxa')
    ind = tblPr.first_child_found_in('w:tblInd')
    if ind is None:
        ind = OxmlElement('w:tblInd'); tblPr.append(ind)
    ind.set(qn('w:w'), '120'); ind.set(qn('w:type'), 'dxa')
    grid = tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn('w:w'), str(int(width * 1440)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'), str(int(width * 1440))); tcW.set(qn('w:type'), 'dxa')
    if table.rows:
        set_repeat_table_header(table.rows[0])

def set_run(run, size=10.5, color='222222', bold=False, italic=False):
    run.font.name = 'Calibri'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color); run.bold = bold; run.italic = italic

styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri'); normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.16
for name, size, color, before, after in [('Heading 1', 17, NAVY, 16, 7), ('Heading 2', 13, BROWN, 11, 5), ('Heading 3', 11.3, NAVY, 8, 3)]:
    s = styles[name]; s.font.name = 'Calibri'; s._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri'); s.font.size = Pt(size); s.font.color.rgb = RGBColor.from_string(color); s.font.bold = True; s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after); s.paragraph_format.keep_with_next = True
if 'Small Label' not in styles:
    s = styles.add_style('Small Label', WD_STYLE_TYPE.PARAGRAPH); s.font.name='Calibri'; s.font.size=Pt(8); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(BROWN); s.paragraph_format.space_after=Pt(2)

# Header and footer
header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run('SAK ERP | Team Acceptance Guide'); set_run(r, 8.5, BROWN, True)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('Internal use - Test environment first | '); set_run(r, 8.5, GREY)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def p(text='', style=None, align=None, before=None, after=None):
    para = doc.add_paragraph(style=style)
    if align is not None: para.alignment = align
    if before is not None: para.paragraph_format.space_before = Pt(before)
    if after is not None: para.paragraph_format.space_after = Pt(after)
    if text:
        rr = para.add_run(text); set_run(rr)
    return para

def title(text, subtitle=None):
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.LEFT; para.paragraph_format.space_before = Pt(18); para.paragraph_format.space_after = Pt(3)
    rr = para.add_run(text); set_run(rr, 28, NAVY, True)
    if subtitle:
        sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(16); rr=sp.add_run(subtitle); set_run(rr, 13, BROWN)

def callout(label, text, fill=BEIGE):
    table = doc.add_table(rows=1, cols=1); table.alignment = WD_TABLE_ALIGNMENT.LEFT; set_table_widths(table,[6.5])
    cell=table.cell(0,0); set_cell_shading(cell,fill); set_cell_margins(cell,130,120,130,120); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    para=cell.paragraphs[0]; para.paragraph_format.space_after=Pt(0)
    r=para.add_run(label.upper()+': '); set_run(r,9.5,BROWN,True)
    r=para.add_run(text); set_run(r,9.5,'333333')
    p('', after=3)

def bullets(items):
    for item in items:
        para = doc.add_paragraph(style='List Bullet'); para.paragraph_format.space_after=Pt(2); para.paragraph_format.left_indent=Inches(.22); para.paragraph_format.first_line_indent=Inches(-.15)
        rr=para.add_run(item); set_run(rr,10.1)

def test_table(rows):
    table=doc.add_table(rows=1, cols=3); table.alignment=WD_TABLE_ALIGNMENT.LEFT; table.style='Table Grid'; set_table_widths(table,[0.55,3.25,2.7])
    headers=['#','Team action','Expected result / evidence']
    for c, txt in zip(table.rows[0].cells,headers):
        set_cell_shading(c,NAVY); set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        q=c.paragraphs[0]; q.paragraph_format.space_after=Pt(0); r=q.add_run(txt); set_run(r,9,'FFFFFF',True)
    set_repeat_table_header(table.rows[0])
    for idx,(action, expected) in enumerate(rows,1):
        cells=table.add_row().cells
        for c in cells: set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if idx%2==0:
            for c in cells: set_cell_shading(c,LIGHT)
        for c, txt, align in [(cells[0],str(idx),WD_ALIGN_PARAGRAPH.CENTER),(cells[1],action,WD_ALIGN_PARAGRAPH.LEFT),(cells[2],expected,WD_ALIGN_PARAGRAPH.LEFT)]:
            par=c.paragraphs[0]; par.alignment=align; par.paragraph_format.space_after=Pt(0); r=par.add_run(txt); set_run(r,9.1)
    set_table_widths(table,[0.55,3.25,2.7])
    p('',after=4)

def module(number, name, purpose, features, rows, evidence):
    doc.add_heading(f'{number}. {name}', level=1)
    p(purpose, after=4)
    lab=p('FEATURES TO CHECK', style='Small Label');
    bullets(features)
    p('CLICK-BY-CLICK ACCEPTANCE', style='Heading 2')
    test_table(rows)
    callout('Evidence to attach', evidence, 'F1EEE8')

# Cover
title('SAK ERP Team Acceptance Guide', 'Step-by-step functional test instructions for the complete application, including recently added controls and intelligence features')
meta = doc.add_table(rows=4, cols=2); meta.style='Table Grid'; set_table_widths(meta,[1.55,4.95])
for i,(k,v) in enumerate([('Audience','Functional leads, super users, finance, stores, production, HR and management'),('Environment','Use Mizantra test first. Do not create demonstration transactions on live.'),('Version','Operational acceptance guide - August 2026'),('Prepared for','SAK ERP implementation team')]):
    a,b=meta.rows[i].cells; set_cell_shading(a,BEIGE); set_cell_margins(a); set_cell_margins(b); pa=a.paragraphs[0]; pb=b.paragraphs[0]; pa.paragraph_format.space_after=pb.paragraph_format.space_after=Pt(0); ra=pa.add_run(k); set_run(ra,9.5,BROWN,True); rb=pb.add_run(v); set_run(rb,9.5)
set_table_widths(meta,[1.55,4.95])
p('',after=8)
callout('How to use this guide','Each module is a test script. Complete the actions in sequence, compare the outcome with the expected result, capture the named evidence, and record failures with screen, timestamp, URL, user, and document number.', 'EFE6D5')
doc.add_heading('Release discipline', level=1)
bullets(['Run every scenario in Mizantra test before client demonstration or live release.', 'Use a dedicated test vendor, customer, employee and warehouse. Label all test notes with TEST and the date.', 'Do not alter approved masters, tax settings, opening balances, or live transactions while testing.', 'For any failure, stop that scenario, keep the document number, take a screenshot, and log the exact action that failed.', 'Hard-refresh the browser (Ctrl+Shift+R) after a deployment before reporting a UI issue.'])
doc.add_heading('Pass criteria', level=1)
test_table([('Complete the workflow with no error banner, failed request, duplicate document, or incorrect status.','Attach the generated document number and screenshot of the final status.'),('Confirm stock, financial value, approval status, and audit trail agree with the event.','Use the relevant trail, ledger, or report as proof.'),('Repeat the relevant role test where maker-checker or approval is configured.','Capture both the maker action and reviewer/approver action.')])
doc.add_page_break()

doc.add_heading('Contents and test ownership', level=1)
contents=[('A. Foundation and governance','Admin / Super User'),('B. Procurement to payment','Purchase, Stores, Finance'),('C. Inventory and traceability','Stores / Quality'),('D. Sales to cash and service','Sales / Service / Finance'),('E. Production, planning and subcontracting','Planning / Production / Stores / Quality'),('F. Finance, HR and statutory operations','Finance / HR'),('G. Intelligence, ROI and integrations','Management / IT')]
test_table(contents)
callout('Suggested team cadence','Assign one owner and one verifier for each module. Review failures daily; rerun only the affected scenario after a fix, then rerun its upstream and downstream integration checks.', 'F1EEE8')

module('1','Access, company settings and control framework','Verify that the tenant identity, access model, notifications, print identity and auditability work before transactions are created.',
 ['Company profile, logo, address, tax identity and document stationery from Master Settings', 'Role-based menu visibility, maker-checker roles, approval delegation and audit trail', 'Notification center, SLA reminders/escalation and responsive navigation'],
 [('Sign in as Super Admin; open Settings > Organization / Company Header. Review name, address, logo, tax registration, currency and market profile.','Dashboard and print previews use the configured company identity; no hard-coded company address.'),('Open Employee Access; compare two users with different roles. Sign in as each user.','Only authorised modules/actions are visible; restricted action is blocked cleanly.'),('Create or view a controlled approval item, then open Audit Trails.','Actor, timestamp, old/new state and document reference are traceable.'),('Create an approval that is due soon; open Notification Center / Action Required.','Reminder is visible, counts match the queue and the link opens the correct work item.'),('Print a PR, PO or GRN preview.','Stationery displays company name/address from master settings and document number/date.')],
 'Screenshots of company header, two role views, one audit trail, notification card and a stationery preview.')

module('2','Executive dashboard and cross-module worklist','Confirm management sees a reliable operational picture and can drill from exceptions to action.',
 ['Operations dashboard, approvals, document ageing, low-stock risk and production WIP', 'Command/operations cockpit, exception worklist and direct drill-down', 'Country-aware currency display based on tenant market profile'],
 [('Open Dashboard and click Refresh.','Metrics refresh without error; last-refreshed timestamp updates.'),('Open each card: approvals, PO exposure, supplier invoice value, stock risk and production WIP.','Each opens a relevant filtered view or worklist; totals are explainable from records.'),('Open an exception worklist item, then return to dashboard.','The linked document opens; no dead link or blank page.'),('Switch only between authorised test tenant contexts if configured.','Currency, statutory labels and tenant data align with the selected market profile.'),('Resize browser or test laptop view.','Sidebar, cards and notification tray remain usable without hidden controls.')],
 'Dashboard screenshot before/after refresh and one drill-down screenshot for each exception type.')

module('3','Master data, item engineering and vendor/customer governance','Prove the controlled master foundation used by procurement, stock, production, sales and finance.',
 ['Item master, UOM, stock minimums, HSN/tax, dimensions and size in MM', 'Vendor/customer master approvals, bank/tax/contact data and duplicate prevention', 'Warehouse, bin, cost centre, employee and document master records'],
 [('Create a TEST raw material with inventory UOM, reorder level, tax/HSN and a description.','Item saves once, appears in item search and is visible to purchase/inventory.'),('Create a TEST finished product produced in NUMBER/PCS, with Size = 30 and Size UOM = MM.','Master retains both count UOM and dimensional size; no forced conversion into one field.'),('Create a test vendor/customer; submit through configured approval.','Maker-checker status changes correctly; unapproved master cannot be used where enforcement applies.'),('Open warehouses and verify source/output warehouse selection.','Only active warehouses are selectable; stock locations are named clearly.'),('Attempt a duplicate item or vendor reference where governance is enabled.','System warns or blocks duplicate according to master-data policy.')],
 'Master record screenshots, approval trail, and item search showing size/UOM.')

module('4','Procure-to-pay: PR, RFQ, PO, GRN, invoice and payment','Test the end-to-end purchasing chain and the financial controls that prevent unmatched spend.',
 ['Purchase requisitions, RFQs, comparative sourcing, PO approvals and amendments', 'GRN and line-level quality, supplier invoice capture, debit notes and payment status', 'Spend intelligence, contract control, import files and service entry sheets'],
 [('Create a TEST purchase requisition with item, quantity, warehouse, need-by date and justification. Submit for approval.','PR gets a unique number and correct approval status.'),('Create RFQ from the approved PR; invite or enter at least two test supplier quotations.','Comparative view shows supplier, price, tax, lead time and chosen source.'),('Convert selected quote to PO; verify tax, freight, discount, terms and approval.','PO total and approval flow are correct; linked PR/RFQ trail is visible.'),('Post a partial GRN, inspect one line in QC, then post the balance.','Received, accepted/rejected and stock balances reconcile at line level.'),('Record a supplier invoice against GRN; open AP/payment status.','Three-way match evidence is available; unmatched value is blocked or flagged.'),('Open Spend Intelligence and create a savings case from a price opportunity.','Case is linked to underlying procurement evidence; it does not alter PO/vendor automatically.')],
 'PR/RFQ/PO/GRN/invoice numbers, comparison view, GRN QC evidence and AP status screenshot.')

module('5','Inventory, stores and UID traceability','Validate reliable stock movements, count controls and accountability for all material issues.',
 ['Stock master, warehouse transfer, adjustment, cycle count, reorder signals and valuation', 'Manual SIV with recipient employee selection', 'Batch/serial/UID tracking, document attachments and stock trail'],
 [('Receive test material through GRN, then open Stock Trail.','Available quantity increases exactly once; trail shows GRN reference and warehouse.'),('Create a manual SIV/issue; select an employee from registered employees and enter purpose.','Issue shows named recipient, employee reference, quantity/UOM and audit history.'),('Transfer a controlled quantity between test warehouses.','Source decreases and destination increases; transfer document and trail agree.'),('Run a cycle count/stock adjustment under the configured approval rule.','Variance is evidenced, approved where required, and visible in audit/stock trail.'),('Open UID Tracking for a tracked item and scan/search its code.','System resolves the item history, documents and current state without cross-tenant data.')],
 'Stock trail before/after, SIV with employee recipient, transfer reference, and UID trace screenshot.')

module('6','Order-to-cash and customer service','Confirm sales, dispatch, collection and service entitlement work as an integrated client journey.',
 ['Quotation, sales order, delivery/dispatch, tax invoice, collections and credit exposure', 'Customer master/credit controls, project linkage and returns/credit notes', 'Service tickets, contracts, entitlement, technician workflow and service finance'],
 [('Create a test quotation and convert it to a sales order.','Customer, items, pricing, tax, delivery terms and audit trail carry through.'),('Check credit control for a customer near/over its limit.','System warns or blocks based on policy and leaves an approval/audit record.'),('Dispatch/fulfil an order and create the invoice.','Stock, delivery status and receivable update once; document trail connects all stages.'),('Record a partial customer payment and open AR ageing/cash application.','Outstanding amount and payment allocation are correct.'),('Create a service ticket against customer/item/contract; progress it to closure.','Entitlement and service status are visible; chargeable work can flow to finance.')],
 'Quote/SO/invoice numbers, stock/AR proof, credit control warning, ticket closure evidence.')

module('7','Production planning and shop-floor execution','Test planning-to-output controls including capacity, MRP, work orders, material consumption, output and loss intelligence.',
 ['BOM/routing, job order, MRP, capacity planning, S&OP and work-centre execution', 'OEE/loss control, digital work instructions, engineering changes and production autonomy', 'Production cost, energy/carbon per unit, maintenance risk and exception management'],
 [('Create or select a BOM/routing and release a test job order.','Required components, operations, dates and quantities are calculated correctly.'),('Run MRP/capacity planning for a known demand.','Shortage/capacity recommendations are explainable and link to demand, stock and supplier lead time.'),('Issue material to job; record production completion, scrap and rework.','WIP, FG stock, consumption and losses update consistently.'),('Open OEE & Loss Control and record/inspect a downtime reason.','Availability/performance/quality values and the loss reason are visible by machine/shift.'),('Open Digital Work Instructions / Engineering Changes for the job.','Correct current revision is shown; acknowledgement and change audit are retained.'),('Review Production Autonomy / Factory Health recommendation.','System proposes a governed action with source evidence; no uncontrolled stock or schedule change occurs.')],
 'Job number, MRP recommendation, production/stock trail, OEE evidence and work-instruction revision.')

module('8','Subcontracting / outside processing - complete acceptance','Verify the complete outside-processing flow, including size/UOM intelligence, one-time raw-material issue, line QC and finance traceability.',
 ['Routes with raw material, output product, quantity UOM, Size (MM) and output size/UOM carried from item master', 'One MOC/material outward challan per subcontract order; stores printout and vendor WIP', 'Multi-output receipts, raw-material backflush/return/scrap, line-by-line QC and finance/AP integration', 'Order trail showing raw material, UOM, GST, freight, deductions, invoice and payable details'],
 [('Open Production > Subcontracting > Routes. Create a TEST route. Select raw material and verify its UOM/available source. Add an output item with NUMBER/PCS and Size 30 MM.','Raw material is selectable; output quantity UOM and size are displayed/stored separately. Route saves once only.'),('Create a subcontracting order from the route with vendor, source warehouse, output warehouse, output lines, price, HSN and tax.','Order is created once with a unique SUB number; no duplicate entry on a single Save click.'),('Open the Order Trail before issue.','Trail displays raw material, quantity and UOM, outputs, financial fields and current finance status.'),('Click Issue Material. Enter external challan/reference and notes, then post. Open Print Issue Slip / stores printout.','Exactly one MOC is created; printout includes company stationery, vendor, RM, quantity, UOM, warehouse, order and challan reference.'),('Open Vendor WIP.','Issued quantity appears against vendor/order; no duplicate stock issue occurs.'),('Receive processed material for multiple output lines. Enter rejected, scrap, unused RM return or approved loss only when applicable.','Each received line keeps its own product/UOM/quantity; raw-material consumption/return calculation is visible.'),('Open QC for the receipt. Approve/reject each line separately and choose rework/scrap for any rejection.','Every FG line requires inspection; accepted/rejected values and QC notes are retained per line.'),('Record supplier invoice details, freight/other charges/deductions; approve QC; open Accounts > Subcontract Payables.','GST/tax, deductions, processing, freight and payable totals are correct. Finance is in Accounts, not a Subcontracting tab.'),('Mark payment with a test reference; reopen order trail.','Payable changes to PAID and trail shows end-to-end flow: order > MOC > GRN > QC > invoice > payment.'),('Test a partial receipt: receive part, QC it, then receive remaining balance and QC it.','Open MOC balance decreases correctly; no premature payable before QC; final balance closes after approved receipt.')],
 'Route, SUB, MOC and SCR numbers; issue-slip PDF; per-line QC screen; vendor WIP; Accounts payables row; final document trail.')

module('9','Quality management','Verify inspection plans, nonconformance, CAPA and evidence-based quality release across purchase, production and subcontracting.',
 ['GRN and subcontract line-level QC, inspection results, rejected disposition and rework/scrap', 'Nonconformance, CAPA, supplier quality and evidence attachment', 'Quality dashboard, alerts and auditability'],
 [('Open a pending incoming GRN QC and inspect each line.','Line values, accept/reject results and notes are mandatory where configured.'),('Reject a line and select a permitted disposition.','System requires valid rework/scrap treatment and adjusts inventory/status correctly.'),('Create a nonconformance and CAPA linked to item, supplier or production order.','Owner, due date, root cause, corrective action and evidence are captured.'),('Close CAPA after verification.','Closure is traceable; quality dashboard and affected document reflect the outcome.')],
 'QC report, rejected disposition proof, CAPA record and closure audit.')

module('10','Maintenance, assets and workforce operations','Validate preventative work, breakdown response, asset history, labour capability and cost visibility.',
 ['Plant maintenance, maintenance plans, work orders, asset register and spare parts', 'Predictive/condition evidence capture and controlled maintenance recommendations', 'Employee skills, attendance/leave, payroll and WPS/India country-specific controls'],
 [('Create a test preventive maintenance plan and generate a work order.','Due date, asset, checklist, technician and spares are populated.'),('Record a breakdown/downtime event against a machine.','Event links to asset and production impact; maintenance work order can be raised.'),('Complete the work order with labour/spares and evidence.','Asset history, cost and status update; no duplicate spare issue.'),('Open HR > My Leaves as an employee and submit a test leave request.','Balance, request status, approver and notification behave correctly.'),('Run payroll/WPS or India payroll only with approved test data.','Country-specific statutory fields and control reports are present; no live payroll posting.')],
 'Maintenance WO, asset history, downtime link, leave request/approval and relevant payroll test output.')

module('11','Accounts, controls and statutory readiness','Confirm finance records and controls are complete from sub-ledgers through ledger, bank and statutory reporting.',
 ['Chart of accounts, journals, maker-checker, tax codes, cost centres, budgets and recurring journals', 'AP/AR, bank reconciliation, cash control, fixed assets, depreciation and multi-currency', 'Opening balances, close checklist, audit trails and UAE/India compliance configuration'],
 [('Open COA, cost centres, tax codes and posting rules. Validate approved test mappings.','Master codes are active only when finance-approved; mappings can be traced.'),('Create a controlled journal and process it through preparer/reviewer/approver/poster roles.','Segregation of duties works; posted journal is immutable except through controlled reversal.'),('Reconcile a test bank statement line to a payment/receipt.','Match, exception and owner are clear; bank balance and ledger reconciliation agree.'),('Create a fixed asset and run a controlled depreciation preview/posting.','Asset value, useful life, depreciation and ledger impact are calculated correctly.'),('Run AP ageing, AR ageing, trial balance, P&L, balance sheet and cash reports.','Reports reconcile to test transactions and support drill-down.'),('Open period-end checklist and attempt a controlled close/reopen scenario.','Controls prevent unauthorised posting and leave a clear audit entry.'),('Check UAE VAT/WPS/e-invoice or India GST/TDS/e-invoice/e-way-bill configuration in the relevant tenant.','Only the selected market labels/rates/mappings appear; statutory returns use approved tax codes.')],
 'Journal approval trail, bank reconciliation, asset/depreciation proof, reports and period-control evidence.')

module('12','ROI, Value Realization and management intelligence','Prove that claimed operational benefits are evidence-based, finance-verifiable and usable in a client conversation.',
 ['Value realization ledger, source-linked baseline/outcome evidence and finance verification', 'Cash vs P&L vs risk classification, payback, subscription value and client ROI statements', 'Anti-double-counting, Value Graph, benefit durability, forecast accuracy and country value libraries'],
 [('Open Value Realization. Create a TEST benefit sourced from procurement, production or collections.','Benefit records baseline, outcome, owner, evidence link, classification and amount.'),('Attach source document/transaction and submit for finance verification.','Evidence is linked; benefit remains unverified until finance action.'),('Create a related benefit from a second module using same source/period.','Duplicate/overlap warning is shown or benefit is flagged for review.'),('Open Value Graph / ROI evidence view.','Flow can be followed from source event to action, outcome, financial proof and ROI statement.'),('Review durability and forecast-vs-realised trend.','System identifies continuation, taper/reversal or accuracy variance.'),('Generate/review monthly ROI statement and value-to-renewal view.','Cash, P&L and risk avoided are clearly separated; payback/subscription value is visible.'),('Compare UAE/India value library selections in the matching test profile.','Relevant country-native use cases/controls are shown without mixing statutory frameworks.')],
 'Benefit record, evidence link, finance verification, duplicate flag, ROI statement and renewal/value cockpit screenshots.')

module('13','Integration hub, automation and governed AI','Confirm integrations and automation help the team without creating uncontrolled external or financial actions.',
 ['Integration hub, connector registration, event mapping and replay/error visibility', 'Approval SLA reminders/escalation, workflow editor and natural-language governed rules', 'Copilot, operational knowledge graph, root-cause briefs and controlled agent proposals'],
 [('Open Integration Hub and review a configured test connector or mock event.','Connector has owner, status, mapping, error/retry history and no exposed secret.'),('Send/replay a permitted test event to a mapped queue.','Event is captured, idempotent and creates only a governed draft/recommendation where configured.'),('Create a workflow rule for a simple test condition (e.g., overdue PO or low stock).','Rule explains trigger, owner, approval path and escalation; audit record is created.'),('Ask Copilot a supported question about finance, supply, production, quality or maintenance.','Answer is grounded in tenant data/reports and links back to source context where available.'),('Open an autonomous/operations agent proposal.','Agent may recommend a controlled action, but external communication or automatic posting remains approval-gated.'),('Review a historical root-cause/change brief.','Brief links affected documents, timeline and supporting evidence.')],
 'Connector status, event log, workflow audit, Copilot response with source context and agent proposal screenshot.')

module('14','Documents, reports, mobile use and final release gate','Ensure the application is ready for a full demonstration and that users can find, print and evidence every major transaction.',
 ['Document management, attachments, printouts, report schedules and exports', 'Mobile/PWA-responsive operation, notification tray, quick search and accessible navigation', 'Release gate, backup/rollback evidence, monitoring and production-readiness checklist'],
 [('Attach a test drawing/invoice/evidence file to a transactional record, then open Documents.','Attachment is accessible to authorised users and linked to the originating document.'),('Print/export a PR, PO, GRN, SIV, subcontract MOC, invoice and ROI statement.','Output contains master-driven stationery, document number, date, values and correct currency.'),('Create a report schedule using a test recipient/configuration.','Schedule is saved, auditable and does not send unapproved external messages.'),('Use Quick Search to locate a document by number and check the result.','Search returns the right module/document and opens it directly.'),('Run release-gate checklist: API health, web health, core flow smoke, error logs, disk capacity, backup point and rollback note.','All required checks are green or an exception is documented before any live change.'),('On a second browser/laptop, repeat dashboard, notification and subcontracting entry.','Behaviour is consistent; browser cache issues are resolved by hard refresh and not misreported as product defects.')],
 'Sample printouts, attachments, report schedule, quick-search proof, release checklist and backup/rollback reference.')

doc.add_page_break()
doc.add_heading('Failure logging template', level=1)
callout('Use this for every failure','Do not report only “error”. Record enough information for the team to reproduce and fix it in one pass.', 'FBEAEA')
failure=doc.add_table(rows=8, cols=2); failure.style='Table Grid'; set_table_widths(failure,[1.6,4.9])
for i,label in enumerate(['Module / scenario','Test environment and URL','User / role','Exact step that failed','Expected result','Actual result / error text','Document number(s)','Screenshot / console / timestamp']):
    a,b=failure.rows[i].cells; set_cell_shading(a,BEIGE); set_cell_margins(a,120,120,120,120); set_cell_margins(b,120,120,120,120); a.paragraphs[0].paragraph_format.space_after=Pt(0); b.paragraphs[0].paragraph_format.space_after=Pt(0); r=a.paragraphs[0].add_run(label); set_run(r,9.5,BROWN,True); r=b.paragraphs[0].add_run(''); set_run(r,9.5)
set_table_widths(failure,[1.6,4.9])
doc.add_heading('Completion sign-off', level=1)
test_table([('All assigned scenarios are passed, evidence is attached and any exceptions are listed.','Functional owner signs off module acceptance.'),('All cross-module document, stock, tax and financial reconciliations are reviewed.','Finance/stores/production joint sign-off is recorded.'),('Only authorised scope is released after a test backup and release gate.','Release owner records build/version, deployment time and rollback location.')])
p('This guide is designed for operational acceptance. Configure real company data, statutory mappings, banks, approval roles and integrations only after the corresponding finance/HR/IT owner signs off.', after=0)

doc.core_properties.title = 'SAK ERP Team Acceptance Guide'
doc.core_properties.subject = 'Step-by-step whole-application functional acceptance testing'
doc.core_properties.author = 'SAK ERP'
doc.save(OUT)
print(OUT.resolve())
