from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path("deliverables/Mizantra_Website_Design_Brief.docx")
NAVY, TEAL, GOLD, INK, MUTED, PALE = "17324D", "0D7C86", "B18445", "243342", "657483", "F3F7F8"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd")) or OxmlElement("w:shd")
    if shd.getparent() is None: tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    m = tcPr.find(qn("w:tcMar")) or OxmlElement("w:tcMar")
    if m.getparent() is None: tcPr.append(m)
    for side, value in (("top",90),("start",120),("bottom",90),("end",120)):
        e = m.find(qn("w:"+side)) or OxmlElement("w:"+side)
        if e.getparent() is None: m.append(e)
        e.set(qn("w:w"), str(value)); e.set(qn("w:type"), "dxa")

def width(cell, value):
    p = cell._tc.get_or_add_tcPr()
    e = p.find(qn("w:tcW")) or OxmlElement("w:tcW")
    if e.getparent() is None: p.append(e)
    e.set(qn("w:w"), str(value)); e.set(qn("w:type"), "dxa")

def make_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; width(c,widths[i]); margins(c); shade(c,NAVY)
        r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(255,255,255)
    trPr=t.rows[0]._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:tblHeader"))
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            c=cells[i]; width(c,widths[i]); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            if len(t.rows)%2==0: shade(c,PALE)
            r=c.paragraphs[0].add_run(str(v)); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(INK)
    return t

def p(doc, text):
    x=doc.add_paragraph(text); x.paragraph_format.space_after=Pt(6); return x
def bullet(doc, text):
    x=doc.add_paragraph(style="List Bullet"); x.paragraph_format.space_after=Pt(3); x.add_run(text); return x
def num(doc, text):
    x=doc.add_paragraph(style="List Number"); x.paragraph_format.space_after=Pt(3); x.add_run(text); return x

doc=Document(); s=doc.sections[0]
s.top_margin=Inches(.75); s.bottom_margin=Inches(.7); s.left_margin=Inches(.8); s.right_margin=Inches(.8)
normal=doc.styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(INK); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1
for name,size,color,before,after in [("Heading 1",17,NAVY,15,7),("Heading 2",13,TEAL,11,5),("Heading 3",11,GOLD,8,3)]:
    st=doc.styles[name]; st.font.name="Calibri"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
s.header.paragraphs[0].text="MIZANTRA | WEBSITE DESIGN BRIEF"; s.header.paragraphs[0].runs[0].font.size=Pt(8); s.header.paragraphs[0].runs[0].font.color.rgb=RGBColor.from_string(MUTED)
f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT; f.add_run("Confidential working brief | Prepared 21 August 2026").font.size=Pt(8)

x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.space_before=Pt(35); x.paragraph_format.space_after=Pt(4)
r=x.add_run("MIZANTRA"); r.font.size=Pt(34); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.space_after=Pt(5)
r=x.add_run("Website Design & Content Brief"); r.font.size=Pt(20); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.space_after=Pt(24)
r=x.add_run("A market-facing website for a practical, auditable manufacturing ERP"); r.italic=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(MUTED)
make_table(doc,["Brief detail","Direction"],[
("Audience","Website designer, UX/UI designer, copywriter and implementation team"),
("Primary goal","Explain Mizantra clearly and generate qualified manufacturing-ERP enquiries"),
("Product context","Cloud ERP for manufacturing, procurement, inventory, subcontracting, sales, service, HR and finance"),
("Brand character","Practical, controlled, trustworthy, industrial and modern"),
("Important split","India and UAE experiences must remain separate; do not mix tax, currency or statutory language")],[2200,7160])

doc.add_heading("1. Executive direction",1)
p(doc,"Mizantra should be presented as an operations-first ERP for manufacturers and service-led industrial businesses. The website must make the product feel dependable and concrete: every process has an owner, every stock movement has a document, every approval is traceable, and management can see the next action.")
p(doc,"Avoid generic “all-in-one software” language. Show the full operational chain: source material or customer request → planning → execution → inspection → inventory/finance → audit trail.")
p(doc,"Recommended core promise: “One controlled system from shop floor to finance.”")
p(doc,"Recommended supporting line: “Mizantra connects manufacturing, procurement, inventory, subcontracting, sales, service and people operations without losing the document trail.”")

doc.add_heading("2. Positioning and target customers",1)
doc.add_heading("Primary segments",2)
for z in ["Manufacturers with raw materials, work orders, production, quality and stock traceability needs.","Engineering, fabrication, machining, anodizing and outside-processing businesses.","Industrial distributors and project businesses needing procurement, inventory, sales and service in one system.","Growing companies replacing disconnected spreadsheets, WhatsApp approvals and separate accounting/stock tools."]: bullet(doc,z)
doc.add_heading("Buyer roles",2)
make_table(doc,["Role","What the website must answer"],[
("Owner / Managing Director","Will this give me control, visibility and accountability without a huge ERP programme?"),
("Operations Head","Can my team run material, production, subcontracting and service workflows end to end?"),
("Finance Head","Are invoices, taxes, payables, receivables, deductions and payment status traceable?"),
("Purchase / Stores","Can we control PR, RFQ, PO, GRN, QC, SIV, stock and vendor documents?"),
("Service Manager","Can we assign technicians, record service work, parts, evidence, warranty and billing?"),
("IT / Admin","Can the system support roles, approvals, audit trails and regional configuration?")],[2200,7160])

doc.add_heading("3. Website objectives and conversion",1)
for z in ["Communicate product value within the first screen.","Use real workflow visuals and document-trail language instead of stock photography alone.","Create separate, clear paths for Indian and UAE prospects.","Convert visitors through Book a Demo, Request a workflow review and Talk to an ERP specialist.","Build trust with security, auditability, implementation and support information.","Make every feature page scannable: problem, workflow, controls, outputs and CTA."]: bullet(doc,z)
doc.add_heading("Primary calls to action",2)
make_table(doc,["CTA","Use"],[("Book a live demo","Primary header, hero and final CTA"),("Request a workflow review","For manufacturers with a defined process problem"),("Explore the platform","Scrolls to modules and workflow overview"),("Download capability brief","Lead-generation asset"),("Contact sales","Secondary CTA on mobile and footer")],[2200,7160])

doc.add_heading("4. Proposed sitemap",1)
make_table(doc,["Page","Purpose and required content"],[
("Home","Positioning, proof, workflow overview, modules, regional entry point and CTA"),
("Platform overview","How master data, transactions, controls and reporting connect"),
("Manufacturing ERP","BOM/routes, work orders, material issue, production, QC and stock update"),
("Subcontracting / outside processing","Route, single-vendor order, one RM issue, vendor WIP, receipt, QC, payable and document trail"),
("Procurement and inventory","PR, RFQ, quotation, PO, GRN, QC, SIV, stock ledger and payables"),
("Sales and dispatch","Customer, quotation, revision, approval, sales order, dispatch, billing, receipt, returns and warranty"),
("Service management","Service call, technician, craft serial, site contact, evidence, parts, warranty, invoice and follow-up"),
("People and HR","Attendance movements, reasons, location, evidence, employee view and management register"),
("Industry / use cases","Machining, fabrication, industrial service, project manufacturing and distribution"),
("India edition","INR, GST, HSN, Indian statutory wording and India contact route"),
("UAE edition","AED, VAT/TRN, UAE statutory wording and UAE contact route"),
("Resources","Capability brief, process guides, FAQs and implementation content"),
("About / contact","Company, support model, offices, enquiry form and legal pages")],[2500,6860])

doc.add_heading("5. Home page wireframe",1)
for z in ["Announcement bar: a short, dismissible manufacturing-control message.","Header: logo, Platform, Solutions, Industries, Resources, India/UAE selector, Contact and Book a Demo.","Hero: headline, two-line explanation, primary/secondary CTA and a clear dashboard/workflow visual.","Trust strip: document trail, role-based approvals, live stock visibility, multi-location readiness and regional tax configuration.","Problem section: duplicate entry, unclear stock, vendor WIP, delayed QC and disconnected finance.","Workflow section: Plan → Purchase → Issue → Produce/Subcontract → Receive/QC → Stock → Finance.","Module cards: Manufacturing, Procurement, Inventory, Subcontracting, Sales, Service, HR and Finance.","Proof section: screenshots or short clips with captions, not decorative mockups with unreadable text.","Regional section: India Edition and UAE Edition cards.","Final CTA: request demo/workflow review with a short form."]: num(doc,z)

doc.add_heading("6. Product feature content",1)
doc.add_heading("Manufacturing and subcontracting",2)
for z in ["BOM/route-style definitions for input materials and output products.","Work orders with vendor, input quantity/UOM, output lines, sizes, pricing, HSN and discount where applicable.","One controlled raw-material issue per subcontract order and vendor-held WIP visibility.","Receipt by actual output, automatic/backflush raw-material accounting, unused material, scrap and approved loss.","QC inspection before final completion and payable release.","Full order trail linking service order, outward challan, GRN, QC, invoice, stock and payment."]: bullet(doc,z)
doc.add_heading("Procurement and inventory",2)
for z in ["PR, approval, RFQ, supplier quotation, PO and document attachments.","GRN with invoice, QC, statutory checks, deductions and payable handoff.","SIV/material issue, stock adjustments, stock ledger and warehouse transfers.","Searchable master data with UOM, category, HSN/tax data, serial/UID controls and stock trail."]: bullet(doc,z)
doc.add_heading("Sales",2)
for z in ["Customer master with contacts, addresses, tax details, credit limits and credit days.","Quotation with editable descriptions, HSN, images/documents, terms, validity, reminders, comments and revisions.","Sales order with approval, price/discount/tax controls, fulfilment and customer document flow.","Dispatch, invoice, receipt, returns, warranty and receivable visibility."]: bullet(doc,z)
doc.add_heading("Service",2)
for z in ["Service call/ticket with craft or asset serial number, site contact, technician assignment, status and SLA dates.","Technician check-in/check-out, work performed, parts used, photos/videos and customer acknowledgement.","Warranty detection with service-charge waiver and parts/labour billing rules.","Service invoice, customer payment, follow-up and service history by customer, asset and technician."]: bullet(doc,z)
doc.add_heading("HR and attendance",2)
for z in ["Check-in, go-out, return-to-office and end-day movement timeline.","Reason codes such as lunch, trials and official work, with admin-managed options.","Location evidence and selfie/photo support for off-premises checkout where configured.","Employee self-view plus HR management history."]: bullet(doc,z)

doc.add_heading("7. Signature workflow visual language",1)
p(doc,"Use a document-trail motif: a horizontal or vertical chain of numbered stages, each with a document name, status badge and owner. This explains Mizantra better than generic feature grids.")
make_table(doc,["Workflow","Stages to visualise"],[
("Procure to pay","PR → RFQ → Supplier quote → PO → GRN → QC → Supplier invoice → AP → Payment"),
("Subcontracting","Route → Work order/vendor → Material outward → Vendor WIP → Receipt → QC → Service invoice → Stock/payable"),
("Order to cash","Customer → Quotation → Revision/approval → Sales order → Dispatch → Sales invoice → Receipt"),
("Service to cash","Service call → Assignment → Visit/evidence → Parts/warranty → Confirmation → Invoice → Collection")],[2100,7260])

doc.add_heading("8. Visual design direction",1)
doc.add_heading("Brand personality",2)
for z in ["Industrial confidence: structured grids, clear borders, strong alignment and useful whitespace.","Human and practical: real people, technicians, stores and production environments.","Controlled but not bureaucratic: status colours should reassure, not alarm.","Premium B2B: restrained palette, crisp typography, diagrams and consistent iconography."]: bullet(doc,z)
doc.add_heading("Suggested visual system",2)
make_table(doc,["Element","Recommendation"],[
("Palette","Deep navy for trust, teal for action/technology, warm gold for Mizantra character, off-white backgrounds, red only for risk/errors"),
("Typography","Inter, Manrope or IBM Plex Sans; one family with clear weights"),
("Shapes","Medium-radius cards, 1px borders, compact status pills, restrained shadows"),
("Icons","Simple line icons for documents, stock, people, tools and finance"),
("Imagery","Industrial photography with natural light; UI screenshots in browser frames; workflow diagrams"),
("Motion","Subtle scroll-reveal and line-progress animation; no dashboard clutter")],[1700,7660])

doc.add_heading("9. Regional and language rules",1)
p(doc,"India and UAE must be separate market experiences. The selector must change content, not only the currency symbol.")
make_table(doc,["Area","India edition","UAE edition"],[
("Currency","INR / ₹","AED / د.إ"),
("Tax wording","GST, GSTIN, HSN, e-way bill","VAT, TRN, UAE VAT wording"),
("Date/time","India locale and Asia/Kolkata","UAE locale and Asia/Dubai"),
("Lead routing","India sales/support contact","UAE sales/support contact"),
("Examples","Indian manufacturing and GST examples","UAE industrial, VAT and TRN examples"),
("SEO","India titles, metadata and schema","UAE titles, metadata and schema")],[1600,3880,3880])
p(doc,"Do not place Indian GST/HSN examples on the UAE edition, or UAE VAT/TRN claims on the India edition. Persist the selected edition while allowing a deliberate switch.")

doc.add_heading("10. Screenshots and diagrams to prepare",1)
for z in ["Executive dashboard: approvals, open PO exposure, stock risk, production WIP and action queue.","Manufacturing/subcontracting order grid: status, planned vs actual, issue, receive/QC and document trail.","Document trail: one order view showing every linked document and financial state.","Sales quotation: revisions, customer comments, HSN, tax, discount, terms and PDF/email actions.","Service ticket: asset serial, technician, customer contact, evidence, parts, warranty and billing.","Stock trail: receipt, issue, transfer, adjustment and current balance."]: bullet(doc,z)
p(doc,"Each screenshot needs a short caption explaining the business outcome. Never show screenshots without context or unreadably small text.")

doc.add_heading("11. Lead forms and information architecture",1)
make_table(doc,["Form","Required fields","After submission"],[
("Book a demo","Name, company, work email, country/edition, industry, employee range, modules, message","Thank-you page, calendar option and sales notification"),
("Workflow review","Name, company, country, current process, biggest control gap, preferred contact","Route to solutions consultant; show response expectation"),
("General enquiry","Name, company, email, category, message, attachment optional","Ticket/reference number and email confirmation")],[1900,4100,3360])
p(doc,"Keep forms short. Use progressive disclosure for process detail. Country/edition should be collected early so routing and content remain separated.")

doc.add_heading("12. Trust, proof and compliance content",1)
for z in ["Explain role-based permissions, maker-checker approvals and audit trails in plain language.","Show how stock changes are tied to documents and users.","State that regional tax/statutory configuration is edition-specific and validated during implementation.","Include implementation: discovery, master data, configuration, migration, training, pilot and go-live support.","State uptime/support commitments only when commercially approved; never invent certifications.","Add privacy policy, terms, cookie notice and data-processing contact before launch."]: bullet(doc,z)

doc.add_heading("13. Responsive and accessibility requirements",1)
for z in ["Desktop-first for operations teams, complete tablet/mobile layouts for technicians and approvals.","Use 12-column desktop, 8-column tablet and 4-column mobile grids.","Keep Book a Demo and edition selector reachable on mobile.","Turn tables into cards or horizontal data regions; never compress labels into unreadable text.","Minimum body text 16px on mobile, strong contrast, keyboard focus and labelled fields.","Give meaningful alt text to screenshots/diagrams; mark decorative icons appropriately.","Do not rely on colour alone; pair statuses with labels such as OPEN, IN PROCESS, COMPLETED and PENDING QC."]: bullet(doc,z)

doc.add_heading("14. SEO and content requirements",1)
make_table(doc,["Page type","Suggested search intent"],[
("Home","manufacturing ERP, manufacturing software, ERP for manufacturing companies"),
("Manufacturing","production planning ERP, BOM and work order software, shop floor ERP"),
("Subcontracting","subcontracting ERP, outside processing inventory, vendor WIP tracking"),
("Procurement","purchase to pay ERP, GRN QC inventory, supplier invoice workflow"),
("Service","field service management for manufacturers, service call parts warranty software"),
("UAE edition","UAE manufacturing ERP, VAT-ready ERP UAE, ERP for Dubai manufacturers"),
("India edition","India manufacturing ERP, GST-ready manufacturing software, HSN inventory ERP")],[2300,7060])
p(doc,"Use descriptive titles, one H1 per page, structured data for Organization/Product/FAQ, fast images, clean URLs and internal links between modules and use cases.")

doc.add_heading("15. Content tone and copy rules",1)
for z in ["Use direct operational language: “Issue raw material”, “Approve QC”, “See the full trail”.","Prefer outcomes and controls over vague claims such as “revolutionary” or “seamless”.","Write for busy operators: short paragraphs, strong subheads and clear verbs.","Explain ERP terms at first use; do not assume visitors know SIV, GRN, WIP or AP.","Use “Mizantra” consistently; keep it distinct from the parent company unless brand architecture says otherwise.","Do not claim SAP certification or official SAP equivalence."]: bullet(doc,z)

doc.add_heading("16. Designer deliverables",1)
for z in ["Low-fidelity sitemap and responsive wireframes for top-level pages.","High-fidelity desktop/tablet/mobile designs for Home, Platform, Manufacturing, Subcontracting, Sales, Service, India and UAE editions.","Design system: colours, type, spacing, buttons, forms, cards, tables, status badges, icons and error states.","Workflow diagrams for procure-to-pay, subcontracting, order-to-cash and service-to-cash.","Screenshot presentation templates and image treatment guidance.","Prototype for edition switching, Book a Demo form, navigation and document-trail interaction.","Handoff package with Figma components, naming conventions, spacing tokens, accessibility notes and exports."]: bullet(doc,z)

doc.add_heading("17. Acceptance checklist",1)
for z in ["A first-time visitor can explain Mizantra within 10 seconds.","A manufacturing buyer can find a relevant workflow within two clicks.","India and UAE pages do not mix currency, tax, statutory or contact content.","Every module page shows a problem, workflow, controls and CTA.","The website works at 1440px, 1024px, 768px and 390px widths.","Forms show validation, success, duplicate prevention and next steps.","Screenshots are legible and captions explain outcomes.","Core pages pass accessibility, performance, metadata and broken-link checks."]: bullet(doc,z)

doc.add_heading("18. Recommended launch sequence",1)
for z in ["Approve brand direction, sitemap and India/UAE information architecture.","Design Home, Platform overview and regional switcher.","Design Manufacturing, Procurement, Sales and Service workflow pages.","Add screenshots, proof, forms, resources and legal pages.","Run content, accessibility, responsive, SEO and analytics QA.","Launch a controlled version, monitor enquiries and iterate based on buyer questions."]: num(doc,z)

doc.add_heading("Appendix: one-paragraph product description",1)
p(doc,"Mizantra is a cloud ERP for manufacturing and industrial service businesses. It connects procurement, inventory, production, subcontracting, quality, sales, service, HR and finance in one controlled operating system. Teams can move from raw-material planning and purchase through issue, production or vendor processing, receipt, inspection, stock update, invoicing and payment while preserving the linked document trail. Regional editions keep India GST/INR workflows and UAE VAT/AED workflows clearly separated.")
doc.core_properties.title="Mizantra Website Design & Content Brief"; doc.core_properties.subject="Designer brief for Mizantra ERP website"; doc.core_properties.author="SAK ERP"; doc.core_properties.keywords="Mizantra, ERP, manufacturing, website design, India, UAE"
OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT); print(OUT)

