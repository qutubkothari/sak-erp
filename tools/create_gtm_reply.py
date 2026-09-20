from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path("GTM_Topic_3_Response.docx")
BLUE, NAVY, GRAY = RGBColor(46,116,181), RGBColor(31,77,120), RGBColor(89,89,89)

def set_font(run, size=11, color=None, bold=None, italic=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    if color: run.font.color.rgb = color
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def shade(cell, fill):
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(node)

def margins(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, value in [("top",130),("start",160),("bottom",130),("end",160)]:
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)

doc = Document()
section = doc.sections[0]
section.top_margin, section.bottom_margin = Inches(.85), Inches(.8)
section.left_margin, section.right_margin = Inches(1), Inches(1)

normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Arial", Pt(11)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(6), 1.15
for name, size, color, before, after in [("Heading 1",16,BLUE,12,6),("Heading 2",13,NAVY,10,5)]:
    style = doc.styles[name]
    style.font.name, style.font.size, style.font.color.rgb, style.font.bold = "Arial", Pt(size), color, True
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("SAK ERP | GTM Strategy"), 9, GRAY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
set_font(p.add_run("RESPONSE"), 10, BLUE, True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
set_font(p.add_run("Industry-wise Pain Points & SAK ERP Solutions"), 22, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
set_font(p.add_run("Topic 3: Proposed vertical GTM direction"), 12, GRAY, italic=True)

for label, value in [("Subject", "Approval of the vertical-specific sales-story approach"),("Status", "Aligned — proceed with Metal Fabrication & Machining first")]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(f"{label}: "), 10.5, bold=True)
    set_font(p.add_run(value), 10.5)
doc.add_paragraph()

p = doc.add_paragraph()
set_font(p.add_run("Thank you for laying out a clear direction for Topic 3. "), bold=False)
set_font(p.add_run("I agree with the proposed approach: our go-to-market message should lead with the operational problems each vertical faces, rather than a generic list of ERP modules."), bold=True)
p = doc.add_paragraph()
set_font(p.add_run("This will make SAK ERP easier for prospects to understand because the conversation starts with their day-to-day challenges and shows a direct path to control, visibility, and measurable business improvement."))

doc.add_paragraph("What I support", style="Heading 1")
for heading, body in [
    ("Vertical-specific story", "Build a distinct narrative for each priority industry, anchored in the customer’s operational and financial pain points."),
    ("Consistent framework", "Use the proposed Pain → Consequence → SAK Intervention → Proof → Outcome model so sales, marketing, and product demonstrations use the same logic."),
    ("Focused launch", "Prioritize Metal Fabrication & Machining for the first UAE campaign because outside processing and material-to-money traceability are compelling SAK ERP differentiators."),
    ("Repeatable expansion", "Use the validated message, proof points, and campaign assets as the foundation for the other four verticals."),
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(f"{heading}: "), bold=True)
    set_font(p.add_run(body))

doc.add_paragraph("Recommended next deliverable", style="Heading 1")
p = doc.add_paragraph()
set_font(p.add_run("The first deliverable should be a Metal Fabrication & Machining vertical playbook that answers all ten questions in the proposed structure. It should include:"))
for text in [
    "The target customer profile and the language management uses to describe its operational challenges.",
    "The specific pain points around outside processing, material leakage, WIP visibility, multiple UOMs, production cost, and quality control.",
    "A mapped SAK ERP intervention, supported by realistic proof points and a focused demonstration flow.",
    "Sales talk tracks, campaign messaging, landing-page copy, cold-outreach angles, and measurable outcome statements.",
]:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text))

table = doc.add_table(rows=1, cols=1)
table.alignment, table.autofit = WD_TABLE_ALIGNMENT.LEFT, False
cell = table.cell(0,0)
cell.width, cell.vertical_alignment = Inches(6.5), WD_CELL_VERTICAL_ALIGNMENT.CENTER
shade(cell, "EAF2F8"); margins(cell)
p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
set_font(p.add_run("Strategic guardrail: "), 10.5, NAVY, True)
set_font(p.add_run("Keep outcome claims specific and supportable. Each promise should connect to a demonstrable workflow and, where possible, a customer-approved proof point or baseline metric."), 10.5, NAVY)

doc.add_paragraph()
p = doc.add_paragraph()
set_font(p.add_run("With this focus, we can establish a credible, repeatable industry message before extending it across the remaining priority verticals."))
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
set_font(p.add_run("Regards,"))
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
set_font(p.add_run("SAK ERP Team"), bold=True)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("SAK ERP — Topic 3 Response"), 8.5, GRAY)
doc.save(OUT)
print(OUT.resolve())

