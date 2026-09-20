from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "documents"
OUT_PATH = OUT_DIR / "Mizantra_Demo_and_Learning_Kit.docx"

INK = "2B2118"
BROWN = "5B3A29"
GOLD = "9A7748"
PALE_GOLD = "F6EFE4"
PALE_BLUE = "E8EEF5"
NAVY = "27445F"
MUTED = "6E6258"
WHITE = "FFFFFF"
GREEN = "DDF4E8"
RED = "FBE5E5"
LINE = "D8C9B6"


MODULES = [
    {
        "name": "Dashboard, MIS & Approvals",
        "audience": "Owners, directors, functional managers and approvers",
        "outcome": "Start from role-relevant KPIs, exceptions and pending approvals rather than long transaction lists.",
        "screens": [
            ("Dashboard", "/dashboard", "Operational landing page and action reminders."),
            ("Business Pulse", "/dashboard/reports/executive/overview", "Executive cross-functional snapshot."),
            ("Finance Manager / Accountant", "/dashboard/reports/finance/manager | /accountant", "Role-specific liquidity, payable and daily-control reporting."),
            ("Sales Manager / Territory / Executive", "/dashboard/reports/sales/...", "Pipeline and performance reporting by responsibility."),
            ("Operations Control", "/dashboard/reports/operations/control", "Procurement, production, inventory and quality exceptions."),
            ("Manager Approvals", "/dashboard/manager", "Central approval worklist for controlled documents."),
        ],
        "flow": "Role login > KPI/exception > drill into source document > approve or return > audit trail",
        "practice": "Open one KPI, trace it to the underlying record, and complete one non-financial approval in the test tenant.",
        "demo": "Emphasize that dashboards are stakeholder-specific and that operational detail remains one click away.",
    },
    {
        "name": "Ask Mizantra / Active Planner",
        "audience": "All authorized users",
        "outcome": "Use natural language to answer tenant-scoped questions and prepare governed workflow drafts.",
        "screens": [
            ("Active Planner", "/dashboard/active-planner", "Conversational analysis and draft preparation."),
            ("Saved conversations", "Active Planner selector", "Return to prior requests and context."),
            ("System-checked preview", "Active Planner right panel", "Shows interpretation, missing information and safe next action."),
        ],
        "flow": "User request > semantic interpretation > permission/data check > answer or missing-detail question > preview > native workflow",
        "practice": "Ask one read-only question, one misspelled question and one draft-creation request; verify that approvals are not bypassed.",
        "demo": "Use business language, not menu names. Show one typo-tolerant query and one controlled hand-off to a native screen.",
    },
    {
        "name": "CRM",
        "audience": "Sales heads, territory managers, sales executives and business-development teams",
        "outcome": "Capture, assign, qualify and progress enquiries until they become customers and commercial opportunities.",
        "screens": [
            ("CRM Overview", "/dashboard/crm", "Pipeline health and work priorities."),
            ("Lead Pipeline", "/dashboard/crm?view=pipeline", "Stage-based opportunity board."),
            ("All Leads", "/dashboard/crm?view=leads", "Searchable lead register."),
            ("Follow-ups", "/dashboard/crm?view=followups", "Due and overdue activities."),
            ("Assignment Rules", "/dashboard/crm?view=rules", "Rule-based lead ownership."),
            ("Customers", "/dashboard/sales?tab=customers", "Handoff into the order-to-cash master."),
        ],
        "flow": "WhatsApp/manual enquiry > deduplication > lead assignment > qualification > follow-up > customer/quotation",
        "practice": "Create a lead, assign it, schedule a follow-up, move its stage and convert it into the sales flow.",
        "demo": "Use a fresh enquiry and explain that repeat WhatsApp messages should remain on the same CRM timeline.",
    },
    {
        "name": "Sales & Order-to-Cash",
        "audience": "Sales, commercial, dispatch, billing and collections teams",
        "outcome": "Control the full customer document chain from quotation through collection and return.",
        "screens": [
            ("Sales Overview", "/dashboard/sales", "Order-to-cash control centre."),
            ("Customers", "/dashboard/sales?tab=customers", "Customer master and commercial context."),
            ("Quotations", "/dashboard/sales?tab=quotations", "Offers, revisions and approval."),
            ("Sales Orders", "/dashboard/sales?tab=orders", "Approved demand and execution status."),
            ("Fulfilment / Dispatch", "/dashboard/sales?tab=fulfilment | dispatch", "Availability, release, pick/pack and dispatch."),
            ("Billing / Collections", "/dashboard/sales?tab=billing | collections", "Invoice and receipt follow-through."),
            ("Returns / Warranties", "/dashboard/sales?tab=returns | warranties", "Reverse flow and coverage."),
            ("Logistics Control", "/dashboard/sales/logistics-control", "Transport and delivery monitoring."),
        ],
        "flow": "Customer > quotation > approval > sales order > ATP/release > fulfilment > dispatch > billing > collection/return",
        "practice": "Create a quotation, submit and approve it, convert it to an order and review downstream fulfilment status.",
        "demo": "Keep one approved quotation and one order ready so the meeting is not dependent on live approvals.",
    },
    {
        "name": "Procurement & Source-to-Pay",
        "audience": "Requestors, buyers, procurement managers, stores, quality and accounts payable",
        "outcome": "Purchase materials and services through auditable demand, sourcing, approval, receipt and settlement controls.",
        "screens": [
            ("Procurement Overview", "/dashboard/purchase", "Purchase pipeline and workload."),
            ("Vendors", "/dashboard/purchase/vendors", "Supplier master and contact/commercial data."),
            ("Purchase Requisitions", "/dashboard/purchase/requisitions", "Internal demand and approval."),
            ("Purchase Orders", "/dashboard/purchase/orders", "Supplier commitment, terms and printable PO."),
            ("Goods Receipt (GRN)", "/dashboard/purchase/grn", "Material receipt, item lines and QC gate."),
            ("Service Entry Sheets", "/dashboard/purchase/service-entries", "Acceptance of service-only POs."),
            ("Supplier Invoices / Payables", "/dashboard/accounts/supplier-invoices | /payables", "Invoice control, open balances and payment."),
            ("Spend / Sourcing / Contracts", "/dashboard/purchase/spend-intelligence | strategic-sourcing | contracts", "Strategic procurement analysis and governance."),
            ("Import Files / Debit Notes", "/dashboard/purchase/import-files | debit-notes", "Bulk intake and supplier recovery."),
        ],
        "flow": "Need > PR > approval > RFQ/comparison > PO > GRN or service entry > QC/acceptance > supplier invoice > payment",
        "practice": "Complete both a material route (GRN) and a service route (service entry); confirm an empty GRN cannot proceed.",
        "demo": "Explain the material-versus-service branch and show that rejected/active receipts protect PO changes.",
    },
    {
        "name": "Inventory, Warehousing & Working Capital",
        "audience": "Stores, warehouse managers, planners, production and finance",
        "outcome": "Maintain reliable stock, governed movements, replenishment and valuation visibility.",
        "screens": [
            ("Inventory Overview", "/dashboard/inventory", "Stock summary, movements and alerts."),
            ("Stock Master", "/dashboard/inventory/items", "Item-level stock and master information."),
            ("Low Stock Planning", "/dashboard/inventory/low-stock", "Replenishment analysis and supplier-grouped PR creation."),
            ("Warehouse Control / Optimization", "/dashboard/inventory/warehouse-control | warehouse-optimization", "Warehouse execution and improvement."),
            ("Working Capital & SLOB", "/dashboard/inventory/working-capital", "Slow/non-moving exposure and capital control."),
            ("Stock Adjustments", "/dashboard/inventory/stock-adjustments", "Controlled correction with traceability."),
            ("SIV / SRV", "/dashboard/inventory/siv | srv", "Material issue and return vouchers."),
        ],
        "flow": "Receipt/production > warehouse stock > reservation/issue > transfer/return/adjustment > valuation and alerts",
        "practice": "Trace one item from receipt through movement history, issue and return, then review low-stock planning.",
        "demo": "Show a single item trace instead of scrolling through a large stock register.",
    },
    {
        "name": "Production Planning & Execution",
        "audience": "Production planners, plant managers, supervisors, stores and quality",
        "outcome": "Translate demand into feasible material and capacity plans, then control execution and variance.",
        "screens": [
            ("Production Overview", "/dashboard/production", "Plant execution snapshot."),
            ("Create / View Job Orders", "/dashboard/production/job-orders/smart-items | job-orders", "Controlled production orders."),
            ("MRP", "/dashboard/production/mrp", "Time-phased demand, supply, netting and recommendations."),
            ("Smart Planning / Control Tower", "/dashboard/production/smart-planning | planning-control-tower", "Plan coordination and exception review."),
            ("MRP & APS Configuration", "/dashboard/production/planning-configuration", "Planning policies and parameters."),
            ("Demand & S&OP / Capacity", "/dashboard/production/demand-planning | capacity-planning", "Demand alignment and resource feasibility."),
            ("Shop Floor / Work Stations", "/dashboard/shop-floor | work-stations", "Execution reporting by operation and resource."),
            ("OEE / Autonomy", "/dashboard/production/oee | autonomy", "Loss analysis and controlled automation."),
            ("Engineering Changes / Maintenance", "/dashboard/production/engineering-changes | maintenance", "Change governance and asset availability."),
        ],
        "flow": "Demand > MRP/APS > planner review > job order > material issue > operation reporting > QC > finished receipt",
        "practice": "Run or inspect an MRP plan, trace pegging to demand, release one recommendation and follow the job order.",
        "demo": "Use one customer demand line to connect sales, MRP, procurement, shop floor, quality and stock.",
    },
    {
        "name": "BOM & Engineering",
        "audience": "Engineering, planning, production and quality",
        "outcome": "Define what is made, how it is routed and how revisions affect controlled execution.",
        "screens": [
            ("BOM", "/dashboard/bom", "Product structure and revision governance."),
            ("BOM Routing", "/dashboard/bom/[id]/routing", "Operations, work centres and sequence."),
            ("Engineering Changes", "/dashboard/production/engineering-changes", "Approval and effectivity of changes."),
        ],
        "flow": "Item > BOM revision > routing > approval/effectivity > MRP and job-order consumption",
        "practice": "Open one BOM, explain its inputs and routing, and identify which revision is effective for a job order.",
        "demo": "Show traceability from a finished product to one component and one operation.",
    },
    {
        "name": "Subcontracting / Outside Processing",
        "audience": "Production, procurement, stores, quality and accounts payable",
        "outcome": "Treat outside processing as a controlled supplier operation with material, receipt, QC and financial traceability.",
        "screens": [
            ("Subcontracting Workspace", "/dashboard/production/subcontracting", "Routes, orders, issue, receipts, GRN register and document trail."),
            ("Material Outward / Stores Slip", "Subcontract order actions", "Printable controlled issue instruction and challan."),
            ("Receive Material / GRN", "Subcontract receipt workspace", "Return of processed material and receipt reference."),
            ("Supplier Payables", "/dashboard/accounts/payables", "Subcontract invoices remain in normal supplier payables with source controls visible."),
        ],
        "flow": "Requirement/route > approval > subcontract order > issue material > vendor WIP > receive/GRN > QC > service invoice > supplier payable",
        "practice": "Create an order, print the issue slip, post outward material, receive output, verify GRN and locate the payable.",
        "demo": "Stress that raw material remains traceable while at the vendor and finance is not isolated in a separate ledger.",
    },
    {
        "name": "Quality Control",
        "audience": "Quality engineers, inspectors, procurement, production and management",
        "outcome": "Plan inspections, record results, prevent uncontrolled stock posting and close root causes.",
        "screens": [
            ("Quality Overview", "/dashboard/quality", "Inspection and NCR workload."),
            ("Inspections", "/dashboard/quality?tab=inspections", "Execution and acceptance/rejection."),
            ("Inspection Plans", "/dashboard/quality/inspection-plans", "Characteristics, sampling and controls."),
            ("NCR / Supplier Quality", "/dashboard/quality?tab=ncr | vendors", "Non-conformance and supplier performance."),
            ("CAPA & Supplier Recovery", "/dashboard/quality/capa", "Corrective action and commercial recovery."),
            ("EHS / Cost of Quality", "/dashboard/quality/ehs-sustainability | cost-of-quality", "Compliance and loss visibility."),
        ],
        "flow": "Inspection trigger > plan/checklist > result > accept/reject > NCR > CAPA/recovery > closure",
        "practice": "Open a GRN inspection, record a result, reject one sample in test and follow the resulting control status.",
        "demo": "Show how QC gates stock and payable readiness without bypassing the native workflow.",
    },
    {
        "name": "Accounts & Finance",
        "audience": "CFO, finance managers, accountants, treasury and commercial teams",
        "outcome": "Connect operating documents to payable, receivable, cash, margin, statutory and financial-control views.",
        "screens": [
            ("Accounting / Cost & Margin", "/dashboard/accounts | costing", "Accounting control centre and profitability."),
            ("Collections / Payment Runs", "/dashboard/accounts/collections | payment-runs", "Receivable follow-up and governed supplier settlement."),
            ("Accounts Payable", "/dashboard/accounts/payables", "Outstanding, overdue and open advances by supplier."),
            ("Cash Forecast / Treasury & FX", "/dashboard/accounts/cash-forecast | treasury-control", "Liquidity and currency exposure."),
            ("Bank Reconciliation", "/dashboard/accounts/bank-reconciliation", "Bank-to-ledger matching."),
            ("Budgets / FP&A", "/dashboard/accounts/budgets | fpna-control", "Plan, scenario and variance management."),
            ("IFRS / IAS Controls", "/dashboard/accounts/lease-accounting | revenue-recognition | ecl-control | provision-control", "Specialized accounting governance."),
            ("Fixed Assets / Statutory / UAE", "/dashboard/accounts/fixed-assets | statutory-returns | uae-compliance", "Asset and compliance control."),
            ("Consolidation / Opening Balances", "/dashboard/accounts/consolidation | opening-balances", "Group and implementation controls."),
        ],
        "flow": "Operational source > approval/receipt > accounting event > payable/receivable > settlement > reconciliation > MIS/compliance",
        "practice": "Trace one PO/GRN invoice to payables, verify due-date logic and advance availability, then review payment control.",
        "demo": "Never present totals without drilling to source documents; explain payment terms and due-date calculation.",
    },
    {
        "name": "Projects",
        "audience": "Project managers, operations, finance and management",
        "outcome": "Connect commercial and operational activity to project cost, margin and earned-value visibility.",
        "screens": [
            ("Project Master", "/dashboard/projects", "Project ownership and execution workspace."),
            ("Margin & EVM Control", "/dashboard/projects/performance", "Planned versus earned value and margin risk."),
        ],
        "flow": "Project setup > linked sales/procurement/production > cost and progress capture > margin/EVM review",
        "practice": "Open a project and identify its customer, commercial value, cost drivers and current variance.",
        "demo": "Use only if the audience manages projects; otherwise mention it as an integrated extension.",
    },
    {
        "name": "Service, Installed Base & Warranty",
        "audience": "Service coordinators, technicians, sales, warranty and billing teams",
        "outcome": "Maintain the customer asset lifecycle after delivery and commissioning.",
        "screens": [
            ("Service Tickets", "/dashboard/service?tab=tickets", "Issue intake, priority and resolution."),
            ("Dispatch Board / Technicians", "/dashboard/service?tab=dispatch | technicians", "Resource assignment and field execution."),
            ("Installed Base", "/dashboard/service?tab=installed-base", "Commissioned customer assets and UID relationship."),
            ("Service Contracts", "/dashboard/service?tab=contracts", "Entitlement and chargeability."),
            ("Maintenance / Warranty Check", "/dashboard/service?tab=maintenance | warranty-check", "Preventive work and coverage validation."),
            ("Service Billing / Reports", "/dashboard/service?tab=billing | reports", "Commercial closure and service performance."),
        ],
        "flow": "Dispatch/UID > confirmed installation or commissioning > installed asset > warranty/contract > ticket > technician > resolution/billing",
        "practice": "Confirm a commissioning event creates/updates an installed asset, then raise a ticket and verify entitlement.",
        "demo": "Explain that asset creation should follow confirmed commissioning, not an unconfirmed shipment.",
    },
    {
        "name": "HR, Attendance & Payroll",
        "audience": "Employees, supervisors, HR and payroll",
        "outcome": "Support employee self-service and management control from attendance through payroll.",
        "screens": [
            ("Employee Self-Service", "/dashboard/hr/employees?tab=attendance", "Check-in, go out, return and end day."),
            ("My Leaves", "/dashboard/hr/employees?tab=leaves", "Employee leave requests and status."),
            ("HR Management", "/dashboard/hr/management?tab=attendance", "Attendance review and reporting."),
            ("Payroll", "/dashboard/hr/management?tab=payroll", "Payroll processing workspace."),
            ("Skills & Capacity Risk", "/dashboard/hr/workforce-skills", "Workforce capability and dependency risk."),
        ],
        "flow": "Employee master > daily attendance/movement > leave and exceptions > approval > payroll > report",
        "practice": "Complete check-in, go out, return and end day on mobile; confirm the report includes date, day and notes source.",
        "demo": "Use a prepared attendance day; do not rely on location/mobile permissions during a boardroom demo.",
    },
    {
        "name": "UID & Document Traceability",
        "audience": "Stores, production, quality, dispatch, service and auditors",
        "outcome": "Trace serialized or uniquely identified material across its complete lifecycle and supporting evidence.",
        "screens": [
            ("UID Management", "/dashboard/uid", "UID creation and current state."),
            ("Trace UID", "/dashboard/uid/trace", "End-to-end event and document history."),
            ("Deployment", "/dashboard/uid/deployment", "Customer deployment/installation context."),
            ("Documents", "/dashboard/documents", "Central document repository and references."),
            ("Audit Trails", "/dashboard/audit-trails", "Who did what and when."),
        ],
        "flow": "Receipt/manufacture > UID > movement/QC > dispatch > deployment > installed asset/service > audit evidence",
        "practice": "Search one UID and narrate every event, source document, location and ownership change.",
        "demo": "UID trace is a strong closing proof because it connects multiple modules in one record.",
    },
    {
        "name": "Administration, Security & Integrations",
        "audience": "Super administrators, tenant administrators, IT and control owners",
        "outcome": "Configure the organization, expose only licensed/relevant screens and maintain segregation and integrations.",
        "screens": [
            ("Users / Roles", "/dashboard/settings?tab=users | roles", "User assignment and action-level permissions."),
            ("Master Feature Access", "/dashboard/settings/feature-access", "Tenant-level module and screen entitlements."),
            ("Organization / Company Header", "/dashboard/settings/organization | company-header", "Legal, structural and document branding."),
            ("Email / Automation", "/dashboard/settings/email-configuration | /dashboard/automation", "Controlled communication and rules."),
            ("Master Data Governance / SoD", "/dashboard/settings/master-data-governance | segregation-of-duties", "Data ownership and conflicting-access prevention."),
            ("Integration Hub", "/dashboard/settings/integration-hub", "Integration health and configuration."),
            ("WhatsApp Business / Automation", "/dashboard/settings/whatsapp | whatsapp/automation", "Connection, inbound CRM capture and governed outbound messages."),
            ("Continuous Controls", "/dashboard/audit-trails/continuous-controls", "Leakage and control monitoring."),
        ],
        "flow": "Tenant entitlement > role permission > user assignment > integration configuration > audit/control monitoring",
        "practice": "Disable one test feature, verify it disappears and is route-blocked, then restore it without affecting data.",
        "demo": "Do not change permissions during the client demo; show a prepared role and explain the two-layer access model.",
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Sequence[int], indent=120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def new_decimal_numbering(doc) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(tabs)
    p_pr.append(ind)
    for node in (start, num_fmt, level_text, level_jc, p_pr):
        level.append(node)
    abstract.append(level)
    numbering.insert(0, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def set_run(run, size=None, color=INK, bold=None, italic=None, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def add_field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run = OxmlElement("w:r")
        run.append(node)
        paragraph._p.append(run)


def add_paragraph(doc, text="", *, size=11, bold=False, color=INK, italic=False, align=None, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    if text:
        set_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text: str, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.375 if level == 0 else 0.625)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text), size=10.5)
    return p


def add_number(doc, text: str, num_id: int):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    p_pr.append(num_pr)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text), size=10.5)
    return p


def heading(doc, text: str, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = level == 1
    set_run(p.add_run(text), size={1: 16, 2: 13, 3: 12}[level], color=NAVY if level < 3 else BROWN, bold=True)
    return p


def callout(doc, title: str, body: str, fill=PALE_GOLD, accent=BROWN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(title + "  "), size=10.5, color=accent, bold=True)
    set_run(p.add_run(body), size=10.5, color=INK)
    add_paragraph(doc, "", after=2)


def add_table(doc, headers: Sequence[str], rows: Iterable[Sequence[str]], widths: Sequence[int], font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(header), size=9.2, color=NAVY, bold=True)
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run(p.add_run(str(value)), size=font_size, color=INK)
    set_table_geometry(table, widths)
    add_paragraph(doc, "", after=1)
    return table


def add_module(doc, module, index: int):
    heading(doc, f"{index}. {module['name']}", 1)
    callout(doc, "Business outcome", module["outcome"], fill=PALE_GOLD)
    add_paragraph(doc, "Primary users", size=10, color=GOLD, bold=True, after=2, keep=True)
    add_paragraph(doc, module["audience"], size=10.5, after=7)
    add_paragraph(doc, "Core process", size=10, color=GOLD, bold=True, after=2, keep=True)
    add_paragraph(doc, module["flow"], size=10.5, bold=True, color=BROWN, after=9)
    add_paragraph(doc, "Screens and purpose", size=10, color=GOLD, bold=True, after=4, keep=True)
    add_table(doc, ["Screen", "Route / location", "What it is for"], module["screens"], [2400, 2940, 4020], font_size=8.8)
    add_paragraph(doc, "What to practise", size=10, color=GOLD, bold=True, after=2, keep=True)
    add_bullet(doc, module["practice"])
    add_paragraph(doc, "How to demonstrate it", size=10, color=GOLD, bold=True, after=2, keep=True)
    add_bullet(doc, module["demo"])


def build_document() -> Document:
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, before, after in ((1, 16, 18, 10), (2, 13, 14, 7), (3, 12, 10, 5)):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY if level < 3 else BROWN)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(hp.add_run("MIZANTRA  |  DEMO & LEARNING KIT"), size=8.5, color=MUTED, bold=True)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(fp.add_run("Confidential - Presenter Edition  |  "), size=8, color=MUTED)
    add_field(fp, "PAGE")

    # Editorial cover
    add_paragraph(doc, "MIZANTRA", size=11, color=GOLD, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=70, after=18)
    add_paragraph(doc, "Demo & Learning Kit", size=30, color=NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_paragraph(doc, "A presenter-ready guide to learn the platform, rehearse the story and deliver a controlled end-to-end ERP demonstration.", size=14, color=BROWN, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
    callout(doc, "Built from the current application", "Module names, routes and workflow descriptions reflect the current Mizantra web application. Visibility may vary by tenant feature access and user role.", fill=PALE_GOLD)
    add_paragraph(doc, "Prepared for", size=9, color=GOLD, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=40, after=2)
    add_paragraph(doc, "Mizantra product demonstrations and internal enablement", size=12, color=INK, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_paragraph(doc, "Version date: 04 September 2026", size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    heading(doc, "How to use this kit", 1)
    add_table(doc, ["If you have...", "Use this path"], [
        ("15 minutes", "Read the Demo Story, Run of Show and Demo-Day Quick Card."),
        ("2 hours", "Add the role map, core workflow chapters and Q&A section."),
        ("1-2 days", "Complete every practice task in the test tenant and conduct two timed rehearsals."),
    ], [1800, 7560], font_size=9.5)
    callout(doc, "Golden rule", "Do not demonstrate every menu. Demonstrate one connected business outcome and open additional screens only when the audience asks.", fill=GREEN, accent=NAVY)

    heading(doc, "Contents", 1)
    contents = [
        "1. Demonstration strategy", "2. Two-day learning plan", "3. Role and stakeholder map",
        "4. The end-to-end Mizantra story", "5. Sixty-minute presenter runbook", "6. Demo data pack",
        "7. Feature handbook by module", "8. Ask Mizantra demonstration guide", "9. Questions and objections",
        "10. Pre-demo verification and recovery", "11. Rehearsal sign-off", "12. Demo-day quick card",
    ]
    for item in contents:
        add_bullet(doc, item)

    heading(doc, "1. Demonstration strategy", 1)
    add_paragraph(doc, "The objective is not to prove that Mizantra has many screens. The objective is to prove that a business event remains connected, permission-controlled and traceable from enquiry to service.", size=11.5, color=BROWN, bold=True, after=10)
    heading(doc, "The five messages to repeat", 2)
    message_num_id = new_decimal_numbering(doc)
    for text in [
        "One connected operational record: commercial, supply, production, quality, finance and service teams work on the same chain.",
        "Role-based simplicity: each user sees the work, KPIs and actions appropriate to that role.",
        "Native controls: AI may interpret, answer and prepare, but approvals, posting, payment and release remain governed.",
        "Traceability by design: source documents, statuses, UID events and audit trails explain every important number.",
        "Tenant-level fit: administrators can enable only the modules and screens relevant to each client.",
    ]:
        add_number(doc, text, message_num_id)
    heading(doc, "What not to do", 2)
    for text in [
        "Do not scroll through long registers without a business question.",
        "Do not create records whose master data or approvals have not been prepared.",
        "Do not claim that AI can approve, post, pay or bypass controls.",
        "Do not open disabled or role-restricted screens in front of the client.",
        "Do not improvise with live WhatsApp, location, email or external services unless they passed the pre-demo check.",
    ]:
        add_bullet(doc, text)

    heading(doc, "2. Two-day learning plan", 1)
    add_table(doc, ["Session", "Duration", "Focus", "Evidence of learning"], [
        ("Day 1 - Session 1", "60 min", "Navigation, roles, dashboards, approvals and feature access", "Explain what each stakeholder sees."),
        ("Day 1 - Session 2", "90 min", "CRM and order-to-cash", "Take a lead to an approved sales order."),
        ("Day 1 - Session 3", "120 min", "Procure-to-pay and inventory", "Complete PR > PO > GRN/service entry > payable."),
        ("Day 1 - Session 4", "90 min", "MRP, production, BOM and quality", "Trace one demand to production and QC."),
        ("Day 2 - Session 1", "75 min", "Subcontracting and supplier finance", "Issue, receive, QC and find the payable."),
        ("Day 2 - Session 2", "75 min", "UID, installed asset, warranty and service", "Trace UID to ticket and entitlement."),
        ("Day 2 - Session 3", "60 min", "Accounts, MIS and Ask Mizantra", "Explain three KPIs and three natural-language queries."),
        ("Day 2 - Session 4", "60 min", "Full timed rehearsal", "Finish the 60-minute script with no dead ends."),
    ], [1650, 950, 2900, 3860], font_size=8.8)
    callout(doc, "Learning method", "For every module, learn four things: business purpose, primary user, source document and next controlled step.", fill=PALE_BLUE, accent=NAVY)

    heading(doc, "3. Role and stakeholder map", 1)
    add_table(doc, ["Stakeholder", "Start here", "Primary scope", "Best demo proof"], [
        ("Owner / CEO / Director", "Business Pulse", "Cross-functional KPIs, exceptions, value and approvals", "KPI drill-down to source record"),
        ("Finance Manager / CFO", "Finance Manager MIS", "Liquidity, AP/AR, margin, treasury, compliance", "Payable due-date and source-document trace"),
        ("Accountant", "Accountant Daily Control", "Invoices, settlements, reconciliation and statutory work", "GRN/invoice/payment chain"),
        ("Sales Manager", "Sales Manager MIS / CRM", "Pipeline, quotation, conversion, fulfilment and collection", "Enquiry to order"),
        ("Territory Manager", "Territory MIS", "Regional pipeline, activity and performance", "Lead assignment and follow-up"),
        ("Sales Executive", "My pipeline / follow-ups", "Owned leads, quotations and customer activity", "Next-best action"),
        ("Procurement Manager", "Procurement Overview", "PR, RFQ, PO, supplier and spend", "Approved PR to PO"),
        ("Stores / Warehouse", "Inventory Overview", "GRN, movements, SIV/SRV and stock", "Item movement trace"),
        ("Production Manager", "Planning Control Tower", "MRP, capacity, job orders and execution", "Demand pegging to job order"),
        ("Quality Manager", "Quality Overview", "Inspection, NCR, CAPA and supplier quality", "QC gate blocking downstream posting"),
        ("Service Manager", "Service workspace", "Installed base, contracts, tickets and technicians", "Commissioning to entitlement"),
        ("HR / Payroll", "HR Management", "Attendance, leave, exceptions and payroll", "Complete daily attendance lifecycle"),
        ("System Administrator", "Settings / Feature Access", "Tenant, roles, screens, integrations and audit", "Feature hidden and route-blocked"),
    ], [1700, 1850, 3220, 2590], font_size=8.4)

    heading(doc, "4. The end-to-end Mizantra story", 1)
    callout(doc, "Primary demonstration story", "ENQUIRY > CRM LEAD > QUOTATION > APPROVAL > SALES ORDER > MRP/PROCUREMENT/PRODUCTION > QUALITY > DISPATCH > BILLING & COLLECTION > INSTALLED ASSET > WARRANTY & SERVICE", fill=PALE_BLUE, accent=NAVY)
    add_paragraph(doc, "Narration", size=10, color=GOLD, bold=True, after=3, keep=True)
    add_paragraph(doc, "A customer enquiry enters through WhatsApp or manual capture. CRM assigns an owner and records follow-ups. The qualified requirement becomes a quotation and, after approval, a sales order. Demand then drives stock allocation, procurement or production. Quality gates receipts and finished output. Dispatch and billing complete the commercial cycle. UID and commissioning create the customer installed base, where warranty, contracts and service tickets continue the lifecycle.", size=11, after=10)
    heading(doc, "Controlled branches", 2)
    add_table(doc, ["Decision", "Route"], [
        ("Material must be purchased", "PR > approval > RFQ/comparison > PO > GRN > QC > supplier invoice > payable"),
        ("Service must be purchased", "PR/PO > service entry sheet > acceptance > supplier invoice > payable"),
        ("Item must be manufactured", "MRP > job order > material issue > operations > QC > finished receipt"),
        ("Operation is outsourced", "Subcontract order > material outward > vendor WIP > receipt/GRN > QC > payable"),
        ("Serialized product is delivered", "UID > dispatch > commissioning > installed asset > warranty/contract > service"),
    ], [2450, 6910], font_size=9.3)

    heading(doc, "5. Sixty-minute presenter runbook", 1)
    add_table(doc, ["Time", "Open", "Demonstrate", "Say / prove"], [
        ("00-05", "Dashboard / Business Pulse", "Role-based KPIs and action reminders", "Mizantra begins with decisions and exceptions, not data entry."),
        ("05-12", "CRM", "New enquiry, assignment, timeline and follow-up", "Every enquiry receives an owner and remains connected to later sales activity."),
        ("12-20", "Sales", "Quotation, approval and sales-order conversion", "Commercial commitments are versioned and controlled."),
        ("20-30", "MRP + Procurement", "Demand pegging, PR, PO and receipt branch", "Demand drives supply; users can trace why a purchase is required."),
        ("30-38", "Production + Quality", "Job order, material, execution and QC gate", "Material and quality status govern what can move next."),
        ("38-43", "Subcontracting", "Issue, vendor WIP, receipt/GRN and payable", "Outside processing has the same document discipline as internal work."),
        ("43-49", "Inventory + UID", "Movement and lifecycle trace", "One UID explains source, state, location and downstream ownership."),
        ("49-54", "Service", "Installed asset, entitlement and ticket", "The relationship continues after delivery."),
        ("54-58", "Ask Mizantra", "Natural-language answer and governed draft", "AI understands business intent but never bypasses native controls."),
        ("58-60", "MIS / close", "Return to outcome and questions", "One connected ERP, simplified by role and intelligence."),
    ], [900, 1700, 3150, 3610], font_size=8.4)
    heading(doc, "Presenter operating pattern", 2)
    presenter_num_id = new_decimal_numbering(doc)
    for text in [
        "State the business question before opening a screen.",
        "Show no more than three fields or KPIs before drilling down.",
        "Identify the source document and next controlled step.",
        "Return to the end-to-end story after every module.",
        "Park deep configuration questions for a scoped follow-up session.",
    ]:
        add_number(doc, text, presenter_num_id)

    heading(doc, "6. Demo data pack", 1)
    add_paragraph(doc, "Prepare and verify these records in the test tenant before the rehearsal. Replace the placeholders with real test references.", size=10.5, after=8)
    add_table(doc, ["Record", "Prepared reference", "Required state", "Fallback"], [
        ("WhatsApp enquiry", "________________", "Inbound and assigned", "Use a pre-captured lead"),
        ("Lead", "________________", "Qualified with follow-up", "Open saved lead"),
        ("Quotation", "________________", "Approved and convertible", "Use approved quotation"),
        ("Sales order", "________________", "Demand visible", "Use existing open order"),
        ("MRP run / recommendation", "________________", "Pegged to sales demand", "Open saved run"),
        ("Purchase requisition", "________________", "Approved", "Use approved PR"),
        ("Purchase order", "________________", "Approved with open quantity", "Use open PO"),
        ("GRN / service entry", "________________", "Valid item lines / accepted service", "Open completed receipt"),
        ("Job order", "________________", "Released or in process", "Open active job"),
        ("Subcontract order", "________________", "Issued and receivable", "Open existing trail"),
        ("UID", "________________", "Dispatched/deployed", "Use known UID"),
        ("Installed asset", "________________", "Commissioned with coverage", "Use existing asset"),
        ("Service ticket", "________________", "Assigned", "Open saved ticket"),
    ], [1900, 1900, 3100, 2460], font_size=8.5)
    callout(doc, "Data safety", "Use demonstration records only. Do not expose client personal data, bank information, employee records or production credentials.", fill=RED, accent="8B1E1E")

    heading(doc, "7. Feature handbook by module", 1)
    add_paragraph(doc, "The following pages are your learning reference. Screen visibility is controlled first by tenant feature access and then by role/user permission.", size=10.5, after=8)
    add_table(doc, ["Module", "Primary users"], [(m["name"], m["audience"]) for m in MODULES], [3400, 5960], font_size=8.5)
    for idx, module in enumerate(MODULES, start=1):
        add_module(doc, module, idx)

    heading(doc, "8. Ask Mizantra demonstration guide", 1)
    add_paragraph(doc, "Use queries that test meaning rather than exact keywords. The result must remain grounded in the tenant's authorized ERP data.", size=10.5, after=8)
    add_table(doc, ["Intent", "Suggested wording", "Expected behaviour"], [
        ("Supplier advances", "Show open supplier advances available by supplier", "Advance availability, not gross outstanding."),
        ("Overdue payables", "Which suppliers have overdue payments and how was due date calculated?", "Due date from document/receipt basis and payment terms; drill to source."),
        ("Latest payment", "Who received our latest supplier payment?", "Most recent authorized settlement with date and reference."),
        ("CRM", "Show new enquiries that still need an owner or follow-up", "Read-only prioritized lead list."),
        ("Production", "whch jobs r late and what is blocking them", "Typo-tolerant delay and constraint analysis."),
        ("Inventory", "what will run out before next week", "Time-aware low-stock or shortage answer."),
        ("Draft action", "Prepare a purchase requisition for the shortage", "Ask only for missing details, then create a preview; no automatic approval."),
    ], [1750, 4200, 3410], font_size=8.8)
    heading(doc, "Live-demo rules", 2)
    for text in [
        "Start with a read-only question whose correct answer you already know.",
        "Use one deliberately misspelled query to demonstrate semantic understanding.",
        "If Mizantra asks for missing data, explain that this is a control, not a failure.",
        "Show the system-checked preview before opening the native workflow.",
        "Never confirm a financial, stock or external-message action in a demo unless it is a disposable test record.",
    ]:
        add_bullet(doc, text)

    heading(doc, "9. Questions and objections", 1)
    add_table(doc, ["Client question", "Recommended answer"], [
        ("Can AI approve or pay automatically?", "No. Ask Mizantra can interpret, analyse and prepare. Approval, posting, payment, release and external communication remain in native controlled workflows."),
        ("Does it understand spelling mistakes?", "The conversational layer is designed around intent and context rather than a fixed regex vocabulary. The answer is still validated against authorized ERP data."),
        ("Can we hide modules we do not buy?", "Yes. Tenant feature access controls modules/screens; roles and users provide a second permission layer. Disabling access does not delete data."),
        ("Can we trace a number on a dashboard?", "Yes. The demonstration should drill from KPI to the source record, status and audit trail."),
        ("Are subcontract invoices separate?", "No. They remain visible in normal supplier payables while retaining the subcontract receipt, QC and service-control context."),
        ("When is an installed asset created?", "After confirmed installation/commissioning, with UID, customer and coverage links. This avoids treating an unconfirmed dispatch as an installed asset."),
        ("Can it work for materials and services?", "Yes. Material POs use GRN and QC; service-only POs use service entry and acceptance before invoicing."),
        ("Is WhatsApp connected to CRM?", "When tenant configuration and inbound capture are enabled, new contacts create/assign leads and repeat messages add to the same timeline."),
    ], [3000, 6360], font_size=9)

    heading(doc, "10. Pre-demo verification and recovery", 1)
    heading(doc, "Day before", 2)
    for text in [
        "Confirm the correct tenant, user role and enabled screens.",
        "Verify every prepared record and copy its reference into the demo data pack.",
        "Test login, Active Planner provider, WhatsApp status, email/PDF generation and key APIs.",
        "Run the full story once using the same laptop, browser and network planned for the meeting.",
        "Prepare a sanitized PDF or screenshot backup for every external dependency.",
    ]:
        add_bullet(doc, text)
    heading(doc, "Thirty minutes before", 2)
    for text in [
        "Restart the browser, sign in and open the required tabs in runbook order.",
        "Set browser zoom to a readable level and close personal tabs and notifications.",
        "Confirm there are no visible 4xx/5xx errors, stale feature menus or failed reminders.",
        "Check the presentation display and mobile screen-sharing method.",
        "Keep this quick card and prepared references beside you.",
    ]:
        add_bullet(doc, text)
    heading(doc, "Recovery language", 2)
    add_table(doc, ["Issue", "What to do", "What to say"], [
        ("External integration slow", "Open prepared captured result", "This integration is asynchronous; here is the resulting ERP record and control trail."),
        ("Approval user unavailable", "Use the prepared approved record", "Approval is deliberately role-separated; I will continue from the approved example."),
        ("AI provider unavailable", "Use native navigation and deterministic workflow", "Core ERP controls remain available even when the AI provider is unavailable."),
        ("Record not visible", "Check tenant/role; use prepared permitted record", "Visibility follows tenant entitlement and user permission."),
        ("Unexpected data", "Stop, do not edit; move to backup", "I will not change live data during a demonstration. Let us continue with the controlled example."),
    ], [1900, 3050, 4410], font_size=8.7)

    heading(doc, "11. Rehearsal sign-off", 1)
    add_table(doc, ["Check", "Pass criteria", "Owner", "Status"], [
        ("Navigation", "All required screens visible; irrelevant screens hidden", "________", "[ ]"),
        ("CRM", "Lead, owner, follow-up and conversion path verified", "________", "[ ]"),
        ("Sales", "Quotation approval and order link verified", "________", "[ ]"),
        ("Procurement", "PR, PO, GRN/service entry and payable verified", "________", "[ ]"),
        ("Production", "Demand, MRP, job order and QC trace verified", "________", "[ ]"),
        ("Subcontracting", "Issue, receipt/GRN, QC and payable verified", "________", "[ ]"),
        ("Service", "Installed asset, coverage and ticket verified", "________", "[ ]"),
        ("Ask Mizantra", "Three known-answer queries and one governed draft pass", "________", "[ ]"),
        ("Documents", "PO/GRN/PDF display and download totals verified", "________", "[ ]"),
        ("Fallback", "Offline screenshots/PDF and prepared records ready", "________", "[ ]"),
        ("Timing", "Full run completes in 55-60 minutes", "________", "[ ]"),
    ], [2400, 4400, 1500, 1060], font_size=8.8)
    add_paragraph(doc, "Rehearsal 1 date/time: _________________________    Result: _________________________", size=10.5, before=10)
    add_paragraph(doc, "Rehearsal 2 date/time: _________________________    Result: _________________________", size=10.5)
    add_paragraph(doc, "Presenter sign-off: _____________________________    Demo date: ______________________", size=10.5)

    heading(doc, "12. Demo-day quick card", 1)
    callout(doc, "Opening line", "Mizantra connects enquiry, commercial commitment, supply, production, quality, finance and service in one controlled record, then adds a natural-language operating layer without weakening approvals.", fill=PALE_GOLD)
    heading(doc, "Ten-screen route", 2)
    route_num_id = new_decimal_numbering(doc)
    for text in [
        "Dashboard / Business Pulse - establish role and priority.",
        "CRM - capture and assign the enquiry.",
        "Sales - approve quotation and show the order.",
        "MRP - explain why supply or production is required.",
        "Procurement - PR, PO and GRN/service entry.",
        "Production / Quality - execution and control gate.",
        "Subcontracting - vendor-held material and receipt trail.",
        "Inventory / UID - end-to-end traceability.",
        "Service - installed base, entitlement and ticket.",
        "Ask Mizantra / MIS - intelligent answer, governed draft and close.",
    ]:
        add_number(doc, text, route_num_id)
    heading(doc, "Four questions for every screen", 2)
    for text in [
        "What business problem does this solve?",
        "Who owns the next action?",
        "Which approval or control applies?",
        "Where does this information flow next?",
    ]:
        add_bullet(doc, text)
    heading(doc, "Close", 2)
    add_paragraph(doc, "Return to the audience's priorities: identify their roles, enable only the modules they need, configure their approval matrix and agree the first end-to-end pilot process.", size=11.5, color=BROWN, bold=True, after=10)
    callout(doc, "Do not over-promise", "State clearly when a feature depends on configuration, master data, an enabled integration or a tenant-specific workflow.", fill=RED, accent="8B1E1E")

    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    core = doc.core_properties
    core.title = "Mizantra Demo & Learning Kit"
    core.subject = "Presenter runbook, feature handbook and demo-day reference"
    core.author = "Mizantra"
    core.keywords = "Mizantra, ERP, demo, training, runbook, feature handbook"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
