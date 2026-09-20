from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "documents"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "203A43"
GOLD = "9A7B48"
BLUE = "2E74B5"
PALE = "E8EEF5"
LIGHT = "F7F4EF"
GRAY = "5B6573"
WIDTH = Inches(6.5)


def font(run, size=11, bold=False, color="1F2937", italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.tcPr.find(qn("w:tcW"))
            if tc_w is not None:
                tc_w.set(qn("w:w"), str(width))
                tc_w.set(qn("w:type"), "dxa")
            set_cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def init_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color, before, after in (("Heading 1", 16, BLUE, 16, 7), ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 11.5, NAVY, 9, 4)):
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
    head = sec.header.paragraphs[0]
    head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = head.add_run("MIZANTRA ERP | PRODUCTION")
    font(r, 8.5, True, GOLD)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Internal working document | September 2026")
    font(r, 8, False, GRAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("MIZANTRA ERP")
    font(r, 10, True, GOLD)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    font(r, 25, True, NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(subtitle)
    font(r, 12, False, GRAY)
    return doc


def para(doc, text, bold_lead=None, style=None, color="1F2937", italic=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        font(r, 10.5, True, color)
        r = p.add_run(text[len(bold_lead):])
        font(r, 10.5, False, color, italic)
    else:
        r = p.add_run(text)
        font(r, 10.5, False, color, italic)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    font(r, 10.5)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    font(r, 10.5)
    return p


def callout(doc, heading, text, fill=LIGHT):
    t = doc.add_table(rows=1, cols=1)
    set_table_widths(t, [9360])
    cell = t.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(heading)
    font(r, 10.5, True, NAVY)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    font(r, 10, False, "384454")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    set_table_widths(t, widths)
    header = t.rows[0]
    set_repeat_table_header(header)
    for c, label in zip(header.cells, headers):
        shade(c, PALE)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        font(r, 9, True, NAVY)
    for data in rows:
        cells = t.add_row().cells
        for c, value in zip(cells, data):
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            font(r, 9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def add_backlog():
    doc = init_doc("Production Development Reference & Roadmap", "MSME-simple operation with enterprise-grade manufacturing intelligence")
    callout(doc, "Purpose", "A living reference for future production-module development. It records the current demo-ready scope, the screw-and-rawl-plug reference scenario, and the next enhancements in priority order.")
    doc.add_heading("1. Product principle", level=1)
    para(doc, "Mizantra should feel simple enough for a small factory team to run from a short daily queue, while producing the controls, evidence, capacity view and recommendations expected from an enterprise manufacturing platform.")
    table(doc, ["Design rule", "Meaning in the product"], [
        ("One obvious next action", "Cockpit and role-based work queues lead the operator, supervisor and planner to the next task."),
        ("Progressive detail", "Planning parameters, BOM detail and advanced capture stay available but do not obstruct normal work."),
        ("Evidence once, reuse everywhere", "Actual production, downtime, QC and tool-use data feed reports, costing and recommendations."),
        ("Controls without friction", "Stock, QC, approval and audit rules remain enforced without turning routine actions into a long form."),
    ], [2520, 6840])
    doc.add_heading("2. Current demo-ready production journey", level=1)
    table(doc, ["Stage", "Current Mizantra experience", "Demo proof"], [
        ("Production Cockpit", "Role-led factory snapshot and exception focus.", "Open Production > Cockpit."),
        ("Plan & job order", "Select product to make; readiness and optional sales-order/BOM detail are progressive.", "Create or show an existing job order."),
        ("Material readiness", "MRP shortage signals and job-focused SIV flow.", "Open Material Requirements, then SIV."),
        ("Shop floor", "My Machine records start/finish, actual output, downtime, scrap and tool evidence.", "Show a representative machine record."),
        ("Receipt & quality", "Job-focused SRV then Quality Workbench for receipt/production inspection.", "Open SRV/QC then Quality."),
        ("Results", "Period views plus action-linked recommendations for downtime, rejects and availability.", "Open Production Results."),
    ], [1800, 4800, 2760])
    doc.add_heading("3. Reference manufacturing scenario: screw + rawl plug", level=1)
    para(doc, "The scenario is a reusable model, not a hard-coded product design. Each customer must configure its own material master, workstations, routings, tools, shifts and policy values.")
    table(doc, ["Area", "Reference facts to configure"], [
        ("Finished kit", "8x60 Screw with Rawl Plug; quantities are in PCS."),
        ("Screw inputs", "Wire-roll input; each roll is approximately 60 kg. Confirm the exact wire grade/diameter and cut allowance by product."),
        ("Blank cutting", "T1: configured 60 screws/min (rated 70), 60-80 mm, max 6 mm; T2: 40/min, 80-160 mm, max 8 mm; CH1: 65/min, 60-100 mm, max 8 mm."),
        ("Threading", "TH1: 40/min, 60-120 mm; TH2: 80/min, 60-160 mm."),
        ("Plating", "3 barrels, each 30 kg/hour; choose in-house or subcontract route."),
        ("Plug moulding", "8x60: 12 cavities x 3 shots/min = 36 plugs/min, 2 g/plug; 8x80: 24 cavities x 3 shots/min = 72 plugs/min, 2.4 g/plug."),
        ("Shift pattern", "Screw line normally 12 hours, expandable to 16; plastic line can run 24 hours."),
        ("Disruptions", "Wire change (20 minutes), mould change, electricity interruption, staffing shortfall, planned maintenance and breakdowns."),
    ], [2100, 7260])
    doc.add_heading("4. Required master data for every customer", level=1)
    table(doc, ["Master-data group", "Must capture", "Why it matters"], [
        ("Item & BOM", "Finished goods, subassemblies, raw materials, UOM, unit weight, yield, scrap and alternates.", "Calculates demand and traceable consumption."),
        ("Routing", "Operation sequence, eligible machines, setup/changeover, run rate, queue/transfer batch and QC gate.", "Creates feasible production timing."),
        ("Machine", "Capacity, size/weight limits, shift calendar, planned maintenance, actual availability and cost/hour.", "Avoids assigning work to unsuitable or unavailable equipment."),
        ("Tools & consumables", "Punch/die/tool compatibility, certified life, expected consumption and actual cycles or kg processed.", "Predicts replacement and measures real tool usage."),
        ("People & calendar", "Skills, minimum crew, holidays, overtime and absence constraints.", "Protects plans from unrealistic labour assumptions."),
        ("Supplier/subcontract", "Lead time, capacity, price, QC performance and backup options.", "Supports reliable make-or-buy decisions."),
    ], [1800, 3900, 3660])
    doc.add_heading("5. Data captured during actual production", level=1)
    bullet(doc, "Start and finish time, machine, operator/crew, job and operation.")
    bullet(doc, "Input quantity/weight, good output, rejected output, scrap, rework and transfer quantity.")
    bullet(doc, "Downtime start/end or duration, reason, evidence and corrective action.")
    bullet(doc, "Wire-roll change count and minutes; production system can calculate expected changes from 60 kg rolls, item weight and planned output once yield is configured.")
    bullet(doc, "Punch/die/tool selection, life counter at issue and close, actual usage, abnormal wear and replacement reason.")
    bullet(doc, "QC result, defects, disposition and any deviation against standard cycle time, yield or consumption.")
    doc.add_heading("6. Roadmap: enhance after the demo", level=1)
    table(doc, ["Priority", "Enhancement", "Definition of done"], [
        ("P0 - demo/near term", "Curated demo data and guided first-use setup", "One clear reference job shows planning, material, production, QC and results with no confusing empty screens."),
        ("P1", "Constraint-aware APS", "Schedules by eligible machine, calendars, shifts, maintenance, tool availability, batch/changeover and subcontract alternative."),
        ("P1", "Tool/consumable intelligence", "Compares expected vs actual consumption; warns before tool-life or chemical-stock risk; measures waste/avoidance."),
        ("P1", "Daily/weekly/monthly production pack", "Product, machine, process, operator/shift, OEE, downtime, rejection, yield, cost and recommendation views."),
        ("P2", "Historical learning", "Uses approved historical actuals to recommend rates, buffers, lead times, overtime and make/subcontract choices; always shows evidence and confidence."),
        ("P2", "IoT ingestion", "Machine events populate the same approved actual-production model while retaining manual fallback and auditability."),
        ("P3", "Advanced engineering control", "Versioned routing/BOM changes with material, customer, work-order, supplier, quality and compliance impact."),
    ], [1260, 3300, 4800])
    doc.add_heading("7. Guardrails", level=1)
    bullet(doc, "Recommendations advise; they never autonomously post inventory, financial or approval transactions.")
    bullet(doc, "Every forecast must expose its source data, assumptions, confidence and owner decision.")
    bullet(doc, "A plan is feasible only when material, machine, tool, people, calendar and quality constraints are evaluated.")
    bullet(doc, "Keep advanced controls behind progressive disclosure; the normal user works from a small daily queue.")
    path = OUT / "Mizantra_Production_Development_Reference_and_Roadmap.docx"
    doc.save(path)
    return path


def add_demo_guide():
    doc = init_doc("Production Module Demo Guide", "Step-by-step walkthrough: 8x60 Screw with Rawl Plug")
    callout(doc, "Demo objective", "Show a small manufacturing team how Mizantra turns one customer requirement into a controlled, visible production journey - without making staff work through a complicated ERP screen.", "FFF8E8")
    doc.add_heading("1. Prepare before presenting", level=1)
    table(doc, ["Check", "What must be ready"], [
        ("Login", "Use the approved Mizantra TEST demo user; do not use live transactions."),
        ("Reference product", "Finished item: 8x60 Screw with Rawl Plug, plus its screw, plug, wire and plastic material masters/BOM/routing."),
        ("Workstations", "T1/TH1 or another representative machine is configured and visible; plating route can be shown as in-house or subcontract."),
        ("Evidence", "At least one representative production record/QC result and one report recommendation are visible."),
        ("Presenter approach", "Show the normal path first. Open advanced detail only when asked."),
    ], [2100, 7260])
    doc.add_heading("2. The story to tell", level=1)
    para(doc, "A customer needs 8x60 screw-and-plug kits. Mizantra checks what must be made, whether material and capacity are ready, issues material to a job, records actual machine work and downtime, completes quality checks, and gives the supervisor an action-focused result view.")
    doc.add_heading("3. Click-by-click demo", level=1)
    steps = [
        ("Open Production Cockpit", "Menu: Production > Cockpit.", "Point to the role-led view: the team sees the next action and exceptions instead of a dense register."),
        ("Show the production plan", "Menu: Production > Plan Production.", "Search 8x60 in the product selector. Explain that the finished kit is selected once; Mizantra derives component, subassembly and material requirements from approved master data."),
        ("Explain feasibility", "Choose the planned quantity and due date; show policy only if asked.", "Mention that planning will consider machine capacity, shift calendar, material safety allowance, lead time, downtime history and cost rules once those master records are configured."),
        ("Open Today’s Orders / Create Job Order", "Menu: Production > Today’s Orders. Use a prepared job or create a representative job.", "The normal screen asks for Product to make. Sales-order linking and detailed BOM/routing are optional, keeping the daily user flow simple."),
        ("Show material readiness", "Menu: Production > Material Requirements.", "Explain that shortages are prioritized so staff procure the critical component rather than over-buying a noncritical one."),
        ("Issue material to the job", "From the job/open request, open SIV. The job-focused SIV appears first.", "Show that wire/plastic are issued against the production demand. Do not post a real issue during the demo unless using labelled demo data."),
        ("Record actual screw production", "Menu: Production > My Machine.", "Use a prepared run for T1 or TH1. Explain Start/Finish, actual input, good output, scrap and downtime. Mention a 60 kg wire roll and 20-minute change time are captured as downtime/consumption evidence."),
        ("Explain tooling and constraints", "On the machine/production record, open advanced detail only if needed.", "Show that compatible punches/dies are tied to the operation and their certified life/actual use can be recorded. T1 cannot take every length or wire thickness; routing eligibility protects the plan."),
        ("Show plating decision", "Open the job route or Outside Processing.", "Explain the planner can use the in-house plating plant (3 barrels x 30 kg/hour) or a controlled subcontract option when capacity or due date makes that safer."),
        ("Receive and inspect", "Open SRV, then Quality.", "SRV is focused on the selected job. In Quality Workbench, choose Inspect a receipt or production output. Explain acceptance/rejection, defect evidence and containment."),
        ("Close with results", "Menu: Production > Results.", "Show daily/weekly/monthly views and the recommendation cards. Link downtime to My Machine, rejection to Quality, and availability to Maintenance."),
    ]
    for heading, action, talk in steps:
        numbered(doc, f"{heading} - {action} {talk}")
    doc.add_heading("4. Use these operational facts in conversation", level=1)
    table(doc, ["Topic", "Reference example"], [
        ("Screw rates", "T1 is configured at 60 screws/min (rated 70); TH1 at 40/min. The plan should select only machines that meet length and thickness limits."),
        ("Plastic output", "8x60 plug mould: 12 cavities x 3 shots/min = 36 plugs/min; it can operate on a 24-hour shift."),
        ("Plating", "In-house capacity is 90 kg/hour across three 30 kg/hour barrels; subcontracting is a controlled alternative."),
        ("Downtime", "Wire-roll replacement takes about 20 minutes. Mould changes, staff availability, power interruption, planned maintenance and breakdowns are recorded as loss reasons."),
        ("Quality", "Good output, rejects, scrap and rework are not hidden: they flow to quality, reporting and recommendations."),
    ], [1900, 7460])
    doc.add_heading("5. Questions and concise answers", level=1)
    table(doc, ["Likely question", "Answer"], [
        ("Why select only the finished kit?", "The daily flow stays simple. The approved product structure determines its screw, plug, material and routing requirements."),
        ("Can we prevent wrong procurement?", "Yes. MRP highlights shortage and priority against a specific job/due date so buying follows constrained demand, not a broad stock guess."),
        ("How is downtime handled?", "The operator records it manually today with time and reason; the same model is ready for future IoT events."),
        ("How is a punch or die controlled?", "It is configured as a compatible tool with certified life and actual usage. Variance is retained as evidence for corrective action."),
        ("Does Mizantra automatically post or buy?", "No. It recommends and prepares controlled work; inventory, financial and approval actions remain governed by the user workflow."),
    ], [2650, 6710])
    doc.add_heading("6. Demo safety", level=1)
    bullet(doc, "Use Mizantra TEST and labelled demo records only.")
    bullet(doc, "Do not post real material issues, receipts, quality acceptance, invoices or payments during the presentation.")
    bullet(doc, "If a detailed screen is empty, return to the Cockpit/Results storyline instead of creating live-looking transactions under pressure.")
    doc.add_heading("7. Closing line", level=1)
    para(doc, "Mizantra gives an MSME one simple operational path, while retaining the evidence needed to improve capacity, downtime, consumption, quality and cash tied up in inventory.", color=NAVY)
    path = OUT / "Mizantra_Screw_Production_Module_Demo_Guide.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(add_backlog())
    print(add_demo_guide())
