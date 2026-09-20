from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("deliverables/GTM_3_Industry_Wise_Pain_Points_Reply.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "5B6573"
LIGHT = "F4F6F9"


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.first_child_found_in("w:tblInd")
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tcW = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tcW.set(qn("w:w"), str(width))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_rule(paragraph, color="D7DBE2"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_font(run, 11, NAVY)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, 13, BLUE, bold=True)
    return p


def add_body(doc, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    set_font(run, 11, NAVY)
    return p


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    header_run = header.add_run("SAK ERP | GTM Topic 3")
    set_font(header_run, 9, MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("Prepared for SAK ERP")
    set_font(footer_run, 9, MUTED)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(6)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("GTM RESPONSE")
    set_font(run, 10, BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("Industry-wise Pain Points & SAK ERP Solutions")
    set_font(title_run, 22, NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle.add_run("Response to GTM Topic 3")
    set_font(subtitle_run, 12, MUTED)
    add_rule(subtitle)

    add_body(doc, "Agreed. The proposed Topic 3 structure is correct and gives us a strong foundation for vertical-specific marketing.")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run("Recommended starting point: ")
    set_font(run, 11, NAVY, bold=True)
    run = p.add_run("Metal Fabrication & Machining in the UAE")
    set_font(run, 11, DARK_BLUE, bold=True)
    run = p.add_run(", where SAK ERP has its clearest differentiation.")
    set_font(run, 11, NAVY)

    for item in [
        "Material-to-money traceability",
        "Raw-material issue and consumption control",
        "Subcontractor and vendor WIP visibility",
        "BOM and production planning",
        "Multi-UOM handling",
        "QC and document traceability",
        "Procurement, inventory, finance and service linkage",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Framework to build")
    add_body(doc, "For this industry, we should build the complete framework:")
    framework = doc.add_paragraph()
    framework.paragraph_format.space_after = Pt(8)
    framework.paragraph_format.left_indent = Inches(0.18)
    framework.paragraph_format.line_spacing = 1.10
    r = framework.add_run("Pain -> Business consequence -> SAK intervention -> Demonstration proof -> Measurable outcome -> Sales message -> Marketing message")
    set_font(r, 11, DARK_BLUE, bold=True)

    add_heading(doc, "First demonstration")
    add_body(doc, "The first demonstration should show the complete operational flow:")
    flow = doc.add_paragraph()
    flow.paragraph_format.space_after = Pt(10)
    flow.paragraph_format.left_indent = Inches(0.18)
    flow.paragraph_format.line_spacing = 1.10
    r = flow.add_run("RFQ/PO -> Raw-material receipt -> Inventory -> Subcontracting issue -> Vendor WIP -> Finished-goods receipt -> QC -> Payable -> Profitability and traceability")
    set_font(r, 11, DARK_BLUE, bold=True)

    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(callout, [9360])
    cell = callout.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("UAE positioning support: ")
    set_font(r, 11, NAVY, bold=True)
    r = p.add_run("AED, VAT, supplier/customer controls, audit trails and management reporting.")
    set_font(r, 11, NAVY)

    add_heading(doc, "Scale the vertical playbook")
    add_body(doc, "Once this vertical playbook is validated, reuse the same structure for:")
    for item in [
        "Electrical and electromechanical manufacturers",
        "Industrial equipment and machinery companies",
        "Marine and defence suppliers",
        "Building products and hardware manufacturers",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Immediate deliverables")
    for item in [
        "Industry pain-point playbook",
        "UAE landing-page messaging",
        "Sales presentation",
        "Demo script",
        "Objection-handling guide",
        "Proof points and ROI calculator",
        "Industry-specific feature matrix",
    ]:
        add_bullet(doc, item)

    close = doc.add_paragraph()
    close.paragraph_format.space_before = Pt(10)
    close.paragraph_format.space_after = Pt(0)
    close.paragraph_format.line_spacing = 1.10
    run = close.add_run("Positioning statement: ")
    set_font(run, 11, NAVY, bold=True)
    run = close.add_run("Avoid generic ERP messaging. SAK ERP should be positioned as a manufacturing control and traceability platform that connects material, production, subcontracting, quality, finance and service in one auditable flow.")
    set_font(run, 11, DARK_BLUE, bold=True)

    doc.core_properties.title = "GTM Topic 3 - Industry-wise Pain Points & SAK ERP Solutions"
    doc.core_properties.subject = "Response for SAK ERP positioning"
    doc.core_properties.author = "SAK ERP"
    doc.save(OUT)


if __name__ == "__main__":
    build()
