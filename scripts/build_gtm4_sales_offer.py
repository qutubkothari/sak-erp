from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('deliverables/SAK_ERP_UAE_Sales_and_Marketing_Offer.docx')
NAVY = '20364C'; BLUE = '2E74B5'; GOLD = '9A7445'; MUTED = '5C6770'; LIGHT = 'F2F4F7'; PALE = 'F7F1E7'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_width(cell, width):
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn('w:tcW'))
    if tcW is None: tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(width)); tcW.set(qn('w:type'), 'dxa')

def set_font(run, size=11, color='000000', bold=None, italic=None):
    run.font.name = 'Calibri'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def para(doc, text='', size=11, color='000000', bold=False, italic=False, style=None, after=6, before=0, align=None):
    p = doc.add_paragraph(style=style); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.10
    if align is not None: p.alignment = align
    r = p.add_run(text); set_font(r, size, color, bold, italic); return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12); p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    r = p.add_run(text); set_font(r, 16 if level == 1 else 13, BLUE if level == 1 else NAVY, True); return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.15
    for r in p.runs: set_font(r, 11)
    p.add_run(text); set_font(p.runs[-1], 11); return p

def numbered(doc, text):
    p = doc.add_paragraph(style='List Number'); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.15
    p.add_run(text); set_font(p.runs[-1], 11); return p

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'; t.autofit = False
    for i, label in enumerate(headers):
        cell = t.rows[0].cells[i]; set_cell_shading(cell, LIGHT); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run(label); set_font(r, 9, NAVY, True)
        if widths: set_cell_width(cell, widths[i])
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths: set_cell_width(cells[i], widths[i])
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            r = p.add_run(str(value)); set_font(r, 9.5, '222222')
    return t

def callout(doc, title, body):
    t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'; t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0,0); set_cell_width(cell, 9360); set_cell_shading(cell, PALE)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3); r = p.add_run(title); set_font(r, 11, NAVY, True)
    p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(2); r = p.add_run(body); set_font(r, 10.5, '3C4043')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def footer(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('SAK ERP | UAE Manufacturing Sales Offer | Confidential'); set_font(r, 8.5, MUTED)

doc = Document(); sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.49); footer(sec)
styles = doc.styles
styles['Normal'].font.name = 'Calibri'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); styles['Normal'].font.size = Pt(11)

# Cover
para(doc, 'SAK ERP', 12, GOLD, True, after=18)
para(doc, 'UAE Sales & Marketing Offer', 30, NAVY, True, after=6)
para(doc, 'A practical customer-acquisition offer for manufacturers seeking tighter material, production, subcontracting, service and financial control.', 14, MUTED, after=24)
table(doc, ['Prepared for', 'Market focus', 'Version'], [['SAK ERP leadership and sales team', 'UAE manufacturing, engineering and service-led businesses', 'GTM Topic 4 | August 2026']], [2500, 4200, 2660])
para(doc, '', after=24)
callout(doc, 'Core market message', 'Find where material, production, subcontracting, service execution and cash are leaking out of the business - then prove the control improvement before committing to a full ERP rollout.')
heading(doc, 'Executive summary', 1)
para(doc, 'SAK should not lead with a generic “buy ERP” message. The UAE offer should begin with a structured Manufacturing Control Assessment, convert the findings into a quantified proof-of-value plan, and then move into a controlled implementation. This creates a lower-risk buying path for owners and operations leaders while differentiating SAK from basic accounting software, disconnected spreadsheets and generic ERP resellers.')
heading(doc, 'Offer architecture at a glance', 2)
table(doc, ['Stage', 'Customer outcome', 'SAK deliverable'], [
    ['1. Attract', 'Recognises a control or profitability problem', 'Industry-specific landing page, diagnostic content and assessment invitation'],
    ['2. Diagnose', 'Understands current leakage and priority processes', 'Manufacturing Control Assessment and executive findings'],
    ['3. Prove', 'Sees SAK against real workflow and data', 'Process map, configured proof-of-value and business case'],
    ['4. Implement', 'Moves priority processes into controlled operations', 'Phased ERP rollout, migration, training and go-live'],
    ['5. Optimise', 'Keeps improving adoption and returns', 'Quarterly control review, automation and expansion roadmap'],
], [1200, 3600, 4560])

heading(doc, '1. Entry offer: Manufacturing Control Assessment', 1)
para(doc, 'The entry offer is a decision-oriented assessment for a manufacturer that suspects losses, delays or visibility gaps but is not yet ready to select an ERP. It should be positioned as an operational and financial control assessment, not as a software demo.')
table(doc, ['Component', 'What SAK does', 'What the prospect receives'], [
    ['Scope', 'Review material receipt, stores, production, subcontracting, sales/service hand-offs, approvals and finance visibility.', 'A current-state control map showing hand-offs, documents and missing controls.'],
    ['Evidence', 'Use a limited sample of transactions, stock movements, reports and stakeholder interviews.', 'A leakage and risk register - with the likely operational and financial impact.'],
    ['Prioritisation', 'Rank issues by business impact, urgency and implementation effort.', 'A 90-day control roadmap and recommended phase-one scope.'],
    ['Commercial bridge', 'Show where SAK workflows can remove or monitor the gaps.', 'A tailored proof-of-value agenda and implementation estimate.'],
], [1450, 3970, 3940])
callout(doc, 'Suggested assessment promise', '“In 10 business days, identify the highest-priority control gaps across material, production, subcontracting and cash visibility - and receive a practical, phased plan to close them.”')
heading(doc, 'Ideal assessment buyer', 2)
for x in ['Owner, Managing Director, COO, Operations Director or Finance Head with authority to sponsor change.', 'Manufacturing or engineering business with meaningful inventory, subcontracting, service operations, custom orders, warranty exposure or multi-stage approvals.', 'A visible trigger: stock mismatch, late job costing, supplier/vendor WIP uncertainty, overdue receivables, poor service visibility, spreadsheet dependence or reporting delays.']:
    bullet(doc, x)

heading(doc, '2. Conversation starters and campaign offers', 1)
para(doc, 'Marketing should lead with a measurable operational outcome rather than feature lists. Each campaign should route to an assessment, a short executive consultation or an industry demonstration.')
table(doc, ['Campaign angle', 'Lead message', 'Best audience', 'Primary CTA'], [
    ['Material control', 'Can you explain every unit from GRN to issue, WIP, return and final stock?', 'Discrete manufacturers and traders with stores complexity', 'Book a material-control review'],
    ['Subcontracting control', 'Know what is with each vendor, what came back, what was accepted and what is payable.', 'Fabrication, machining, anodising and job-work firms', 'Request a subcontracting workflow demo'],
    ['Order-to-cash visibility', 'Connect quotation, sales order, dispatch, invoice, receivable and warranty without blind spots.', 'Project and product businesses', 'Review your order-to-cash controls'],
    ['Field service proof', 'Know who attended, what was done, parts used, evidence captured and what is billable.', 'Equipment manufacturers and service organisations', 'Request a service operations review'],
    ['Finance control', 'Move from operational transactions to auditable journals, AR/AP ageing, banking and period-close discipline.', 'Finance leaders replacing spreadsheets/Tally-only processes', 'Book an accounting-control consultation'],
], [1500, 3000, 3000, 1860])

heading(doc, '3. Discovery offer and qualification process', 1)
para(doc, 'Discovery must establish whether the prospect has a real control problem, a credible sponsor, usable data and a timeframe. It is not a free consulting engagement; it is a structured decision stage.')
table(doc, ['Discovery area', 'Questions to answer', 'Qualification signal'], [
    ['Business model', 'Make-to-order, make-to-stock, project, repair/service, distribution or hybrid?', 'Clear process pattern and identifiable value stream.'],
    ['Operational pain', 'Where do stock, production, vendor, service or collection issues occur today?', 'At least one measurable recurring pain point.'],
    ['Current systems', 'Which tools hold accounting, stock, job, service and approval data?', 'Fragmented systems or insufficient process discipline.'],
    ['Decision path', 'Who owns the business case, budget, process ownership and final decision?', 'Executive sponsor plus operational/finance champions.'],
    ['Change readiness', 'What needs to improve within the next 90-180 days?', 'A meaningful business event, growth plan, audit need or risk trigger.'],
], [1650, 4200, 3510])
heading(doc, 'Discovery outputs', 2)
for x in ['One-page process map from source transaction to financial effect.', 'Prioritised use-case backlog with value hypothesis, business owner and acceptance criteria.', 'Data-readiness view: masters, opening stock/balances, documents and integrations.', 'Decision briefing: recommended scope, implementation phasing, risks, investment range and next action.']:
    bullet(doc, x)

heading(doc, '4. Live proof-of-value demonstration', 1)
para(doc, 'A generic click-through demo should be avoided. The proof-of-value should follow the prospect’s transaction trail and show how SAK controls the operational event, resulting inventory position, approval and financial visibility.')
table(doc, ['Demonstration path', 'Evidence to show', 'Business proof'], [
    ['Procure to pay', 'PR/RFQ/PO, GRN, quality controls, supplier invoice, payable status and document trail.', 'Commitment, receipt and liability are linked.'],
    ['Material to production', 'Stock master, inward/outward movement, work order, material issue, WIP and traceability.', 'Material cannot disappear between store and production.'],
    ['Subcontracting', 'Route, one vendor work order, total RM outward challan, receipt/GRN, QC, payable and order trail.', 'Vendor-held material and outcome remain accountable.'],
    ['Quote to cash', 'Customer, quotation, revisions, sales order, delivery, invoice, receivable and follow-up.', 'Commercial pipeline is connected to cash collection.'],
    ['Service to billing', 'Service request, assigned technician, craft serial number, start/end time, photos, parts, warranty and billing.', 'Field work is captured as evidence and revenue opportunity.'],
    ['Finance control', 'Chart of accounts, journals, bank reconciliation, ageing, suspense and period-close checklist.', 'Management reports have an auditable accounting basis.'],
], [1600, 4200, 3560])
callout(doc, 'Proof-of-value rule', 'Use the customer’s terminology and a representative sample of their own process/data. The desired output is a signed scope decision, not a long product tour.')

heading(doc, '5. Paid engagement and business-case offer', 1)
para(doc, 'For qualified prospects with complex operations, SAK should offer a paid Process & Control Blueprint. The fee is credited against implementation if the project proceeds within an agreed period. This protects SAK’s delivery capacity and gives the customer a concrete decision asset.')
table(doc, ['Workstream', 'Included deliverables', 'Decision value'], [
    ['Process blueprint', 'Future-state process maps, roles, approvals, document trail and exception controls.', 'Confirms operational scope before configuration.'],
    ['Data and migration plan', 'Master-data ownership, cleanup needs, opening data and migration sequence.', 'Reduces go-live uncertainty.'],
    ['Solution design', 'Configured use cases, reports, integrations, security roles and branch/market controls.', 'Prevents scope ambiguity and hidden rework.'],
    ['Business case', 'Baseline, target KPIs, expected savings/time release, implementation investment and adoption measures.', 'Allows management to approve against ROI, not software features.'],
], [1750, 4400, 3210])

heading(doc, '6. ERP implementation offer', 1)
para(doc, 'Implementation should be sold as a phased operating-model rollout. Phase one focuses on the controls that create immediate visibility and transaction discipline; subsequent phases extend depth, automation and analytics.')
table(doc, ['Phase', 'Typical scope', 'Exit criteria'], [
    ['Foundation', 'Tenant/company setup, roles, approval controls, chart of accounts, master data, opening data and document numbering.', 'Controlled access, verified masters and agreed operating procedures.'],
    ['Core operations', 'Purchase, inventory, GRN/QC, production, subcontracting and operational dashboards.', 'Priority transactions run end-to-end with traceability.'],
    ['Commercial and service', 'Customer, quotation, sales order, dispatch, billing, warranty and service management.', 'Order-to-cash and service-to-billing are live.'],
    ['Finance and control', 'Double-entry journals, AR/AP, bank reconciliation, tax, assets, budgets and period close.', 'Management reporting and period-close controls are adopted.'],
    ['Optimisation', 'Automations, exception worklists, AI-assisted analysis, reports and integration expansion.', 'KPIs show sustained improvement and ownership is transferred.'],
], [1500, 4650, 3210])
heading(doc, 'Implementation principles', 2)
for x in ['Phased deployment over “big bang” implementation for operational risk control.', 'Named process owner for every module and measurable acceptance criteria for every phase.', 'Configuration before custom development; customisation only where it protects a true competitive workflow.', 'Training by role and scenario, backed by live document trails and operational checklists.', 'Formal go-live readiness review with data, process, user and support sign-off.']:
    bullet(doc, x)

heading(doc, '7. Ongoing support and optimisation offer', 1)
table(doc, ['Offer layer', 'Cadence', 'Included outcome'], [
    ['Go-live hypercare', 'Daily/weekly during stabilisation', 'Rapid issue triage, adoption support and transaction-control monitoring.'],
    ['Managed support', 'Monthly', 'Incident handling, release management, user guidance and usage reporting.'],
    ['Control optimisation review', 'Quarterly', 'Review exceptions, close-cycle health, ageing, stock risk, service performance and automation candidates.'],
    ['Continuous improvement backlog', 'Quarterly planning', 'Prioritised enhancements tied to measurable business value.'],
    ['Executive business review', 'Quarterly or half-yearly', 'Progress against ROI, adoption, risk reduction and next-phase roadmap.'],
], [1900, 1900, 5560])

heading(doc, '8. Pricing structure', 1)
para(doc, 'Price should be presented as a transparent combination of platform subscription, one-time implementation and optional managed services. Final amounts depend on entities, users, modules, data quality, integrations and rollout complexity. Avoid publishing a single “all-inclusive ERP” price that creates false expectations.')
table(doc, ['Commercial component', 'Pricing basis', 'Customer explanation'], [
    ['Manufacturing Control Assessment', 'Fixed-fee diagnostic; optionally credited against Blueprint/implementation.', 'A low-risk way to obtain a real control roadmap before ERP commitment.'],
    ['Process & Control Blueprint', 'Fixed fee or tightly bounded time-and-materials.', 'Produces the approved scope, migration plan and business case.'],
    ['Platform subscription', 'Per tenant/entity, module bundle and named/role user tier.', 'Predictable recurring software and platform value.'],
    ['Implementation', 'Milestone-based fixed scope with change-control for additions.', 'Payment follows defined delivery outcomes and acceptance criteria.'],
    ['Managed support', 'Monthly service tier with included response/service hours.', 'Ongoing operational assurance after go-live.'],
    ['Integrations/custom development', 'Separate estimate after technical discovery.', 'Only charged when a business-approved requirement needs additional build.'],
], [2200, 3300, 3860])
callout(doc, 'Recommended pricing discipline', 'Publish “starting from” bands only after validating UAE market positioning. Use an assessment or Blueprint to convert complex opportunities into a scoped implementation proposal with named assumptions and change-control boundaries.')

heading(doc, '9. Guarantee and risk-reduction mechanism', 1)
para(doc, 'SAK should reduce buyer risk with transparent delivery controls, not unqualified ROI promises. The offer should make commitments that are observable, contractible and within SAK’s control.')
table(doc, ['Mechanism', 'How it works'], [
    ['Assessment credit', 'Credit a defined portion of the paid Blueprint/assessment fee against implementation when the customer proceeds within the stated validity period.'],
    ['Phase gates', 'Each implementation phase has documented scope, data readiness, user training, acceptance criteria and sign-off.'],
    ['Configured proof-of-value', 'Before full rollout, demonstrate priority workflows against agreed scenarios and sample data.'],
    ['No-surprise governance', 'Weekly delivery status, risk register, decision log, scope-change approval and named executive sponsors.'],
    ['Adoption commitment', 'Provide role-based training and hypercare; customer provides process owners and timely decisions.'],
], [2300, 7060])

heading(doc, '10. Calls to action and sales assets', 1)
table(doc, ['Funnel stage', 'Primary CTA', 'Sales asset'], [
    ['Awareness', 'Take the 10-minute Manufacturing Control Diagnostic.', 'Short diagnostic, pain-point landing page, industry insight post.'],
    ['Interest', 'Book a 30-minute control consultation.', 'Executive one-pager and sample control map.'],
    ['Consideration', 'Request a tailored proof-of-value.', 'Role-based demo agenda and relevant transaction trail.'],
    ['Decision', 'Commission a Process & Control Blueprint.', 'Scope statement, outcomes, sample deliverables and fee-credit terms.'],
    ['Implementation', 'Approve the phased rollout plan.', 'Implementation proposal, milestone plan, governance pack and commercial schedule.'],
], [1600, 3200, 4560])
heading(doc, 'Suggested website copy', 2)
para(doc, 'Headline: “Control every material, job, vendor and service event - from transaction to financial outcome.”', 11, NAVY, True)
para(doc, 'Supporting line: “SAK ERP helps UAE manufacturers replace spreadsheet blind spots with traceable operations, disciplined finance and real-time management control.”', 11)
para(doc, 'Primary CTA: “Book a Manufacturing Control Assessment.”', 11, GOLD, True)
para(doc, 'Secondary CTA: “See a live process trail.”', 11, GOLD, True)

heading(doc, '11. Go-to-market operating cadence', 1)
table(doc, ['Weekly activity', 'Purpose', 'Owner'], [
    ['Target-account outreach', 'Create relevant conversations with qualified manufacturing and service prospects.', 'Sales / founder-led growth'],
    ['Industry content and diagnostic promotion', 'Build credibility around controllable operational pain points.', 'Marketing'],
    ['Discovery and assessment reviews', 'Convert interest into evidence-based opportunities.', 'Solutions / operations'],
    ['Proof-of-value sessions', 'Show the customer’s workflow, not generic software screens.', 'Solutions + product'],
    ['Pipeline review', 'Advance qualified opportunities or disqualify early.', 'Sales leadership'],
    ['Customer success insights', 'Turn adoption and outcome data into proof and referral assets.', 'Customer success'],
], [2600, 4900, 1860])

heading(doc, '12. Success metrics', 1)
table(doc, ['Measure', 'Why it matters', 'Target use'], [
    ['Assessment-to-Blueprint conversion', 'Tests whether the entry offer identifies genuine priority pain.', 'Offer quality and qualification.'],
    ['Blueprint-to-implementation conversion', 'Tests whether discovery produces a credible business case.', 'Sales effectiveness.'],
    ['Time from first meeting to signed scope', 'Indicates decision velocity and process friction.', 'Pipeline forecasting.'],
    ['Phase acceptance and adoption rate', 'Ensures implementation produces usable control, not shelfware.', 'Delivery quality.'],
    ['Customer control KPI improvement', 'Links SAK to material accuracy, vendor WIP, close cycle, ageing, service evidence or margin visibility.', 'ROI proof and case studies.'],
    ['Expansion and retained-support rate', 'Shows whether SAK becomes an operating platform rather than a one-off project.', 'Long-term account health.'],
], [2300, 4360, 2700])

heading(doc, 'Recommended next actions', 1)
for x in ['Create a one-page Manufacturing Control Assessment offer and a short website landing page for UAE manufacturers.', 'Prepare four role-based proof-of-value demo scripts: owner/COO, operations/stores, finance and service head.', 'Define three commercial packages: assessment, Blueprint and phased implementation; validate final AED pricing with market interviews.', 'Build a qualification scorecard and CRM stage definitions aligned to the offer funnel.', 'Create two sector-specific campaign variants first: discrete manufacturing/subcontracting and equipment/service operations.', 'Collect customer-approved evidence, outcomes and screenshots only after consent, then develop case-study and referral assets.']:
    numbered(doc, x)

heading(doc, 'Important positioning boundaries', 1)
para(doc, 'This offer describes the intended UAE go-to-market architecture. Commercial terms, tax treatment, data residency, integrations, service levels and industry compliance requirements must be confirmed in each customer proposal. SAK should not claim guaranteed savings, statutory compliance or integration coverage unless they are documented and contractually agreed for that customer.')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = 'SAK ERP UAE Sales and Marketing Offer'
doc.core_properties.subject = 'GTM Topic 4 - UAE Sales Offer and Customer Acquisition Funnel'
doc.core_properties.author = 'SAK ERP'
doc.save(OUT)
print(OUT)
