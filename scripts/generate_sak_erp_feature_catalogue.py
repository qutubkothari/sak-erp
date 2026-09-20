from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
DOCX_PATH = OUT_DIR / "SAK_ERP_Complete_Feature_Catalogue.docx"

# compact_reference_guide preset, with a named customer_pack brand override.
PAGE_WIDTH = 12240
PAGE_HEIGHT = 15840
MARGIN = 1440
CONTENT_WIDTH = 9360
TABLE_INDENT = 120
CELL_MARGIN_V = 80
CELL_MARGIN_H = 120

BROWN = "4A3020"
GOLD = "A67C3D"
CREAM = "F7F2E9"
BLUE = "1F4E79"
DARK_BLUE = "17365D"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "F3E8D7"
LIGHT_GRAY = "F5F6F7"
MID_GRAY = "667085"
GREEN = "1D6B4F"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=CELL_MARGIN_V, start=CELL_MARGIN_H,
                     bottom=CELL_MARGIN_V, end=CELL_MARGIN_H) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_table_geometry(table, widths: Sequence[int], indent: int = TABLE_INDENT) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)


def set_borders(table, color="D8C8B3", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)
        borders.append(tag)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_field(paragraph, field: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    for node in (begin, instr, separate, text, end):
        run.append(node)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.26 + 0.20 * level)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def add_note(doc: Document, title: str, text: str, fill: str = PALE_GOLD) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH - TABLE_INDENT])
    set_borders(table, "D8C8B3")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BROWN)
    p.add_run(text)


def add_feature_table(doc: Document, rows: Iterable[tuple[str, str]],
                      headers=("Capability", "What it provides"), widths=(3250, 6110)) -> None:
    data = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, list(widths))
    set_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, label in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        keep_with_next(p)
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(BLUE)
    for n, row_data in enumerate(data):
        row = table.add_row()
        set_row_cant_split(row)
        if n % 2:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_GRAY)
        for idx, value in enumerate(row_data):
            p = row.cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if idx == 0:
                run = p.add_run(value)
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(BROWN)
            else:
                p.add_run(value)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_module(doc: Document, number: str, title: str, purpose: str,
               rows: Sequence[tuple[str, str]], controls: Sequence[str] = ()) -> None:
    doc.add_page_break()
    doc.add_heading(f"{number}  {title}", level=1)
    p = doc.add_paragraph(purpose)
    p.paragraph_format.space_after = Pt(8)
    add_feature_table(doc, rows)
    if controls:
        doc.add_heading("Controls and operational safeguards", level=2)
        for control in controls:
            add_bullet(doc, control)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("252525")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 28, BROWN, 0, 12),
        ("Subtitle", 15, GOLD, 0, 10),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        styles[name].font.name = "Calibri"
        styles[name].font.size = Pt(10.5)

    # customer_pack header: logo/brand at left and document identity at right.
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.48), WD_TAB_ALIGNMENT.RIGHT)
    logo = ROOT / "apps" / "web" / "public" / "pwa-icon-192.png"
    if logo.exists():
        p.add_run().add_picture(str(logo), width=Inches(0.27))
        p.add_run("  ")
    r = p.add_run("SAK ERP")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(BROWN)
    p.add_run("\tCOMPLETE FEATURE CATALOGUE").font.color.rgb = RGBColor.from_string(MID_GRAY)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.48), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("Client presentation • Current implemented scope • 19 August 2026")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MID_GRAY)
    p.add_run("\tPage ")
    add_field(p, "PAGE")


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    # Cover
    doc.add_paragraph().paragraph_format.space_after = Pt(55)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = ROOT / "apps" / "web" / "public" / "pwa-icon-512.png"
    if logo.exists():
        p.add_run().add_picture(str(logo), width=Inches(1.05))
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("SAK ERP")
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Complete Feature Catalogue")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Integrated enterprise operations platform")
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(MID_GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(35)
    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [2500, 4300], indent=1280)
    set_borders(table, "D8C8B3")
    cover_rows = [
        ("Prepared for", "Client presentation and solution review"),
        ("Release scope", "Current implemented ERP capability"),
        ("Document date", "19 August 2026"),
        ("Classification", "Client-facing product information"),
    ]
    for idx, (label, value) in enumerate(cover_rows):
        set_cell_shading(table.cell(idx, 0), PALE_GOLD)
        p = table.cell(idx, 0).paragraphs[0]
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(BROWN)
        table.cell(idx, 1).paragraphs[0].add_run(value)
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    add_note(doc, "Positioning", "SAP-aligned process discipline, document controls and traceability within the SAK ERP application. SAK ERP is an independent product and is not represented as SAP-certified.")

    # Contents and executive summary
    doc.add_page_break()
    doc.add_heading("Contents", level=1)
    contents = [
        "1. Executive overview and solution landscape",
        "2. Executive cockpit, workflow and notifications",
        "3. Procurement and procure-to-pay",
        "4. Inventory and warehouse management",
        "5. Production, BOM, routing and shop-floor control",
        "6. Subcontracting and outside processing",
        "7. Quality management",
        "8. Sales and order-to-cash",
        "9. Customer service and service-to-cash",
        "10. Finance, payables and receivables",
        "11. Human resources, attendance and payroll",
        "12. Projects, documents, UID and warranty traceability",
        "13. Reporting, analytics and auditability",
        "14. Security, administration and platform capabilities",
        "15. End-to-end workflow catalogue",
        "16. ERP control framework and implementation assurance",
        "Appendix A. Document and numbering catalogue",
        "Appendix B. Feature availability and terminology",
    ]
    for item in contents:
        add_bullet(doc, item)
    add_note(doc, "Scope note", "This catalogue describes implemented capabilities visible in the current SAK ERP codebase and deployment. Individual buttons, records and approvals appear according to user permission, document status and tenant configuration.", PALE_BLUE)

    doc.add_page_break()
    doc.add_heading("1  Executive overview and solution landscape", level=1)
    doc.add_paragraph(
        "SAK ERP brings procurement, inventory, production, subcontracting, quality, sales, service, finance, human resources and traceability into one controlled operating system. It is designed around document-driven workflows: each transaction carries a number, status, responsible party, approval state and linked document trail."
    )
    add_feature_table(doc, [
        ("Integrated master data", "Common item, vendor, customer, employee, warehouse, UOM, tax and organization records reduce duplicate entry across departments."),
        ("End-to-end document flow", "Requisitions, orders, issues, receipts, inspections, invoices, payments and reversals remain linked for operational and audit review."),
        ("Status-driven processing", "Actions are enabled only at the correct lifecycle stage; posted stock or finance documents cannot be casually deleted."),
        ("Role-based operations", "Navigation, create/view/edit/approve/download and administrative actions are controlled by user role and permission."),
        ("India-ready commercial data", "GST, HSN/SAC, discounts, Indian currency formatting, vendor/customer statutory fields and commercial terms are supported."),
        ("Browser and mobile-ready", "Responsive web screens, progressive web application assets, offline fallback and print/PDF outputs support office and field usage."),
        ("Live and test separation", "Production and isolated test deployments support controlled validation before business release."),
    ])
    doc.add_heading("Functional landscape", level=2)
    add_feature_table(doc, [
        ("Source-to-pay", "Vendor master → requisition → RFQ/quotation → purchase order → GRN/QC or service entry → supplier invoice → payable/payment."),
        ("Plan-to-produce", "Item/BOM/route → job order → material issue/return → shop-floor execution → quality → finished stock."),
        ("Subcontract-to-pay", "Route → service order → single RM outward challan → vendor WIP → receipt/backflush → QC → service payable."),
        ("Order-to-cash", "Customer → quotation → sales order → ATP/release → dispatch/PGI → invoice → receipt/allocation → return/credit note."),
        ("Service-to-cash", "Installed asset/contract → ticket → assignment/SLA → parts and work confirmation → service invoice → receipt."),
        ("Hire-to-pay", "Employee master → attendance and movements → work-hour controls → payroll and pay-day records."),
    ])

    add_module(doc, "2", "Executive cockpit, workflow and notifications",
        "A single operational landing area gives management visibility into approvals, procurement, quality, stock and module health.", [
        ("Executive KPIs", "Live counters and summaries for pending approvals, purchase activities, QC reminders, inventory alerts and module status."),
        ("Action Required panel", "Consolidated reminders for PO approval and QC-pending records with direct navigation to the respective transaction."),
        ("Collapse and Later", "Collapse minimizes the reminder panel; Later dismisses it for the current session without completing the underlying business action."),
        ("Re-openable actions", "Pending items remain accessible from their source grid even after a reminder card has been opened and closed."),
        ("Role-aware dashboard", "Users see only the modules and actionable information allowed by their assigned roles."),
        ("Operational search", "Global and grid-level search helps users find document numbers, vendors, customers, items and statuses."),
    ], [
        "Reminder dismissal never changes transaction status or bypasses approval/QC.",
        "Deep links retain the business document context and return users to the correct operational screen.",
    ])

    add_module(doc, "3", "Procurement and procure-to-pay",
        "Procurement covers demand initiation, supplier sourcing, commercial ordering, goods/service receipt, quality and accounts payable.", [
        ("Vendor master", "Supplier identity, contacts, statutory information, addresses, payment context and active status."),
        ("Purchase requisitions", "Department demand capture with project/requirement context, line items, approval state and traceable PR numbering."),
        ("Approval workflow", "Pending approvals appear in operational reminders; approved records progress into sourcing and ordering."),
        ("RFQ and vendor responses", "Multiple suppliers, quoted lines, response status, attachments, delivery/commercial terms and comparison support."),
        ("Purchase order creation", "POs can originate from PR/RFQ or direct procurement, with vendor selection, items, UOM, quantity, rate, discount, HSN/SAC, GST, delivery and payment terms."),
        ("Quotation inheritance", "When a received RFQ response is applied, its quotation context and supporting response are linked to the PO to avoid redundant entry."),
        ("PO view, print and download", "Formatted purchase-order view/PDF is available to authorized view users; download permission can be granted without edit rights."),
        ("Discount and totals", "Line discount, taxable amount, GST and total amount are presented in order grids and printed documents using Indian currency grouping."),
        ("Goods receipt note", "Partial and complete GRNs record received quantities against the PO and create a controlled receipt document."),
        ("Incoming QC", "GRN lines can remain pending QC until accepted/rejected quantities and inspection decisions are posted."),
        ("Service entry sheets", "Service procurement is acknowledged through a service entry process rather than physical stock receipt."),
        ("Import files", "Import documentation, commercial references and customs-related handling are maintained for imported procurement."),
        ("Supplier invoice and AP", "QC/service acceptance feeds supplier invoices, advances, debit notes, deductions, payable status and payment history."),
        ("Debit notes", "Commercial or quality deductions can be documented and linked to the vendor transaction."),
    ], [
        "A posted GRN, QC decision or payment is preserved in the document flow rather than silently overwritten.",
        "PO item totals are visible in detail view and financial totals use consistent rounding and Indian lakh/crore formatting.",
        "Vendor quotation attachment requirements account for source RFQ documentation and purchasing route.",
    ])

    add_module(doc, "4", "Inventory and warehouse management",
        "Inventory maintains quantity ownership, warehouse balances and a chronological movement ledger for raw material, WIP and finished goods.", [
        ("Item master", "Item code/name, descriptions, category, UOM, HSN/SAC, status and related engineering/commercial attributes."),
        ("Category and UOM masters", "Controlled classification and units used consistently by procurement, stock, production, sales and service."),
        ("Warehouse stock master", "Available quantity by item and warehouse, including main stores and subcontracting WIP locations."),
        ("Opening balance and adjustments", "Authorized stock initialization and correction with reference, note and ledger impact."),
        ("SIV / material issue", "Stores issue voucher records controlled outbound material movements to production or other approved consumers."),
        ("SRV / material return", "Stores return voucher records material coming back from production or another internal holder."),
        ("GRN stock receipt", "Accepted purchase receipts increase inventory through a numbered and traceable receipt process."),
        ("Transfer and WIP movement", "Warehouse-to-warehouse and vendor-WIP transfers preserve both source and destination context."),
        ("Stock trail", "Per-item chronological ledger shows reference document, movement type, warehouse, inward, outward and running balance."),
        ("Low-stock planning", "Reorder risk and availability alerts support purchasing and production planning."),
        ("Stock reconciliation", "Ledger and balance checks identify inconsistencies without creating duplicate opening/adjustment entries."),
        ("UID/serial linkage", "Serialized units can be traced through stock, deployment, warranty and service events."),
    ], [
        "Duplicate movement-number and duplicate-post protection prevents the same business event being applied twice.",
        "Negative-inventory guards and exact warehouse selection protect issue transactions.",
        "Posted stock movement documents remain auditable and cannot be removed through ordinary grid delete actions.",
    ])

    add_module(doc, "5", "Production, BOM, routing and shop-floor control",
        "Production converts approved material definitions into controlled job orders, issues, work execution and finished output.", [
        ("Bill of materials", "Product structures define input components, output relationship and engineering quantity basis."),
        ("Process routing", "Operations and production sequence establish how an item moves through internal or external processing."),
        ("Job-order creation", "Production orders identify product, planned quantity, BOM/route, schedule and responsible execution context."),
        ("Smart item planning", "Job-order item planning presents required inputs and supports material preparation."),
        ("Material issue and return", "Production consumes stores stock through numbered issue documents and returns unused material through controlled vouchers."),
        ("Workstations and shop floor", "Workstation records and shop-floor views support operation-level execution visibility."),
        ("WIP status", "Open, in-process, partially completed and completed states communicate production progress."),
        ("Output stock update", "Completed and accepted production output is reflected in finished-goods inventory."),
        ("Production traceability", "Job order, source material, movement, operation and output relationships remain reviewable."),
    ], [
        "Issue quantities are validated against availability and the order requirement.",
        "Lifecycle controls distinguish editable plans from posted stock transactions.",
    ])

    add_module(doc, "6", "Subcontracting and outside processing",
        "The subcontracting module follows a clear SAP-aligned job-work flow while supporting multiple finished products from one issued raw material.", [
        ("Route/BOM template", "The route defines one input raw material and the permitted output products/sizes. Vendor, order quantities and prices are intentionally entered later on the service order."),
        ("Flexible output definition", "A route can contain multiple finished or co-product items; duplicate selections are controlled while valid distinct outputs are supported."),
        ("Subcontract service order", "User selects route and one vendor, then records compulsory raw-material weight and length/UOM plus per-output quantity, UOM, size, price, HSN and discount."),
        ("Rapid line entry", "Column-level fill-down arrows copy a first-row UOM, quantity, size, price, HSN or discount through subsequent output rows."),
        ("Single RM issue", "One order creates one material outward challan for the complete route-level raw-material issue—not one issue per finished product."),
        ("Inventory reduction", "The outward challan reduces the selected source warehouse and moves company-owned material into subcontracting/vendor WIP."),
        ("Vendor WIP", "One WIP row summarizes the service order, vendor, input item, remaining material and available next action."),
        ("Receipt/GRN", "A tabular receipt displays all planned outputs, planned/previously received/remaining quantities and the current receipt quantity."),
        ("Partial and excess receipt rules", "Partial receipts leave the order open; controlled over-receipt can be handled when business reality exceeds the original plan."),
        ("Automatic backflush", "Raw-material consumption is calculated from received output quantity and the line-item size/route basis; the user is not required to type RM consumed for every line."),
        ("Mass balance", "Issued RM is reconciled as consumed output basis + scrap + unused return + approved process loss. Open balance remains available for later receipt or settlement."),
        ("Scrap and unused material", "Scrap and unused return are separately recorded; unused balance is derived from issue less calculated consumption and other accounted quantities."),
        ("QC inspection", "Subcontract GRN remains pending QC until the inspection decision is posted. Accepted output then updates finished-goods stock."),
        ("Service invoice", "Line pricing from the service order determines processing value; GST, custom deductions, invoice number/date and vendor invoice attachment determine payable."),
        ("Accounts payable", "QC-approved service liability appears in Finance/AP with processing, tax, deductions, payable and payment status."),
        ("Unified order trail", "One Order Trail displays the service order, MOC, receipts, QC state, stock effects, service invoice and payment progression."),
        ("Professional grids", "Routes, Orders, Vendor WIP and Finance provide search, sorting, filtering, pagination, column controls, status-aware row actions and closeable detail panels."),
    ], [
        "Edit is available before issue; RM Issue is disabled after posting; receipt/QC actions close when the order is fully completed.",
        "Idempotency and document-number controls prevent duplicate challans, receipts and inventory postings.",
        "Partial receipt preserves future receipt capability against the same open outward challan.",
        "Finished stock increases only through the controlled receipt/QC lifecycle; raw stock reduces through the outward challan.",
    ])

    add_module(doc, "7", "Quality management",
        "Quality gates incoming, subcontract and return transactions before inventory or financial completion.", [
        ("Incoming inspection", "Purchase GRNs can be queued for QC with accepted/rejected quantities and inspection status."),
        ("Subcontract QC", "Processed outputs are inspected after GRN; order and WIP statuses transition when all required QC is complete."),
        ("Sales-return inspection", "Returned goods can be inspected and assigned a disposition before stock/credit handling."),
        ("Inspection evidence", "Reference, notes and transaction links retain the context of the quality decision."),
        ("QC reminders", "Pending inspections appear in Action Required and quality views for timely follow-up."),
        ("Status propagation", "QC outcome updates the originating receipt/order and enables the next permitted stock or finance step."),
    ], [
        "QC buttons are disabled after a final decision to avoid repeat posting.",
        "Rejected, accepted and pending quantities remain distinct in operational and document-trail views.",
    ])

    add_module(doc, "8", "Sales and order-to-cash",
        "Sales manages the commercial lifecycle from customer master and quotation through delivery, billing, collection, returns and warranty.", [
        ("Customer master", "Customer code, corporate/government/other type, contact, GST/PAN, billing/shipping addresses, location, credit limit, credit days and active status."),
        ("Data validation", "Contact-title, statutory and date validation is shown inside the active customer/quotation dialog so correction is immediate."),
        ("Quotation", "Customer, quotation/validity dates, payment terms and multiple item lines with quantity, unit price and discount."),
        ("Quotation revision and approval", "Controlled revision/approval status supports commercial negotiation and conversion readiness."),
        ("Quotation view and PDF", "Client-ready quotation view, print/save PDF and Indian lakh/crore number formatting."),
        ("Future validity dates", "Quotation validity correctly supports future expiry dates while preventing invalid date sequences."),
        ("Sales order", "Direct or quotation-based order creation with customer, item lines, expected delivery, payment terms, delivery terms and notes."),
        ("Order item editing", "Authorized pre-fulfilment edit displays and updates the order item lines as well as header terms."),
        ("Pricing and tax", "Quantity, rate, discount, taxable value, GST and order total are visible in transaction detail and output documents."),
        ("Approval, credit and ATP", "Release controls can validate approval state, customer credit and available-to-promise inventory."),
        ("Dispatch / delivery note", "Warehouse and UID/serial selection, dispatch quantity, delivery context and numbered delivery note."),
        ("Post goods issue", "Dispatch posting reduces stock and advances fulfilment while protecting against duplicate or excessive issue."),
        ("Proof of delivery", "Delivery completion context can be recorded against the dispatch."),
        ("Billing", "Sales invoice is generated from fulfilled quantities with GST and linked sales-order/delivery references."),
        ("Receivables", "Open invoice, due date, outstanding amount, collection status and customer exposure."),
        ("Customer receipts", "Receipt entry, allocation to invoices, part payment, unallocated balance and controlled reversal."),
        ("Returns and credit notes", "Return request/receipt, QC disposition, stock decision and customer credit note."),
        ("Warranty", "Sold UID/product warranty status and downstream service linkage."),
        ("Document flow", "Quotation → sales order → dispatch/PGI → invoice → receipt/allocation → return/credit note is available as one traceable chain."),
        ("Customer statements", "Commercial and collection history supports customer-account review."),
    ], [
        "Posted delivery, invoice and receipt records use reversal/correction rather than destructive deletion.",
        "All commercial amounts use consistent rounding and Indian currency grouping in screen and PDF views.",
        "Status-aware actions prevent dispatch, invoice or receipt steps from being repeated incorrectly.",
    ])

    add_module(doc, "9", "Customer service and service-to-cash",
        "Service management connects installed equipment, contractual entitlement, tickets, field work, parts, SLA and billing.", [
        ("Installed base / equipment", "Customer-owned deployed assets, product/UID identity, location, commissioning and service context."),
        ("Service contracts", "AMC, warranty and on-call agreements with customer, covered assets, validity, included visits/labour, value and tax."),
        ("SLA definition", "Response and resolution targets are stored on the contract and frozen onto the ticket at creation for historical integrity."),
        ("Contract terms snapshot", "The ticket preserves the actual entitlement/terms used even if the contract is later changed."),
        ("Service tickets", "Customer, asset, UID/product, location, contact, issue, priority and service type with numbered ticket lifecycle."),
        ("Entitlement", "Ticket classification distinguishes CONTRACT, WARRANTY and CHARGEABLE service."),
        ("Assignment", "Technician/team assignment, ownership and execution status."),
        ("Technician capacity planning", "Daily technician capacity is maintained in hours and compared with assigned workload for dispatch planning."),
        ("Service workflow", "Open → assigned → in progress → resolved/completed with timestamps and controlled transitions."),
        ("Escalation control", "Ticket escalations record severity, owner, due date, reason and resolution status for auditable management follow-up."),
        ("Failure-code master", "Standard failure codes classify symptom, cause and corrective context so recurring equipment problems can be analysed consistently."),
        ("Parts issue and return", "Spare parts are issued/returned through inventory-linked transactions with stock movement history."),
        ("Work confirmation", "Technician activity, diagnosis, failure code, resolution, labour/visit context and customer-facing job sheet."),
        ("Field evidence and site contacts", "Technician, client site contact, visit timestamps, GPS coordinates, notes, photos and videos are retained against the service visit."),
        ("Contract consumption", "Completed work automatically records visits, labour and chargeable usage against the applicable service-contract entitlement."),
        ("Entitlement balance", "Contract views show included, consumed and remaining visits/labour for operational and commercial control."),
        ("RMA lifecycle", "Return-material authorization records customer, item/asset, quantity, reason, received date, disposition and vendor/customer return status."),
        ("Service billing", "Chargeable parts/labour/visit value can generate a service invoice linked to the ticket."),
        ("Service receipt", "Customer payment entry, allocation and reversal support service-to-cash completion."),
        ("Warranty validation", "Asset/UID warranty eligibility is checked before commercial treatment."),
        ("SLA monitoring", "Overdue response/resolution reporting, open ticket workload and contract-expiry visibility."),
        ("Reliability and cost insight", "Warranty cost and product/asset service history support reliability analysis."),
        ("Service document trail", "Installed asset/contract → ticket → assignment → parts/work confirmation → invoice → receipt."),
    ], [
        "Ticket entitlement and SLA snapshots preserve the basis used when the ticket was opened.",
        "RMA dates, dispositions and quantities use controlled validation; future or chronologically invalid movements are rejected.",
        "Parts stock movements and financial postings are linked to the service document rather than entered as disconnected adjustments.",
    ])

    add_module(doc, "10", "Finance, payables and receivables",
        "Finance consolidates liabilities and collections generated by procurement, subcontracting, sales and service.", [
        ("Supplier invoices", "Purchase-linked invoice detail, tax, payable, outstanding and payment status."),
        ("Service invoices", "Service-entry and subcontract processing liabilities with invoice number, attachment and source document."),
        ("Accounts payable", "Open payable register, due context, deductions, payment tracking and paid history."),
        ("Advances", "PO and blanket advances with payment detail/history and later adjustment visibility."),
        ("Deductions and debit notes", "Commercial, quality or custom deductions reduce the vendor payable while retaining the reason and source."),
        ("Mark paid / payment status", "Authorized payment completion and status progression from pending to paid."),
        ("Sales receivables", "Customer invoice, due amount, outstanding and collection status."),
        ("Customer receipts", "Receipt entry, allocation across invoices, partial settlement and unallocated amount."),
        ("Ageing and dunning", "Receivable ageing and collection follow-up views support working-capital control."),
        ("Reversals", "Payment and receipt corrections are recorded as controlled reversals rather than record deletion."),
        ("Tax visibility", "Processing/base value, GST/tax, deductions and final payable/receivable are separately displayed."),
        ("Document linkage", "Every financial row links back to its PO, GRN, service order, subcontract receipt, sales invoice or service ticket."),
    ], [
        "Payables arise only after the applicable receipt/service acceptance and QC gate.",
        "Source-order pricing drives subcontract provisional payable; receipt quantity and approved deductions determine the final liability.",
    ])

    add_module(doc, "11", "Human resources, attendance and payroll",
        "HR combines employee administration, personal attendance visibility, movement security and payroll-related controls.", [
        ("Employee master", "Employee code, identity, department/role, contact and employment context."),
        ("Employee self-service", "Personal attendance and movement history is visible to the individual employee."),
        ("HR management view", "Authorized HR/supervisors can review attendance history, evidence, locations and movement timelines."),
        ("Check In", "Starts the workday with time, GPS/location, notes and supporting evidence where configured."),
        ("Go Out", "Records an intermediate exit and pauses in-office work time."),
        ("Return to Office", "Records return and resumes in-office work time; multiple exit/return cycles are supported."),
        ("End Day", "Performs final checkout and closes the attendance day."),
        ("Outing reasons", "Configured reasons such as Trials and Lunch; Super Admin can maintain additional dropdown choices."),
        ("Lunch rule", "Standard one-hour lunch allowance with excess time deducted from the day’s calculated working hours."),
        ("Movement timeline", "Gone-out and returned-at times, reason, location and notes appear in HR and employee views."),
        ("Office geofencing", "Office coordinates are recognized and displayed with concise location labels: Saif Seas - APIS (Vizag) and EAC (Kolkata)."),
        ("Outside-premises selfie", "Checkout outside configured office coordinates requires selfie evidence for security and accountability."),
        ("Work-hour calculation", "Only completed in-office intervals are totaled, with lunch excess and movement pauses correctly handled."),
        ("Payroll support", "Attendance hours, pay days, travel/per-diem context and payroll records support salary processing."),
    ], [
        "Every movement records time and evidence context; intermediate movements are not lost between check-in and final checkout.",
        "Location labels reduce clutter while the underlying coordinates remain available for evidence and audit.",
    ])

    add_module(doc, "12", "Projects, documents, UID and warranty traceability",
        "Supporting modules connect project context, controlled files and serialized product history to core ERP transactions.", [
        ("Project master", "Project identity and context can be referenced by procurement and operational records."),
        ("Document repository", "Upload, categorize, view and maintain business documents and supporting attachments."),
        ("Transaction attachments", "RFQ responses, vendor invoices, delivery/quality evidence and other files remain associated with their source records."),
        ("Organization documents", "Company header, branding and document formatting support consistent client/vendor outputs."),
        ("UID generation and management", "Serialized identity records support item-level traceability."),
        ("UID deployment", "Deployment/installation records connect a serialized unit to customer, location and date."),
        ("Trace UID", "A single trace view follows UID history across inventory, delivery, warranty and service."),
        ("Public warranty lookup", "Authorized/public-facing warranty validation can confirm product coverage."),
        ("Audit trail", "User and transaction events support accountability across the system."),
    ], [
        "Attachment availability follows document access permission.",
        "UID history is additive and traceable; identity is not silently reassigned after business movement.",
    ])

    add_module(doc, "13", "Reporting, analytics and auditability",
        "Operational reports convert transaction history into actionable management information while retaining drill-back to source documents.", [
        ("Executive summaries", "Open orders, active routes/contracts, WIP, payables, receivables and operational backlog."),
        ("Procurement reporting", "PR/PO/GRN progress, pending approval/QC, vendor activity and procurement document status."),
        ("Inventory reporting", "Stock by warehouse/category, low stock, movement trail, negative-risk and reconciliation indicators."),
        ("Production reporting", "Job-order status, material issue/return, WIP and output progress."),
        ("Subcontract reporting", "Vendor WIP, open RM balance, output receipt, QC state, scrap/return/loss, service payable and document trail."),
        ("Sales reporting", "Quotation/order conversion, dispatch, invoice, receivable, ageing, returns and warranty."),
        ("Service reporting", "Ticket workload, SLA overdue, installed base, active/expiring contracts, entitlement and service costs."),
        ("HR reporting", "Attendance register, work hours, movements, evidence and pay-day information."),
        ("Grid standards", "Search, column sorting, status filtering, configurable columns, resizing, pagination and row-level actions."),
        ("Export/print", "Operational documents and client/vendor outputs support view, print and PDF/download where authorized."),
        ("Audit drill-down", "Document numbers in summaries lead to the originating order, movement, receipt, inspection, invoice or payment."),
    ], [
        "Counts and financial totals are derived from transaction state, not manually typed dashboard values.",
        "Indian currency grouping and consistent decimal rounding are used in client-facing financial views.",
    ])

    add_module(doc, "14", "Security, administration and platform capabilities",
        "The platform separates tenants, users and duties while providing configuration and resilient web delivery.", [
        ("Authentication", "Secure login, reset-password flow and authenticated API access."),
        ("Role-based access control", "Module and action permissions for view, create, edit, approve, download and administration."),
        ("Segregation of duties", "Request, approval, receipt, QC and payment actions can be separated across responsible roles."),
        ("Tenant isolation", "Organization/tenant context limits data visibility and configuration."),
        ("User and role administration", "Super Admin maintains users, roles and permission assignments."),
        ("Company settings", "Organization profile, document header/branding and operational configuration."),
        ("Email configuration", "Administrative email setup and notification integration."),
        ("Migration reliability", "Database migrations execute through the direct PostgreSQL connection rather than depending on an optional remote SQL RPC."),
        ("PWA support", "Installable web application assets, service worker and offline fallback page."),
        ("Responsive UI", "Desktop and mobile layouts use consistent navigation, modal close controls and grid patterns."),
        ("API architecture", "Modular controllers/services support procurement, inventory, production, quality, sales, service, HR and traceability."),
        ("Environment separation", "Live ERP and Mizantra test deployments use separate databases and release validation."),
    ], [
        "Production deployment follows test-first validation and controlled release.",
        "Sensitive infrastructure and database credentials are not exposed in user-facing documents or screens.",
    ])

    # Workflows
    doc.add_page_break()
    doc.add_heading("15  End-to-end workflow catalogue", level=1)
    workflows = [
        ("Procure-to-pay", "PR → approval → RFQ/vendor response → PO → GRN → QC → supplier invoice → AP → payment"),
        ("Service procurement", "PR/PO → service entry sheet → acceptance → supplier/service invoice → AP → payment"),
        ("Production", "Item/BOM/route → job order → material issue → operation/shop floor → QC → finished stock"),
        ("Subcontracting", "Route → service order → single RM issue/MOC → vendor WIP → partial/final GRN → backflush/mass balance → QC → service payable → payment"),
        ("Order-to-cash", "Customer → quotation → approval/conversion → sales order → ATP/release → dispatch/PGI → invoice → receipt/allocation"),
        ("Sales return", "Return request/receipt → QC disposition → stock decision → credit note → customer account adjustment"),
        ("Service-to-cash", "Installed asset/contract → ticket → entitlement/SLA → assignment → parts/work confirmation → service invoice → receipt"),
        ("Attendance", "Check In → zero or more Go Out/Return cycles → End Day → hours/pay-day calculation"),
        ("UID lifecycle", "Item/UID creation → stock → dispatch/deployment → warranty → ticket/service history"),
    ]
    add_feature_table(doc, workflows, headers=("Business process", "Controlled document flow"), widths=(2350, 7010))
    doc.add_heading("Lifecycle action rules", level=2)
    add_feature_table(doc, [
        ("Draft / planning", "Create, view and edit are permitted according to role; validation is completed before posting."),
        ("Approved / released", "Operational execution is enabled; edits are limited to fields that do not invalidate approval."),
        ("Posted stock movement", "No destructive delete. Correction occurs through return, reversal or controlled adjustment."),
        ("Pending QC", "Stock/finance completion waits for inspection where the configured process requires it."),
        ("Financially posted", "Payment/receipt is preserved; cancellation or reversal creates a traceable counter-event."),
        ("Completed", "Execution actions are disabled; record, attachments and document trail remain available for review."),
    ])

    doc.add_page_break()
    doc.add_heading("16  ERP control framework and implementation assurance", level=1)
    add_feature_table(doc, [
        ("Document identity", "Business transactions receive readable, module-specific numbers and success messages display the created document number."),
        ("Document flow", "Each core order has a single trail that shows upstream and downstream documents rather than disconnected popups."),
        ("Idempotent posting", "Duplicate clicks/retries are protected so one receipt or issue does not create multiple stock/financial entries."),
        ("Quantity validation", "Planned, issued, received, accepted, rejected, consumed, scrap, returned and outstanding quantities are validated independently."),
        ("Stock accounting", "Raw material decreases on issue, WIP records company material at the vendor, and accepted finished goods increase stock at receipt/QC."),
        ("Financial accounting", "Pricing, discount, tax, deductions and payable/receivable remain separate and reconcile to the commercial document."),
        ("Approval and QC gates", "Next-step actions remain unavailable until the required approval or inspection is complete."),
        ("Permission control", "The same screen presents different allowed actions to viewer, operator, approver, QC, finance and administrator roles."),
        ("Error placement", "Validation messages are displayed in the active transaction screen/modal with a clear correction message."),
        ("Grid consistency", "Registers use a common pattern for search, sort, filter, page size, pagination, columns and row actions."),
        ("Evidence and attachments", "Source files remain linked to the business document and can be viewed/downloaded by authorized users."),
        ("Release assurance", "Changes are built, deployed to the isolated test environment, smoke-tested, then promoted to live."),
    ])
    doc.add_heading("Verification approach", level=2)
    for text in [
        "API and database verification for schema, document creation, status transitions, stock movements and financial calculations.",
        "Browser click-through verification for create, view, edit, post, close, search, filter, pagination and PDF actions.",
        "End-to-end smoke tests across Sales and Service, including document flow and reversal controls.",
        "Subcontracting smoke coverage from route and order through RM issue, receipt, QC, stock update and payable.",
        "Live validation uses non-destructive checks unless an explicit production transaction test is authorized.",
    ]:
        add_bullet(doc, text)

    # Appendix
    doc.add_page_break()
    doc.add_heading("Appendix A  Document and numbering catalogue", level=1)
    add_feature_table(doc, [
        ("PR", "Purchase requisition"),
        ("RFQ", "Request for quotation / supplier response"),
        ("PO", "Purchase order"),
        ("GRN", "Purchase goods receipt note"),
        ("SIV / SRV", "Stores issue voucher / stores return voucher"),
        ("JOB / production order", "Internal production authorization"),
        ("SUB", "Subcontract service/work order"),
        ("MOC", "Material outward challan to subcontractor"),
        ("SCR", "Subcontract goods receipt / GRN"),
        ("QT", "Sales quotation"),
        ("SO", "Sales order"),
        ("DN", "Delivery/dispatch note"),
        ("INV", "Sales or service invoice"),
        ("CR", "Customer receipt / credit-related document according to context"),
        ("Ticket / contract / asset", "Service request, entitlement agreement and installed equipment record"),
    ], headers=("Document", "Business meaning"), widths=(2500, 6860))

    doc.add_page_break()
    doc.add_heading("Appendix B  Feature availability and terminology", level=1)
    add_feature_table(doc, [
        ("Implemented scope", "Capability is present in the current SAK ERP application. Exact data and buttons depend on role, configuration and document status."),
        ("SAP-aligned", "Uses familiar enterprise controls such as document flow, status gates, material issue/receipt, QC, backflush, AP/AR and reversals. It does not mean SAP certification."),
        ("CRUD", "Create, read/view, update/edit and delete. In ERP transactions, delete is intentionally restricted after a stock or financial posting."),
        ("Backflush", "Automatic calculation of raw-material consumption from accepted/received output and the route/order quantity basis."),
        ("Mass balance", "Reconciliation of issued raw material into consumption, scrap, unused return and approved loss."),
        ("ATP", "Available-to-promise stock check before sales-order release or dispatch."),
        ("PGI", "Post goods issue: the dispatch event that reduces inventory."),
        ("SLA", "Committed service response and resolution target."),
        ("UID", "Unique identifier/serial identity used to trace a physical unit."),
    ])
    add_note(doc, "Client note", "This document is intended for solution presentation, scope confirmation and user orientation. Detailed configuration values, role matrices and transaction procedures can be supplied as separate implementation documents.")

    # Core properties
    props = doc.core_properties
    props.title = "SAK ERP — Complete Feature Catalogue"
    props.subject = "Client-facing catalogue of implemented ERP capabilities"
    props.author = "SAK Solution"
    props.keywords = "SAK ERP, feature catalogue, procurement, inventory, production, subcontracting, sales, service, HR, finance"
    props.comments = "Generated from the current implemented SAK ERP scope as of 19 August 2026."

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build())
