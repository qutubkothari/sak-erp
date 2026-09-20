from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("deliverables/GTM_3_Industry_Wise_Pain_Points_Completed_Questionnaire.docx")
NAVY = "0B2545"
BLUE = "2E74B5"
DARK = "1F4D78"
MUTED = "5B6573"
LIGHT = "F4F6F9"
GOLD = "8A641C"


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        n = tc_mar.find(qn(f"w:{side}"))
        if n is None:
            n = OxmlElement(f"w:{side}")
            tc_mar.append(n)
        n.set(qn("w:w"), str(val))
        n.set(qn("w:type"), "dxa")


def table_geometry(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    for grid_col, width in zip(tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell_margins(cell)


def bottom_rule(paragraph, color="D7DBE2"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def setup_section(section):
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    h = header.add_run("SAK ERP | GTM Topic 3")
    set_font(h, 9, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    f = footer.add_run("Industry-wise Pain Points & SAK ERP Solutions")
    set_font(f, 9, MUTED)


def para(doc, text="", size=11, color=NAVY, bold=False, italic=False, before=0, after=6, line=1.10, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if text:
        r = p.add_run(text)
        set_font(r, size, color, bold, italic)
    return p


def question_block(doc, number, question, answer):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{number}. {question}")
    set_font(r, 11, DARK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(answer)
    set_font(r, 10.7, NAVY)


DATA = [
    {
        "name": "Metal Fabrication & Machining",
        "tag": "Primary UAE launch vertical",
        "summary": "The strongest first campaign for SAK ERP because outside processing and material-to-money traceability are visible, high-value problems.",
        "answers": [
            ("Who is the customer?", "UAE metal fabricators, CNC and machining shops, sheet-metal manufacturers, aluminium/steel processors, job-work subcontractors and multi-site fabrication businesses supplying industrial, construction, marine or energy customers."),
            ("What is happening in their business?", "Material moves repeatedly between stores, machines and outside processors. Production is often driven by urgent job orders, manual spreadsheets and phone calls, while stock is tracked in mixed units such as KG, MTR, MM, PCS and NOS."),
            ("What frustrates management?", "Management cannot quickly prove where a specific lot of raw material went, how much is still with a subcontractor, whether a work order is profitable, or why material consumption and finished output do not reconcile."),
            ("What is the operational obstacle?", "There is no single controlled flow linking BOM/route, raw-material issue, material outward challan, vendor WIP, finished-goods receipt, QC, scrap/return settlement and vendor payable."),
            ("What is the financial consequence?", "Leakage and unrecorded scrap inflate material cost; vendor-held WIP is understated; delays in receipt/QC delay invoicing and payment; and margin is estimated rather than measured per order or product."),
            ("What does SAK ERP do differently?", "SAK ERP connects the full material-to-money chain: multi-UOM inventory, route/BOM driven subcontracting, one controlled material issue per service order, vendor WIP, GRN/QC, document trail, service pricing and payable controls."),
            ("What should we demonstrate?", "RFQ/PO to raw-material GRN; stock by UOM; route creation; subcontracting order; material outward challan; vendor WIP; partial/full receipt; QC inspection; scrap/unused material reconciliation; service invoice and payable traceability."),
            ("What measurable outcome should we promise?", "A baseline-led target: reduce manual follow-up and reconciliation time, improve visibility of vendor-held material, account for 100% of issued material, and measure order-level material and service cost. Promise exact percentages only after a discovery baseline."),
            ("What should the salesperson say?", "SAK ERP shows every kilogram, metre or piece from receipt to production, subcontractor, QC and finance. You can see what was issued, what is still with the vendor, what came back and what the job really cost."),
            ("What should marketing communicate?", "Stop losing control of material outside your factory. SAK ERP gives fabricators one auditable flow from raw material to finished goods, vendor WIP, QC and payable."),
        ],
    },
    {
        "name": "Electrical & Electromechanical Manufacturing",
        "tag": "Complex BOM and traceability vertical",
        "summary": "A fit for businesses that assemble panels, electrical products, controls, harnesses, motors or electromechanical equipment with many purchased components.",
        "answers": [
            ("Who is the customer?", "Panel builders, switchgear manufacturers, electrical equipment assemblers, wiring-harness makers, automation companies, OEM suppliers and electromechanical product manufacturers in the UAE."),
            ("What is happening in their business?", "Teams manage long component lists, alternates, shortages, drawings and frequent engineering changes. Production and procurement need to react to customer-specific specifications and delivery commitments."),
            ("What frustrates management?", "They cannot confidently answer whether all required components are available, whether the correct revision was built, where a serialised unit is, or whether a delayed purchase line will affect promised dispatch."),
            ("What is the operational obstacle?", "BOM, purchase, stock, production, QC and serial/UID records are disconnected. Users re-enter component demand and rely on manual follow-ups to find shortages and trace finished units."),
            ("What is the financial consequence?", "Expediting, excess components, rework, missed delivery dates and warranty exposure reduce margin. Incorrect or incomplete traceability creates costly support and compliance risk."),
            ("What does SAK ERP do differently?", "SAK ERP connects BOM and production demand to inventory/procurement visibility, controls issue and receipt movements, supports UID/serial traceability, quality checkpoints, document trails and service/warranty follow-up."),
            ("What should we demonstrate?", "A customer order converted into a configured BOM demand; material availability and shortage view; controlled issue to production; serial/UID creation; QC result; dispatch; warranty registration and a linked future service call."),
            ("What measurable outcome should we promise?", "A baseline-led target: reduce manual component chasing, increase build-to-correct-revision discipline, improve serial-level traceability and improve on-time delivery predictability."),
            ("What should the salesperson say?", "SAK ERP gives your team one connected view from component availability to the exact serialised unit delivered and serviced. It replaces spreadsheet coordination with controlled production and traceability."),
            ("What should marketing communicate?", "Build the right electrical product, with the right components, revision and traceability - then support it after delivery from the same system."),
        ],
    },
    {
        "name": "Industrial Equipment & Machinery",
        "tag": "Make-to-order lifecycle vertical",
        "summary": "A strong fit where every order has engineering, procurement, manufacturing and after-sales dependencies.",
        "answers": [
            ("Who is the customer?", "Manufacturers and assemblers of industrial machinery, pumps, process equipment, automation systems, custom skids, material-handling equipment and capital goods."),
            ("What is happening in their business?", "Sales commitments initiate customised engineering, procurement and build activities. Long-lead bought-out items, revisions, drawing approvals and site commissioning make every customer order a project-like execution flow."),
            ("What frustrates management?", "Management lacks a single view of order progress, bought-out risk, production status, dispatch readiness, warranty exposure and service history. Customers repeatedly ask for updates that require manual coordination."),
            ("What is the operational obstacle?", "Quotation, sales order, BOM, purchase requirement, job order, inspection, dispatch, installation and service records are not linked as one commercial and operational trail."),
            ("What is the financial consequence?", "Margin erodes through late procurement, rework, unbilled site work, missed spares opportunities and warranty costs that are not tied back to the original equipment or order."),
            ("What does SAK ERP do differently?", "SAK ERP links order-to-cash with procurement, stock, BOM/production, QC, dispatch, serial/UID, warranty and field service. It gives the business a traceable lifecycle rather than disconnected transactions."),
            ("What should we demonstrate?", "A quotation and revision; sales order; demand into BOM/procurement/production; item or equipment serialisation; QC and dispatch; installed-base record; service call; technician work log; parts used and service billing."),
            ("What measurable outcome should we promise?", "A baseline-led target: shorten order-status reporting, improve visibility of procurement and production blockers, improve installed-base traceability and capture more chargeable service/spares work."),
            ("What should the salesperson say?", "SAK ERP lets you manage the machine beyond the sale: from quotation and build through delivery, warranty, service, parts and billing - with one customer and equipment history."),
            ("What should marketing communicate?", "Turn make-to-order machinery delivery into a controlled lifecycle: quote, build, deliver, service and grow recurring revenue from the same ERP."),
        ],
    },
    {
        "name": "Marine & Defence Suppliers",
        "tag": "Controlled traceability and compliance vertical",
        "summary": "Designed for suppliers where evidence, controlled materials, quality records and chain-of-custody matter as much as transaction speed.",
        "answers": [
            ("Who is the customer?", "Marine equipment suppliers, shipyard vendors, defence subcontractors, controlled-component manufacturers and project suppliers serving regulated customers or critical assets."),
            ("What is happening in their business?", "They must manage material certificates, approved vendors, quality evidence, inspections, drawings, serial numbers, controlled movement and contractual documentation across long project cycles."),
            ("What frustrates management?", "When a customer or auditor asks for a trace, the team searches email folders and spreadsheets. It is difficult to prove the supplier, lot, inspection record, document and movement history for a delivered item."),
            ("What is the operational obstacle?", "Quality, procurement, production, subcontracting, documents and dispatch evidence are captured in separate places with incomplete links to the relevant order, material or serialised asset."),
            ("What is the financial consequence?", "Document gaps create delayed acceptance, rejected deliveries, rework, withheld payments and compliance risk. Time spent reconstructing an audit trail is expensive and disrupts project delivery."),
            ("What does SAK ERP do differently?", "SAK ERP provides controlled workflows across procurement, material receipt, UID/serial tracking, quality, production/subcontracting, documents, delivery and service history, all with user and document trails."),
            ("What should we demonstrate?", "An approved procurement record and attachment; incoming GRN/QC; a UID-controlled material; manufacturing/subcontracting movements; final inspection; dispatch document trail; and retrieval of linked evidence from one screen."),
            ("What measurable outcome should we promise?", "A baseline-led target: faster evidence retrieval, improved completion of inspection/document records, lower acceptance delay risk and stronger traceability across high-value or controlled materials."),
            ("What should the salesperson say?", "When your customer asks for proof, SAK ERP gives you the trace - material, supplier, quality, movement, document and delivery - without rebuilding the story from email and spreadsheets."),
            ("What should marketing communicate?", "For marine and defence supply, traceability is not a report. It is the evidence behind every material, inspection and delivery."),
        ],
    },
    {
        "name": "Building Products & Hardware Manufacturing",
        "tag": "High-volume stock and production control vertical",
        "summary": "A fit for manufacturers and distributors managing large SKU counts, multiple units of measure, fast-moving stock and outside processing.",
        "answers": [
            ("Who is the customer?", "Manufacturers and distributors of architectural hardware, building accessories, fittings, fasteners, fabricated building products, aluminium/steel accessories and construction-related product lines."),
            ("What is happening in their business?", "They manage hundreds or thousands of SKUs, fast sales movement, multiple warehouses, conversions between pieces/weight/length, seasonal or project demand and a mix of in-house and subcontracted operations."),
            ("What frustrates management?", "Stock numbers look available but are not reliable by warehouse, UOM, quality status or vendor WIP. Teams struggle to reconcile production consumption, scrap, returns and finished stock."),
            ("What is the operational obstacle?", "Inventory, sales, production, subcontracting and quality are not governed by a shared master-data and movement discipline. Reorder decisions and availability commitments depend on manual checks."),
            ("What is the financial consequence?", "Overstock, stock-outs, incorrect procurement, unplanned substitutions and material leakage tie up cash and hurt customer service. Margin is distorted when production and scrap are not accurately captured."),
            ("What does SAK ERP do differently?", "SAK ERP combines item master controls, multi-UOM handling, stock trails, reorder visibility, sales demand, BOM/production, subcontracting/vendor WIP, QC and finance traceability in one operating system."),
            ("What should we demonstrate?", "Item master with correct UOM; stock by warehouse and status; a sales order; availability/reorder view; production/subcontract route; material issue; finished receipt/QC; and a stock trail showing every movement."),
            ("What measurable outcome should we promise?", "A baseline-led target: improve inventory accuracy and availability confidence, reduce manual stock reconciliation, reduce avoidable stock-outs/over-buying and improve visibility of material consumption and scrap."),
            ("What should the salesperson say?", "SAK ERP gives you a live, traceable answer to what is in stock, what is committed, what is in production or with vendors, and what it cost to make or procure."),
            ("What should marketing communicate?", "Control high-volume stock without losing production traceability. SAK ERP connects warehouse, sales, manufacturing, subcontracting and quality in one real-time view."),
        ],
    },
]


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_section(doc.sections[0])
    add_header_footer(doc.sections[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    p = para(doc, "COMPLETED QUESTIONNAIRE", 10, BLUE, True, after=4)
    title = para(doc, "Industry-wise Pain Points & SAK ERP Solutions", 22, NAVY, True, after=3)
    subtitle = para(doc, "GTM Topic 3 | Vertical-specific sales and marketing responses", 12, MUTED, False, after=12)
    bottom_rule(subtitle)

    para(doc, "Purpose", 13, BLUE, True, before=10, after=4)
    para(doc, "This document answers the ten GTM questions for each priority industry. It is designed to support discovery, UAE vertical messaging, demos, landing pages and sales enablement. Measurable outcomes are intentionally baseline-led so that SAK ERP does not make unsupported customer-specific ROI claims.", 10.7, NAVY, after=10)

    table = doc.add_table(rows=1, cols=2)
    table_geometry(table, [2700, 6660])
    for i, text in enumerate(("Priority", "Recommended approach")):
        c = table.cell(0, i)
        shade(c, LIGHT)
        r = c.paragraphs[0].add_run(text)
        set_font(r, 10.5, DARK, True)
    for priority, approach in [
        ("First UAE campaign", "Metal Fabrication & Machining - lead with outside processing and material-to-money traceability."),
        ("Replicate next", "Electrical/Electromechanical, Industrial Equipment & Machinery, Marine/Defence Suppliers, and Building Products & Hardware."),
        ("Positioning rule", "Do not lead with generic ERP modules. Lead with each vertical's operational control problem and demonstrate the connected workflow."),
    ]:
        row = table.add_row().cells
        for c, text in zip(row, (priority, approach)):
            r = c.paragraphs[0].add_run(text)
            set_font(r, 10.2, NAVY, c is row[0])

    for index, industry in enumerate(DATA):
        doc.add_page_break()
        p = para(doc, f"VERTICAL {index + 1}", 10, GOLD, True, after=3)
        para(doc, industry["name"], 18, NAVY, True, after=3)
        para(doc, industry["tag"], 11, MUTED, False, italic=True, after=8)
        lead = doc.add_table(rows=1, cols=1)
        table_geometry(lead, [9360])
        c = lead.cell(0, 0)
        shade(c, LIGHT)
        r = c.paragraphs[0].add_run(industry["summary"])
        set_font(r, 10.5, NAVY, False)

        for number, (question, answer) in enumerate(industry["answers"], 1):
            question_block(doc, number, question, answer)

    doc.core_properties.title = "Completed Questionnaire - GTM Topic 3"
    doc.core_properties.subject = "Industry-wise Pain Points & SAK ERP Solutions"
    doc.core_properties.author = "SAK ERP"
    doc.save(OUT)


if __name__ == "__main__":
    build()
