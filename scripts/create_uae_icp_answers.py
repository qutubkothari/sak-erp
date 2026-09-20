from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"deliverables/SAK_ERP_UAE_ICP_Qualification_Answers.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def set_table_width(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None: tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(w)); grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'),str(widths[i])); tcW.set(qn('w:type'),'dxa')
            set_cell_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(4); p.add_run(text)

def add_answer(doc, number, question, answer):
    h = doc.add_paragraph(style='Heading 2'); h.add_run(f'{number}. {question}')
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
    r=p.add_run('Answer: '); r.bold=True; r.font.color.rgb=RGBColor(31,58,95)
    p.add_run(answer)

doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(0.75); sec.bottom_margin=Inches(0.75); sec.left_margin=Inches(0.85); sec.right_margin=Inches(0.85)
styles=doc.styles
styles['Normal'].font.name='Calibri'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.1
for s,size,color,before,after in [('Heading 1',16,'2E74B5',16,8),('Heading 2',13,'2E74B5',12,6)]:
    st=styles[s]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
styles['List Bullet'].font.name='Calibri'; styles['List Bullet'].font.size=Pt(10.5)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('SAK ERP'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(122,90,0)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('UAE Ideal Customer Profile (ICP) & Qualification Criteria'); r.bold=True; r.font.size=Pt(22); r.font.color.rgb=RGBColor(11,37,69)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Completed answers for GTM Topic 2 | 22 August 2026'); r.italic=True; r.font.color.rgb=RGBColor(90,90,90)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(8)
r=p.add_run('Executive conclusion. '); r.bold=True; r.font.color.rgb=RGBColor(122,90,0)
p.add_run('SAK ERP should focus first on growing UAE manufacturers and industrial service businesses whose material, subcontracting, service, and finance processes have outgrown spreadsheets or disconnected tools. The strongest beachhead is discrete manufacturing with frequent outside processing, multi-UOM inventory, and a management need for one auditable flow from purchase/issue through receipt, QC, billing, and payment.')

h=doc.add_paragraph(style='Heading 1'); h.add_run('Round 1 - Identify the best customer')
answers=[
(1,'Which manufacturing segments should we attack first?','Priority order: (1) metal fabrication and machining, (2) electrical/electromechanical products, (3) industrial equipment and machinery, (4) marine and defence suppliers, and (5) building products and hardware. These segments have the fastest product-market fit because SAK already addresses raw-material traceability, length/weight/number UOMs, subcontractor WIP, QC, procurement controls, service follow-up, and document trails. Furniture hardware and fasteners are good follow-on segments; plastics and packaging should be pursued after adding more process-specific planning and production controls.'),
(2,'Which outside processes should we target?','Target customers that regularly send company-owned material outside for machining/job work, anodizing or plating, and powder coating/heat treatment. These processes create the clearest pain around outward issue, vendor-held WIP, finished-goods receipt, unused material, scrap, QC, and service-charge reconciliation. The sales message should lead with “know exactly what each subcontractor is holding and what came back.”'),
(3,'What company size should we target?','Choose option B: AED 20-100M turnover and approximately 30-250 ERP users, with a practical sweet spot around AED 20-75M and 30-150 users. These are guide rails, not hard gates: operational complexity, multiple warehouses, subcontracting, and service workload matter more than turnover alone. Companies above AED 100M can be pursued selectively when the buying unit is a focused plant or business division.')
]
for a,b,c in answers: add_answer(doc,a,b,c)

h=doc.add_paragraph(style='Heading 1'); h.add_run('Round 2 - Find the pain')
answers=[
(4,'What events trigger an ERP search?','Top five triggers, in order: (1) inventory problems and unreliable stock/WIP balances, (2) too much subcontracting with weak reconciliation, (3) heavy Excel or WhatsApp dependency, (4) rapid growth, new factory, or expansion to another location, and (5) the current ERP failing to reflect production and service reality. Audit/ISO pressure, customer complaints, and CFO demands for better controls are strong accelerators that increase urgency.'),
(5,'What are the five strongest pain signals?','1) “We send material to several vendors but cannot state the exact balance, scrap, or unused return by vendor.” 2) “Stores, production, purchase, and finance report different stock numbers.” 3) “The ERP exists, but production, subcontracting, or field service still runs on spreadsheets.” 4) “GRN/QC/AP, sales billing, or service invoices are delayed because documents are not linked.” 5) “The owner or CFO needs people to prepare manual reports before knowing profitability, WIP, receivables, or service performance.”')
]
for a,b,c in answers: add_answer(doc,a,b,c)

h=doc.add_paragraph(style='Heading 1'); h.add_run('Round 3 - Existing systems')
answers=[
(6,'Which systems should UAE sales focus on first?','Start with (1) Excel/manual processes, (2) multiple disconnected applications, and (3) Tally, Zoho, or QuickBooks used mainly for accounting while operations remain outside the system. These prospects understand the need for control but have the largest execution gap. Odoo and SAP Business One replacement opportunities should be pursued selectively where production, subcontracting, service, or usability is demonstrably weak; do not lead with a “rip and replace” pitch unless a process proof is clear.'),
(7,'Should we target companies that already have an ERP?','Yes. Make option C - transformation - the primary motion: companies with an ERP but production, subcontracting, inventory, or field service still outside it. Greenfield Excel/manual customers are the secondary motion because the implementation story is simpler. Replacement opportunities are selective and require a quantified control or adoption failure. The overall strategy is D (all three), with C receiving the highest priority.')
]
for a,b,c in answers: add_answer(doc,a,b,c)

h=doc.add_paragraph(style='Heading 1'); h.add_run('Round 4 - Decision maker')
answers=[
(8,'Who feels the pain first and who approves?','Pain is usually discovered by the factory/production manager, stores or inventory manager, COO, or service head. The CFO becomes involved when stock, invoice, margin, or audit exposure is visible. The champion is normally the COO, factory manager, production manager, or finance controller. Final approval is usually the Owner/MD, CFO, or COO; IT validates security, integration, and deployment but is rarely the economic buyer. Sales should multi-thread the opportunity rather than rely on one contact.'),
(9,'What would make the Owner/MD say “I need this”?','Top two messages: (1) “I cannot see what my subcontractors are holding or what my material really became,” and (2) “My ERP does not reflect how my factory and service operation actually work.” A supporting financial message is: “I cannot trust true profitability because stock, service cost, scrap, billing, and collections are disconnected.”')
]
for a,b,c in answers: add_answer(doc,a,b,c)

h=doc.add_paragraph(style='Heading 1'); h.add_run('Round 5 - Qualification')
add_answer(doc,10,'Complete: “The perfect SAK ERP customer is a manufacturer that …”','“The perfect SAK ERP customer is a growing UAE manufacturer or industrial service business that moves material through multiple warehouses, subcontractors, production or field teams, and finance controls; has outgrown spreadsheets or disconnected systems; and has an owner, COO, or CFO willing to standardize the process and measure the result.”')

h=doc.add_paragraph(style='Heading 1'); h.add_run('Formal UAE ICP')
sections=[
('Primary ICP',['UAE discrete manufacturers, metal/electrical/industrial equipment/marine suppliers, AED 20-100M turnover, 30-250 users or equivalent operational complexity.','Frequent outside processing; multiple warehouses; material traceability by weight, length, number, or serial/UID; visible QC and payable controls.','Owner/MD, COO, CFO, factory, production, stores, or service leader has an active control problem and a project sponsor.']),
('Secondary ICP',['Industrial service, maintenance, and after-sales organizations with technicians, serialised assets, contracts/warranties, parts consumption, customer-site evidence, and service billing.','Smaller businesses with unusually high subcontracting or field-service complexity, even if below the turnover band.','ERP users seeking an operational layer or transformation rather than a full finance replacement.']),
('Disqualification / nurture criteria',['No accountable business sponsor or no willingness to provide process owners and master data.','Very simple single-site operations with no meaningful inventory, production, subcontracting, or service complexity.','Requirement is only a low-cost accounting package or a heavily bespoke global rollout with no standardization appetite.','No urgency, no budget path, or expectation that software alone will fix undocumented processes.'])
]
for title, bullets in sections:
    h=doc.add_paragraph(style='Heading 2'); h.add_run(title)
    for b in bullets: add_bullet(doc,b)

h=doc.add_paragraph(style='Heading 1'); h.add_run('Lead qualification scorecard (0-5 each)')
doc.add_paragraph('Use the existing 40-point model. Score evidence, not optimism. A score of 30-40 is Priority A, 20-29 Priority B, and below 20 is Nurture/Reject. A prospect should not be Priority A without a sponsor and a defined pain, even if its company size is attractive.')
rows=[('Manufacturing complexity','0 simple trading; 5 multi-stage production, variants, multiple UOMs, QC, and traceability'),('Outside processing','0 none; 5 frequent vendor processing with material balances, scrap, returns, and multiple vendors'),('Inventory complexity','0 one location; 5 multi-warehouse, batch/serial/UID, length/weight/number, and stock adjustments'),('Current ERP weakness','0 trusted integrated system; 5 Excel/disconnected tools or ERP not used by operations'),('Growth/expansion','0 stable; 5 rapid growth, new plant, new location, new product, or investment/audit pressure'),('Management visibility problem','0 no issue; 5 daily manual reports and no trusted WIP, margin, receivable, or service view'),('Company size','0 outside target; 5 within AED 20-100M / 30-250 users or equivalent complexity'),('Decision-maker access','0 no access; 5 sponsor and economic buyer engaged with a dated decision process')]
t=doc.add_table(rows=1, cols=2); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.style='Table Grid'; set_table_width(t,[2700,6660])
for i,hdr in enumerate(['Factor','Scoring guidance']):
    c=t.rows[0].cells[i]; c.text=hdr; shade(c,'E8EEF5')
    for run in c.paragraphs[0].runs: run.bold=True
for a,b in rows:
    cells=t.add_row().cells; cells[0].text=a; cells[1].text=b

h=doc.add_paragraph(style='Heading 1'); h.add_run('Sales qualification workflow')
for s in ['Confirm the trigger and quantify the current cost: stock variance, WIP leakage, delayed billing, scrap, rework, or service response time.','Map one real end-to-end process, preferably subcontracting or field service, and identify the people, documents, and systems involved.','Score the eight factors using evidence and record the champion, economic buyer, timeline, budget path, and data readiness.','Run a focused proof of value using the customer’s own route/BOM, material issue, receipt/QC, service call, or billing example.','Advance only when the customer accepts success measures, implementation owners, and a decision date.']: add_bullet(doc,s)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12)
r=p.add_run('Scope note. '); r.bold=True; r.font.color.rgb=RGBColor(155,28,28)
p.add_run('This ICP is intentionally focused on UAE go-to-market fit. India and UAE deployments should remain separated by tenant, database, legal entity, currency, tax/localization configuration, and access rules; do not mix customer data or market assumptions in qualification reporting.')

for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=footer.add_run('SAK ERP | UAE ICP & Qualification Criteria'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(120,120,120)

doc.save(OUT)
print(OUT)

