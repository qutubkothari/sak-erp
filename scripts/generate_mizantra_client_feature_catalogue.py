from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

import generate_sak_erp_feature_catalogue as base


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "documents"
DOCX_PATH = OUT_DIR / "Mizantra_Client_Feature_Catalogue_2026.docx"


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.86)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("29231F")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in (
        ("Title", 30, base.BROWN, 0, 9),
        ("Subtitle", 15, base.GOLD, 0, 8),
        ("Heading 1", 16, base.BLUE, 16, 8),
        ("Heading 2", 13, base.BLUE, 12, 6),
        ("Heading 3", 11.5, base.DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        styles[name].font.name = "Calibri"
        styles[name].font.size = Pt(10)

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.82), WD_TAB_ALIGNMENT.RIGHT)
    icon = ROOT / "apps" / "web" / "public" / "pwa-icon-192.png"
    if icon.exists():
        p.add_run().add_picture(str(icon), width=Inches(0.25))
        p.add_run("  ")
    r = p.add_run("MIZANTRA")
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(base.BROWN)
    r = p.add_run("\tCLIENT CAPABILITY CATALOGUE")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(base.MID_GRAY)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.82), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("Mizantra | Intelligent, governed enterprise operations | 6 September 2026")
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(base.MID_GRAY)
    p.add_run("\tPage ")
    base.add_field(p, "PAGE")


def bullet(doc: Document, text: str, level: int = 0) -> None:
    base.add_bullet(doc, text, level)


def feature_table(doc: Document, rows, headers=("Capability", "Client value"), widths=(3100, 6260)) -> None:
    base.add_feature_table(doc, rows, headers=headers, widths=widths)


def module(doc: Document, number: str, title: str, promise: str, rows, controls=()) -> None:
    base.add_module(doc, number, title, promise, rows, controls)


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    # Editorial customer-pack cover.
    doc.add_paragraph().paragraph_format.space_after = Pt(48)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    icon = ROOT / "apps" / "web" / "public" / "pwa-icon-512.png"
    if icon.exists():
        p.add_run().add_picture(str(icon), width=Inches(1.08))
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("MIZANTRA")
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Intelligent Enterprise Operations Platform")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Client-facing feature catalogue, differentiation and value guide")
    r.font.size = Pt(12.5)
    r.font.color.rgb = RGBColor.from_string(base.MID_GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(28)

    table = doc.add_table(rows=4, cols=2)
    base.set_table_geometry(table, [2450, 4550], indent=1100)
    base.set_borders(table, "D8C8B3")
    rows = [
        ("Purpose", "Client evaluation, solution positioning and scope discovery"),
        ("Platform", "Integrated ERP, CRM, planning, service and management intelligence"),
        ("Delivery model", "Modular, tenant-isolated and permission-controlled"),
        ("Document date", "6 September 2026"),
    ]
    for idx, (label, value) in enumerate(rows):
        base.set_cell_shading(table.cell(idx, 0), base.PALE_GOLD)
        run = table.cell(idx, 0).paragraphs[0].add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(base.BROWN)
        table.cell(idx, 1).paragraphs[0].add_run(value)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    base.add_note(
        doc,
        "Mizantra promise",
        "One connected operating system that turns business data into controlled decisions, then carries approved decisions into traceable execution.",
        base.PALE_BLUE,
    )

    # MOAT and benefits are deliberately first.
    doc.add_page_break()
    doc.add_heading("The Mizantra advantage", level=1)
    p = doc.add_paragraph()
    r = p.add_run("Mizantra is not merely a collection of ERP screens. ")
    r.bold = True
    p.add_run(
        "Its differentiation is the closed loop between operational evidence, intelligent recommendations, human approval, native transaction execution and measurable business value."
    )

    doc.add_heading("Our MOAT: intelligence with enterprise control", level=2)
    feature_table(doc, [
        ("Evidence-led intelligence", "Recommendations originate from tenant-scoped ERP records, exceptions and document history, with source context and confidence rather than unsupported answers."),
        ("Decision-to-execution loop", "Mizantra identifies an exception, routes it to an owner, prepares an action, preserves maker-checker approval and returns the user to the native workflow."),
        ("Deterministic safety core", "Stock, tax, accounting, planning and approval rules remain authoritative. AI improves interpretation and drafting; it does not silently post business transactions."),
        ("One operational graph", "Customers, suppliers, items, orders, receipts, production, quality, service, finance and value evidence are linked instead of being isolated departmental records."),
        ("Manufacturing depth plus CRM", "Advanced MRP, capacity, shop-floor, subcontracting, quality and traceability sit in the same platform as multichannel CRM, sales and service."),
        ("Micro-level tenant control", "Administrators can enable only the modules and screens a client needs, then apply role and action permissions inside that entitled scope."),
        ("Explainable health and priorities", "Management views expose exceptions, impact, urgency, ownership and evidence. Forecast claims are withheld until sufficient client history exists."),
        ("Value realization", "Operational improvements can be tied to baselines, financial verification and benefits evidence, moving the conversation from software usage to verified outcomes."),
    ])

    doc.add_heading("Why clients onboard Mizantra", level=2)
    feature_table(doc, [
        ("Fewer disconnected hand-offs", "Shared master data and end-to-end document flow reduce spreadsheet reconciliation and repeated entry across departments."),
        ("Earlier exception visibility", "Shortages, overdue actions, capacity conflicts, quality loss, service risk and cash exposure become owned work rather than late surprises."),
        ("Faster, safer decisions", "Role-specific worklists and intelligent summaries shorten analysis time while approvals, audit and segregation of duties remain intact."),
        ("Better working capital", "Inventory ageing, slow/non-moving stock, supplier exposure, collections, payment planning and cash forecasting are visible in one system."),
        ("More reliable delivery", "Demand, material, capacity, maintenance, production and dispatch constraints are considered together before commitments are released."),
        ("Stronger customer continuity", "Enquiry, opportunity, quotation, order, installed asset, contract, ticket and payment history remain connected to the same relationship."),
        ("Lower control leakage", "Duplicate prevention, quantity checks, QC gates, approval controls, reversals and immutable audit trails protect operational and financial integrity."),
        ("Adoption by role", "Each stakeholder receives a focused workspace and reports relevant to their responsibilities, including responsive mobile workflows."),
    ])

    doc.add_heading("How Mizantra differs", level=2)
    feature_table(doc, [
        ("Traditional ERP", "Records transactions. Mizantra also prioritizes exceptions, explains impact and prepares governed next actions across modules."),
        ("Standalone BI", "Reports what happened. Mizantra connects insight to the responsible native workflow and captures the decision outcome."),
        ("Generic AI assistant", "Generates answers from prompts. Mizantra limits context by tenant and permission, uses deterministic fallbacks and blocks direct autonomous posting."),
        ("Point CRM or MRP tool", "Optimizes one department. Mizantra links demand, supply, execution, quality, cash, service and customer outcomes."),
        ("One-size-fits-all suite", "Exposes every screen. Mizantra supports tenant-level module/screen entitlements plus role-level action permissions."),
    ])

    doc.add_heading("The Mizantra compounding advantage", level=2)
    for text in [
        "Connect operational evidence across departments instead of copying data into a separate intelligence silo.",
        "Detect and rank the exceptions that threaten delivery, margin, quality, cash or customer commitments.",
        "Route each decision to the right stakeholder with evidence, impact and a governed recommended action.",
        "Execute through the native controlled workflow, preserving approval, segregation of duties and auditability.",
        "Measure the verified outcome and use that history to improve future prioritization, planning and adoption.",
    ]:
        bullet(doc, text)
    base.add_note(
        doc,
        "Why this matters",
        "Every completed transaction and verified outcome strengthens the tenant's operating context. The advantage compounds through connected data, workflow history, controls and client-specific evidence - not through a generic chatbot alone.",
        base.PALE_GOLD,
    )

    doc.add_page_break()
    doc.add_heading("End-to-end operating model", level=1)
    doc.add_paragraph(
        "Mizantra supports complete business value streams. The intelligence layer sits above these controlled document chains and does not replace them."
    )
    feature_table(doc, [
        ("Lead to cash", "Inbound enquiry -> classification -> assignment -> opportunity -> quotation -> approval -> sales order -> fulfilment -> invoice -> collection"),
        ("Source to pay", "Requirement -> PR -> approval -> RFQ/comparison -> PO -> GRN or service entry -> QC -> supplier invoice -> payable -> payment"),
        ("Plan to produce", "Demand -> S&OP/MRP -> material and capacity decision -> job order -> issue -> shop-floor execution -> QC -> finished output"),
        ("Subcontract to pay", "Outside-process requirement -> order -> material outward -> vendor WIP -> receipt -> QC -> reconciliation -> supplier payable"),
        ("Install to serve", "UID dispatch -> deployment/commissioning -> installed asset -> warranty/contract -> ticket -> field service -> billing"),
        ("Hire to pay", "Employee and role -> attendance/movement -> leave -> payroll controls -> payslip/pay-day evidence"),
        ("Insight to value", "Operating evidence -> exception -> priority -> owner -> approved action -> outcome -> financially verified benefit"),
    ], headers=("Value stream", "Connected flow"), widths=(2400, 6960))

    doc.add_heading("Catalogue navigation", level=2)
    for item in [
        "1. Mizantra Intelligence, management control and automation",
        "2. Intelligent CRM and revenue operations",
        "3. Sales and order-to-cash",
        "4. Procurement and source-to-pay",
        "5. Inventory, warehouse and UID traceability",
        "6. Production planning, MRP and capacity",
        "7. Production execution, engineering and maintenance",
        "8. Subcontracting and outside processing",
        "9. Quality, CAPA, EHS and supplier recovery",
        "10. Service, installed base, warranty and contracts",
        "11. Accounts, finance and enterprise control",
        "12. Human resources, attendance and payroll",
        "13. Projects, documents and governed onboarding",
        "14. MIS, stakeholder dashboards and reporting",
        "15. Security, administration, integrations and platform",
    ]:
        bullet(doc, item)

    module(doc, "1", "Mizantra Intelligence, management control and automation",
        "An evidence-led intelligence layer converts cross-functional operating data into role-visible priorities and controlled action.", [
        ("Ask Mizantra / Active Planner", "Natural-language requests across supported ERP workflows; structured interpretation, missing-information prompts, validated draft preview and hand-off to the correct native screen."),
        ("Permission-scoped answers", "Responses use only the data and workflows the signed-in user is entitled to access."),
        ("Deterministic fallback", "Core parsing and supported reports remain available when an external AI provider is unavailable."),
        ("Exception register", "Cross-module operational, quality, finance, maintenance, sales and procurement exceptions with severity, impact, urgency and source evidence."),
        ("Priority and ownership", "Decision queues assign responsibility and make overdue or high-impact work visible."),
        ("Management briefs", "Live daily briefing plus immutable historical snapshots and evidence-only period comparison."),
        ("Factory Health", "Transparent operating-health score with configurable factors, drill-down and history; predictions appear only after an approved evidence threshold."),
        ("Controlled agents", "Operations, collections, supplier and customer agents can analyse and propose within approved policies."),
        ("Governed action tools", "Allowlisted actions create requests or drafts with one-time confirmation, idempotency, approval and audit."),
        ("Knowledge graph", "Tenant-scoped relationships connect business partners, items, documents, events, exceptions and verified value evidence."),
        ("Business Transformation", "Evidence-backed initiatives, baseline/target benefits, financial verification, overlap control and realization tracking."),
        ("Observability and graceful degradation", "Provider metrics, circuit protection and safe fallback prevent an AI outage from becoming an ERP outage."),
    ], [
        "AI does not directly approve, post, pay, release inventory or contact an external party without the configured governed workflow.",
        "Unsupported or weakly evidenced questions return a limitation or request for information instead of invented certainty.",
    ])

    module(doc, "2", "Intelligent CRM and revenue operations",
        "A connected CRM manages the relationship from inbound communication to opportunity, quotation, order, collection and service history.", [
        ("CRM overview and pipeline", "Lead and opportunity stages, aging, ownership, priorities, next actions and conversion visibility."),
        ("Unified inbox", "Inbound WhatsApp and email intake is consolidated for review, classification and lead association."),
        ("Smart message classification", "Messages can be classified as new enquiry, ongoing discussion, order-related, service-related, irrelevant or spam instead of treating every message as a new lead."),
        ("Conversation continuity", "Repeat messages from a known contact can be attached to the same lead timeline rather than creating duplicates."),
        ("Lead scoring", "Priority, completeness, engagement and commercial context support focused follow-up."),
        ("Auto or manual assignment", "Administrators choose manual ownership or round-robin assignment across eligible salespeople; assignment rules remain auditable."),
        ("Lead hygiene", "Search, filters, duplicate detection, merge and authorized single/bulk deletion support a clean active pipeline."),
        ("Accounts and contacts", "Business accounts, multiple contacts, salutations, roles, departments, primary contact and communication consent."),
        ("Opportunity management", "Deal value, currency, probability, expected close, products, competitors, requirement, next step and stage progression."),
        ("Sales targets and scorecards", "Monthly salesperson or territory targets for revenue, collections, wins, leads, calls, visits, quotations and product quantity/value."),
        ("Territory management", "Territory definitions, ownership and territory-level performance."),
        ("Forecasting", "Weighted pipeline, target attainment, stale deals and overdue next actions."),
        ("Follow-up cadences", "Structured call, task, email and WhatsApp follow-up steps with explicit-send controls."),
        ("Campaigns", "Campaign setup, budget/expected revenue and selected lead participation."),
        ("Templates and workflow control", "Reusable communication templates and governed workflow rules."),
        ("Communication ledger", "Traceable inbound/outbound activity and delivery status."),
    ], [
        "Email and WhatsApp capture require client-approved provider credentials, mailbox/number connection, consent and classification settings.",
        "Automated external messages remain limited to approved templates/rules and configured consent controls.",
    ])

    module(doc, "3", "Sales and order-to-cash",
        "Commercial execution connects customer master data, quotations, fulfilment, billing, receipts, returns and warranty.", [
        ("Customer master", "Commercial and statutory identity, addresses, contacts, GST/tax data, credit and relationship context."),
        ("Quotations", "Revision-controlled offers with products, quantity, price, discount, tax, validity and supporting documents."),
        ("Quotation approval and conversion", "Approval status and controlled conversion into sales orders without rekeying commercial lines."),
        ("Sales orders", "Order lifecycle, customer references, delivery and payment terms, commercial totals and traceable source quotation."),
        ("ATP and release", "Availability checks and controlled release before fulfilment."),
        ("Pick, pack and dispatch", "Fulfilment preparation, delivery records, transport context and post-goods-issue stock impact."),
        ("Billing", "Invoice generation from governed fulfilment with taxes, discounts, freight/transport and gross total."),
        ("Collections", "Receipts, allocation, open receivables, aging and follow-up visibility."),
        ("Sales returns", "Return receipt, QC disposition, stock decision and credit adjustment."),
        ("Warranty linkage", "Dispatched/deployed UIDs can carry warranty and service history."),
        ("Logistics control", "Delivery risk, status and transport visibility."),
        ("Document outputs", "Role-controlled view, print, download and branded commercial PDFs."),
    ])

    module(doc, "4", "Procurement and source-to-pay",
        "Demand initiation, sourcing, purchasing, receipt, inspection and supplier settlement remain one controlled document chain.", [
        ("Vendor master", "Supplier identity, contacts, address, statutory fields, category, active status and commercial context."),
        ("Purchase requisitions", "Department/project demand with item or service lines, need dates, justification and approval."),
        ("RFQ and response capture", "Multiple supplier quotations, attachments and commercial comparison."),
        ("Strategic sourcing", "Supplier evaluation and sourcing decision support."),
        ("Purchase orders", "Material/service orders with line UOM, rate, discount, HSN/SAC, GST, delivery, payment, freight and approval."),
        ("PO print/PDF", "Branded output with contact salutation, commercial breakup, terms and governed download permission."),
        ("Goods receipts (GRN)", "Partial/full material receipt against eligible PO quantity with duplicate prevention and document linkage."),
        ("Service entry sheets", "Acceptance of service-only procurement without creating a physical stock receipt."),
        ("Incoming QC hand-off", "Received, accepted and rejected quantities with value impact and inspection controls."),
        ("Supplier invoices", "GRN/SES invoice workspace, payment sanction, gross invoice visibility and supporting invoice evidence."),
        ("Debit notes", "Supplier debit/recovery documentation linked to commercial or quality events."),
        ("Contract control", "Supplier agreement dates, terms and compliance context."),
        ("Spend intelligence", "Supplier/category spend, exposure and actionable procurement insight."),
        ("Import files", "Import and supporting-document management for overseas procurement."),
    ], [
        "A material PO cannot be materially changed after an active downstream receipt; reversal/cancellation rules protect the audit chain.",
        "Empty/failed receipts are blocked from QC, stock and payables and can be system-cancelled without creating material or financial movement.",
    ])

    module(doc, "5", "Inventory, warehouse and UID traceability",
        "Planning-grade inventory visibility combines quantity, location, movement, reservation, value and unique-unit genealogy.", [
        ("Item and stock master", "Item code, description, classification, UOM, HSN/SAC, planning and stock attributes."),
        ("Multi-warehouse visibility", "Balances and movements by warehouse/location with controlled access."),
        ("Stock movement ledger", "Receipt, issue, return, adjustment, production and dispatch movement history."),
        ("Low-stock planning", "Shortage awareness and replenishment context."),
        ("Reservations and availability", "Demand allocation and planning-grade available quantity."),
        ("Warehouse control", "Receipt, put-away, picking, counting and operational status."),
        ("Warehouse optimization", "Slotting/flow insight and warehouse exception management."),
        ("Working capital and SLOB", "Ageing, slow/non-moving stock and capital exposure."),
        ("Stock adjustments", "Permission-controlled correction with reason and audit trail."),
        ("SIV and SRV", "Stores issue and return vouchers linked to responsible demand or job order."),
        ("UID management", "Serial/unique identity creation and lifecycle status."),
        ("UID trace", "Searchable genealogy across receipt, production, dispatch, deployment, warranty and service."),
        ("Deployment", "Customer/site deployment and commissioning evidence linked to the physical UID."),
    ])

    module(doc, "6", "Production planning, MRP and capacity",
        "Mizantra converts demand into explainable material and capacity recommendations before any planner-controlled release.", [
        ("Demand planning and S&OP", "Confirmed orders, forecasts and approved consensus demand with source and target date."),
        ("Planning configuration", "Make/buy, lead time, safety stock, MOQ, lot size, scrap/yield, supplier, calendar and time-fence parameters."),
        ("Multi-level MRP", "Approved BOM explosion and netting against on-hand stock, reservations, scheduled receipts, open POs and work orders."),
        ("Dated recommendations", "Required dates, planned supply, shortages, excess/residual demand and policy quantity application."),
        ("Planner workbench", "Review, change, justify and approve recommendations before preparing draft PRs or production orders."),
        ("Exception queue", "Material shortage, supplier lead-time and promise-date risks with recommended action."),
        ("Smart production planning", "Constraint-aware planning packs and scenario comparison."),
        ("Finite capacity", "Work-centre, machine, labour, shift and maintenance availability."),
        ("What-if scenarios", "Alternative quantities, dates, resources and supply choices without silently changing live transactions."),
        ("Planning control tower", "Delivery risk, shortages, overloads, WIP and release status in one review space."),
        ("Controlled release", "Approved planning decisions prepare native drafts; purchasing, stock and accounting controls remain separate."),
        ("Planning autonomy policy", "Bounded automation policies may propose plans within configured safety thresholds."),
    ], [
        "MRP recommendations are advisory until an authorized planner releases them.",
        "Client BOMs, routings, calendars, stock and policy parameters must be accepted before operational reliance.",
    ])

    module(doc, "7", "Production execution, engineering and maintenance",
        "Approved plans progress into traceable work orders, material movement, operations, quality and actual performance.", [
        ("BOM management", "Versioned product structures, quantities, UOM, effectivity and revision control."),
        ("Routing", "Operations, sequence, work centres, processing standards and outside-processing steps."),
        ("Engineering changes", "Controlled revision/effectivity changes with traceable impact."),
        ("Job order creation", "Item-driven job orders linked to BOM/routing, planned quantity, dates and project/demand context."),
        ("Material issue and return", "SIV/SRV quantities tied to the job order and inventory ledger."),
        ("Shop-floor execution", "Operation start/completion, station progress, output, rejection and actual time/consumption."),
        ("Backflush controls", "Controlled consumption calculation based on accepted output and approved product structure."),
        ("WIP visibility", "Open/in-process work and execution variance."),
        ("OEE and loss control", "Availability, performance, quality and downtime/loss analysis."),
        ("Plant maintenance", "Preventive and breakdown maintenance, asset availability, spares and downtime context."),
        ("Tooling and calibration context", "Tool/die life and calibration restrictions can feed planning and quality blocks."),
        ("Production reports", "Plan versus actual, WIP, output, consumption, loss and delivery risk."),
    ])

    module(doc, "8", "Subcontracting and outside processing",
        "Outside processing follows PO-style approval, receipt, quality and payable discipline while preserving material ownership and reconciliation.", [
        ("Subcontract routes", "External operations, vendors, input/output items and conversion basis."),
        ("Subcontract requisition/order", "Demand, approval and service-order workflow with printable commercial document."),
        ("Material outward", "Issue slip/challan for company material sent to the subcontractor."),
        ("Vendor WIP", "Quantity and value visibility for material held outside the plant."),
        ("Partial and final receipt", "Processed material receipts with outstanding balance control."),
        ("Subcontract GRN", "Receipt number and document flow surfaced inside the subcontracting workspace."),
        ("QC and disposition", "Accepted, rejected, rework and scrap decisions."),
        ("Backflush and mass balance", "Issued material reconciled into consumption, finished receipt, scrap, approved loss and unused return."),
        ("Cost reconciliation", "Processing charge, material implication and quantity variance."),
        ("Supplier payable integration", "Subcontract service invoices appear in normal supplier payables instead of a disconnected finance silo."),
        ("Print and document trail", "Order, outward slip, receipt and downstream payable remain printable and linked."),
    ])

    module(doc, "9", "Quality, CAPA, EHS and supplier recovery",
        "Quality is embedded in receiving, production, returns and supplier performance rather than treated as a separate after-the-fact report.", [
        ("Quality overview", "Open inspections, failures, NCRs, supplier issues and quality KPIs."),
        ("Inspection plans", "Revision-controlled characteristics, parameters, limits and sampling context."),
        ("Incoming QC", "GRN accepted/rejected quantity and value, evidence document and inspection notes."),
        ("In-process and final QC", "Production-stage and finished-output decisions."),
        ("NCR", "Non-conformance classification, source, severity, disposition and ownership."),
        ("CAPA", "Root-cause, corrective/preventive actions, owner, due date and effectiveness follow-up."),
        ("Supplier quality", "Vendor performance, recovery and debit-note linkage."),
        ("Rework and scrap", "Traceable disposition and cost/quantity impact."),
        ("Cost of quality", "Prevention, appraisal, internal failure and external failure visibility."),
        ("EHS and sustainability", "Environmental, health, safety and sustainability control workspace."),
        ("UID genealogy", "Affected units remain traceable through inspection, dispatch, warranty and service."),
    ])

    module(doc, "10", "Service, installed base, warranty and contracts",
        "Post-sale operations retain the connection between the customer, physical asset, coverage, ticket, technician, parts and billing.", [
        ("Service tickets", "Customer, ship/site, location, product/UID, priority, category, description and lifecycle."),
        ("Installed base", "Customer-owned/deployed assets with UID, product, location, commissioning and service history."),
        ("Automatic installed-asset creation", "Confirmed installation/commissioning can create or update the installed asset and preserve the deployment source."),
        ("Warranty lookup", "Coverage validation against UID/product and warranty dates."),
        ("Service contracts and entitlements", "Contract period, coverage, SLA and chargeability linked to customer/assets."),
        ("Dispatch board", "Technician assignment, schedule and operational status."),
        ("Technician workspace", "Assigned work, visit progress, evidence and completion."),
        ("Parts and work confirmation", "Parts usage and service activity linked to the ticket."),
        ("Preventive maintenance", "Scheduled service against installed assets/contracts."),
        ("SLA and escalation", "Response/resolution tracking and escalation visibility."),
        ("Service billing", "Chargeable work can progress into invoice and receipt."),
        ("Service reports", "Ticket aging, technician workload, SLA, repeat failure and contract/warranty insight."),
    ])

    module(doc, "11", "Accounts, finance and enterprise control",
        "Finance connects operational source documents to payables, receivables, cash, compliance, costing and management control.", [
        ("Accounts control centre", "High-level accounting status, exposures and period actions."),
        ("Supplier payables", "Goods, service and subcontract invoices; sanction, deductions, advances, due status and payment."),
        ("Accounts receivable and collections", "Customer open items, aging, receipt allocation and follow-up."),
        ("Payment runs", "Grouped payable selection and controlled payment preparation."),
        ("Cash forecast", "Expected inflows/outflows and liquidity outlook."),
        ("Treasury and FX", "Liquidity, currency exposure and treasury controls."),
        ("Bank reconciliation", "Book-to-bank matching and reconciliation status."),
        ("Cost and margin", "Material, processing, production and commercial margin analysis."),
        ("Margin-to-cash", "Margin quality through billing, collection and cash realization."),
        ("Budgets and FP&A", "Budget control and driver-based scenario analysis."),
        ("Fixed assets", "Asset register and financial control context."),
        ("Cost centres", "Responsibility-based cost classification and reporting."),
        ("Opening balances", "Controlled initialization of financial position."),
        ("Statutory and regional compliance", "Statutory returns plus UAE-focused compliance workspace where configured."),
        ("IFRS controls", "IFRS 16 leases, IFRS 15 revenue, IFRS 9 expected credit loss and IAS 37 provisions."),
        ("FX revaluation", "Period currency remeasurement controls."),
        ("Group consolidation", "Legal-entity and intercompany reconciliation/elimination support."),
        ("Report schedules", "Recurring finance reporting calendar."),
    ], [
        "Client chart of accounts, tax rules, opening balances and statutory outputs require finance-owner validation before production use.",
        "Posted financial events are corrected through controlled reversal/adjustment, not destructive deletion.",
    ])

    module(doc, "12", "Human resources, attendance and payroll",
        "Employee self-service and HR management connect identity, attendance, movement, leave and payroll controls.", [
        ("Employee master", "Personal, organizational, role, employment and access context."),
        ("Mobile attendance", "Check-in, go-out, return-to-office and end-day sequence."),
        ("Location and selfie controls", "Policy-based geolocation and camera evidence for off-site actions."),
        ("Movement reasons and notes", "Lunch, visit or other movement context captured with attendance events."),
        ("Attendance history", "Daily timing, movement and worked-hours records."),
        ("Leave requests", "Future-dated leave application, balance/status and manager action."),
        ("HR management", "Attendance review, exceptions, employee administration and approvals."),
        ("Payroll", "Pay-period preparation, attendance-linked controls and payroll records."),
        ("Skills and capacity risk", "Workforce capabilities, coverage gaps and production-capacity risk."),
        ("Exports", "Attendance and management reports, including date and weekday context."),
    ])

    module(doc, "13", "Projects, documents and governed onboarding",
        "Cross-functional records remain tied to project outcomes and source evidence, while onboarding data is staged safely.", [
        ("Project master", "Project identity, customer, owner, dates, commercial and execution context."),
        ("Project margin and EVM", "Budget, actual cost, earned value, schedule/cost performance and margin control."),
        ("Document repository", "Upload, categorize, search, view, download and link documents to business records."),
        ("Document intelligence", "Classify and extract supported documents into staged data for review."),
        ("Approval evidence", "Supporting quotation, invoice or document can open alongside the transaction decision."),
        ("Governed onboarding", "Mapping inference, validation, duplicate detection and exception-only review for supported datasets."),
        ("Independent approval", "Approved staging is handed to native import/workflow controls; no silent master-data write."),
        ("Branding and templates", "Company header, letterhead and document presentation controls."),
    ])

    module(doc, "14", "MIS, stakeholder dashboards and reporting",
        "Crisp, role-specific dashboards keep decision makers focused while detailed registers remain available through structured submenus.", [
        ("Executive Business Pulse", "Enterprise health, high-priority decisions, delivery, cash, inventory, quality and value indicators."),
        ("Finance Manager dashboard", "Liquidity, receivables/payables, close, exposure, margin and control exceptions."),
        ("Accountant daily control", "Transaction completeness, reconciliation, due actions and operational accounting worklist."),
        ("Sales Manager dashboard", "Pipeline, target attainment, forecast, quotation/order movement, collections and team exceptions."),
        ("Territory Manager dashboard", "Territory pipeline, conversion, performance and regional ownership."),
        ("Sales Executive dashboard", "Personal leads, follow-ups, targets, opportunities, quotations and collections."),
        ("Operations Control dashboard", "Procurement, inventory, production, quality and service exceptions."),
        ("Manager approvals", "Role-visible maker-checker worklist."),
        ("Production reports", "Plan/actual, shortages, capacity, WIP, output, loss and delivery risk."),
        ("Operational registers", "Search, sort, filter, paging, configurable columns, saved views and export where authorized."),
        ("Natural-language reporting", "Supported questions resolve to bounded evidence or the correct report/workflow."),
        ("Scheduled reporting", "Finance and management report calendar where configured."),
        ("Actionable drill-down", "KPIs link to the underlying exception or transaction instead of ending at a static chart."),
    ])

    module(doc, "15", "Security, administration, integrations and platform",
        "Mizantra is designed as a configurable multi-tenant platform with enterprise controls at feature, screen, role and action level.", [
        ("Tenant administration", "Create/manage client tenants and isolate their data and configuration."),
        ("Feature entitlements", "Enable or disable modules and individual screens per client without deleting data."),
        ("Role and action permissions", "View, create, edit, delete, approve and download rights by module/screen."),
        ("Segregation of duties", "Separate request, approval, receiving, QC, posting and payment responsibilities."),
        ("Safe landing and navigation", "Users are routed to an allowed workspace; disabled screens are hidden and server-protected."),
        ("Audit trail", "Actor, tenant, action, record, time and relevant evidence retained for governed operations."),
        ("Continuous controls", "Duplicate, leakage, authorization and process-control monitoring."),
        ("Master-data governance", "Controlled stewardship of critical masters and changes."),
        ("Integration Hub", "Tenant-scoped connector catalogue, protected references, mappings and test-event ledger."),
        ("WhatsApp Business", "Connection status/number, CRM capture, governed messaging, automation rules and message ledger."),
        ("Email configuration", "Mailbox/provider configuration for approved inbound and outbound workflows."),
        ("Device/event gateway", "Authenticated, idempotent telemetry ingestion with mapping approval and review-required routing."),
        ("API architecture", "Modular REST services for enterprise workflows and controlled integration."),
        ("Responsive web/PWA", "Desktop, tablet and mobile browser support with installable application assets."),
        ("Environment and release controls", "Separated test/live targets, additive migrations, pre-release backup and targeted deployment practice."),
    ], [
        "External providers, bank/statutory filing services, physical devices and communication channels require client-specific credentials and activation acceptance.",
        "Feature visibility does not grant transaction authority; server-side permission and workflow rules remain the enforcement layer.",
    ])

    doc.add_page_break()
    doc.add_heading("Recommended stakeholder scope", level=1)
    feature_table(doc, [
        ("Owner / CEO / Director", "Business Pulse, Factory Health, high-priority decisions, value realization, cash, delivery, margin and approvals."),
        ("Finance leader", "Finance Manager MIS, treasury, cash forecast, AP/AR, margin, compliance, close and enterprise finance controls."),
        ("Accountant", "Daily control, supplier invoices, payables, receipts, reconciliation, cost centres and authorized postings."),
        ("Sales leader", "CRM overview, team pipeline, forecast, targets, territories, quotations, orders and collections."),
        ("Salesperson", "Assigned leads/opportunities, follow-ups, personal scorecard, quotations and customer activity."),
        ("Purchase manager", "PR/RFQ/PO approval, supplier comparison, contracts, spend and delivery exposure."),
        ("Stores / warehouse", "GRN, stock, locations, SIV/SRV, counts, adjustments and UID movement."),
        ("Production planner", "Demand, MRP, capacity, exceptions, plan scenarios and controlled release."),
        ("Production supervisor/operator", "Job orders, material, shop-floor operations, output, rejection, downtime and WIP."),
        ("Quality manager/inspector", "Inspection plans, incoming/in-process/final QC, NCR, CAPA and supplier recovery."),
        ("Service manager/technician", "Installed base, contracts, dispatch, tickets, SLA, parts, work confirmation and billing."),
        ("HR / employee", "HR management and payroll controls / personal attendance, movement, leave and assigned work."),
        ("System administrator", "Tenant, feature entitlements, users, roles, integrations, communication, master governance and audit."),
    ], headers=("Stakeholder", "Recommended Mizantra workspace"), widths=(2500, 6860))

    doc.add_heading("Implementation and activation approach", level=2)
    for text in [
        "Discover: confirm business value streams, stakeholders, controls and success measures.",
        "Configure: tenant, organization, masters, feature entitlements, roles, workflows and document templates.",
        "Migrate safely: stage data, validate mappings and duplicates, approve, then use native import controls.",
        "Prove end to end: execute representative lead-to-cash, source-to-pay, plan-to-produce and service scenarios in UAT.",
        "Activate integrations: connect approved email, WhatsApp, device, bank or statutory providers with credentials, consent and field mappings.",
        "Go live by role: train users on their focused workspace and preserve maker-checker responsibilities.",
        "Measure value: agree baselines and track adoption, delivery, working capital, productivity, quality and cash outcomes.",
    ]:
        bullet(doc, text)

    base.add_note(
        doc,
        "Availability note",
        "This catalogue describes the current Mizantra product capability. The exact client scope is determined by subscription, tenant feature entitlements, role permissions, master-data readiness, country requirements and approved external integrations.",
        base.PALE_GOLD,
    )

    doc.add_heading("Suggested client value measures", level=2)
    feature_table(doc, [
        ("Delivery", "On-time delivery, promise-date adherence, shortage lead time and schedule attainment."),
        ("Working capital", "Inventory days, SLOB value, receivable days, payable planning and cash forecast accuracy."),
        ("Productivity", "Approval cycle time, planner touch time, transaction rework and manual reconciliation effort."),
        ("Quality", "First-pass yield, rejection, cost of quality, CAPA closure and supplier recovery."),
        ("Commercial", "Lead response, conversion, pipeline hygiene, target attainment, margin and collection realization."),
        ("Service", "Response/resolution SLA, repeat failure, first-time fix and warranty/contract recovery."),
        ("Governance", "Overdue approvals, unauthorized attempts, duplicate prevention, reversal traceability and audit completeness."),
    ], headers=("Outcome area", "Example measures"), widths=(2200, 7160))

    doc.add_heading("What a successful first 90 days looks like", level=2)
    for text in [
        "The agreed stakeholder groups see only the modules, dashboards and actions relevant to their role.",
        "Core masters and opening controls are approved, with representative end-to-end transactions accepted in UAT.",
        "Management receives a concise exception and decision view instead of manually combining departmental reports.",
        "Priority integrations are connected only after credentials, consent, mappings and owners are approved.",
        "Baseline measures are captured and the first value-realization review is scheduled with accountable owners.",
    ]:
        bullet(doc, text)
    base.add_note(
        doc,
        "Next step",
        "Run a structured discovery workshop to select the initial value streams, stakeholder scope, integrations and measurable outcomes. Mizantra can then be configured as a focused operating system for the client rather than an oversized menu of unused software.",
        base.PALE_BLUE,
    )

    props = doc.core_properties
    props.title = "Mizantra - Client Feature Catalogue 2026"
    props.subject = "Client-facing Mizantra differentiation, business benefits and complete feature catalogue"
    props.author = "SAK Solution"
    props.keywords = "Mizantra, ERP, CRM, MRP, manufacturing, intelligence, feature catalogue, business benefits"
    props.comments = "Prepared from the current application navigation, feature entitlement registry and enterprise capability audit."

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build())
