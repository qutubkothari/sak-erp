from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("output/docs/Mizantra_Stakeholder_Module_Scope_Handbook.docx")

NAVY = "243B53"
TEAL = "167D7F"
GOLD = "B88931"
PALE_TEAL = "E8F4F2"
PALE_GOLD = "FFF4D9"
PALE_GREY = "F3F6F8"
TEXT = "253246"
MUTED = "5F6B7A"


def shade(cell, colour):
    tc_pr = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), colour)
    tc_pr.append(fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def border_table(table, color="D5DEE7", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def text(cell, value, size=8.5, bold=False, color=TEXT):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(value)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell)


def bullet(doc, value, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.22 + (0.18 * level))
    p.paragraph_format.first_line_indent = Inches(-0.12)
    r = p.add_run(value)
    r.font.name = "Aptos"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(TEXT)


def heading(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(8)
        rr = p2.add_run(subtitle)
        rr.font.name = "Aptos"
        rr.font.size = Pt(9.5)
        rr.font.color.rgb = RGBColor.from_string(MUTED)


def add_callout(doc, title, body, colour=PALE_TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    border_table(table, "C8E4E1" if colour == PALE_TEAL else "ECDCA7")
    cell = table.cell(0, 0)
    shade(cell, colour)
    set_cell_margins(cell, 120, 180, 120, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(TEAL if colour == PALE_TEAL else "775500")
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.name = "Aptos"
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor.from_string(TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    border_table(table)
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        shade(c, NAVY)
        set_cell_width(c, widths[i])
        text(c, h, 8.5, True, "FFFFFF")
    repeat_header(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_width(cells[i], widths[i])
            if row_index % 2 == 1:
                shade(cells[i], PALE_GREY)
            text(cells[i], value, 8.1)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run("MIZANTRA  |  Stakeholder Scope Handbook  |  Controlled distribution")
    r.font.name = "Aptos"
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p.add_run("  •  Page ")
    p._p.append(fld)


def add_page_break(doc):
    doc.add_page_break()


def add_role_section(doc, title, purpose, rows):
    heading(doc, title, purpose)
    add_matrix(doc,
               ["Stakeholder / role", "Modules in scope", "Key features and responsibilities", "Recommended authority"],
               rows,
               [1.35, 1.55, 3.15, 1.35])


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.58)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.62)
    sec.right_margin = Inches(0.62)
    add_footer(sec)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor.from_string(TEXT)

    # Customer-pack cover
    top = doc.add_paragraph()
    top.paragraph_format.space_after = Pt(3)
    rr = top.add_run("MIZANTRA")
    rr.bold = True
    rr.font.name = "Aptos Display"
    rr.font.size = Pt(18)
    rr.font.color.rgb = RGBColor.from_string(TEAL)
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(12)
    label.paragraph_format.space_after = Pt(4)
    r = label.add_run("ROLE & ACCESS REFERENCE")
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(GOLD)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("Stakeholder, Module\nand Feature Scope Handbook")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(29)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    r = sub.add_run("A practical guide for defining what each Mizantra user should see, do and approve.")
    r.font.name = "Aptos"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    add_callout(doc, "How to use this handbook",
                "Use it during onboarding and access approval. Select only the roles that match a person’s job; module visibility and transaction authority must be granted separately.",
                PALE_GOLD)
    cover_table = doc.add_table(rows=4, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cover_table.autofit = False
    border_table(cover_table, "D5DEE7")
    for i, (a, b) in enumerate([
        ("Document owner", "Mizantra Implementation / Business Owner"),
        ("Audience", "Business owners, department heads, end users and IT administrators"),
        ("Purpose", "Role-scope design, training, onboarding and access sign-off"),
        ("Version", "1.0  |  September 2026"),
    ]):
        shade(cover_table.cell(i, 0), PALE_GREY)
        text(cover_table.cell(i, 0), a, 9, True, NAVY)
        text(cover_table.cell(i, 1), b, 9)
        set_cell_width(cover_table.cell(i, 0), 1.55)
        set_cell_width(cover_table.cell(i, 1), 5.85)

    add_page_break(doc)
    heading(doc, "1. Access model", "A consistent operating model for every team and module.")
    add_callout(doc, "Core rule", "A role gives a person a business scope. It does not automatically give the person approval, posting, payment or configuration authority.")
    add_matrix(doc,
               ["Access level", "Meaning", "Typical examples"],
               [
                   ("View", "Read data, reports and documents only.", "Management view, auditor, customer-support reference."),
                   ("Create / edit", "Prepare a draft or maintain approved master data within the assigned scope.", "Create quotation, PR, job order, service ticket, GRN draft."),
                   ("Submit", "Send a completed draft into the controlled workflow.", "Submit PR, PO, supplier invoice, expense, QC record."),
                   ("Approve / reject", "Make the designated business decision; approvals are separate from preparation.", "PR approver, PO approver, QC decision, payment approver."),
                   ("Post / execute", "Create the controlled operational or financial effect after all prerequisites are met.", "Post accepted GRN, inventory issue, approved payment run."),
                   ("Admin", "Configure users, access, settings and integrations. Use only for named administrators.", "Organisation settings, role setup, audit controls, integrations."),
               ], [1.22, 3.15, 3.85])
    heading(doc, "2. Scope design principles")
    for item in [
        "Least privilege: start with the smallest role that lets the person complete their work.",
        "Segregation of duties: the person preparing a commercial or financial document should not approve their own document unless an exceptional policy is formally approved.",
        "Native controls remain in force: GRN, QC, inventory, invoices and payments stay in their respective workflows.",
        "Use the Mizantra prompt workspace for guided queries and drafts; it does not bypass permissions, approvals or posting controls.",
        "Review access when a user changes department, manager, role or employment status.",
    ]:
        bullet(doc, item)

    add_page_break(doc)
    add_role_section(doc, "3. Leadership, governance and analysis",
                     "Recommended roles for the people responsible for business performance, controls and platform ownership.",
                     [
                         ("Business owner / Super Admin", "Executive dashboard; Active Planner; all module overviews; Settings; Reports", "Owns organisational setup, access design, feature activation, cross-functional review, audit visibility and escalation. Uses intelligence for summaries and governed workflow hand-offs.", "View all; configure access/settings; approve only where formally nominated."),
                         ("Managing Director / Leadership", "Executive dashboard; Reports; planning, sales, finance and production dashboards", "Reviews KPIs, cash and margin exposure, approvals, delivery risks, production plans and operational exceptions. Drills into source screens when required.", "View and designated approval authority; normally no transactional posting."),
                         ("Department Head", "Their department workspace; Approvals; Reports; planning dashboards", "Owns team performance, backlog and exception handling. Reviews and approves documents within delegated value / responsibility limits.", "View, approve / reject within policy; limited create where operationally needed."),
                         ("Internal auditor / Compliance viewer", "Audit trails; Documents; Reports; controlled module views", "Checks evidence, approval history, exception logs and compliance status without altering transactions.", "Read-only. No creation, approval, posting or settings."),
                     ])

    add_page_break(doc)
    add_role_section(doc, "4. Sales, customer and delivery roles",
                     "From lead and quotation through order fulfilment, dispatch, invoice hand-off and customer service.",
                     [
                         ("Sales Executive", "Sales; Customers; Quotations; Sales Orders; Documents", "Maintains customer information, prepares quotations, converts approved quotations to orders, follows order status and stores supporting documents.", "Create/edit/submit sales drafts. No price override or approval unless delegated."),
                         ("Sales Manager", "Sales; Delivery / fulfilment; Collections view; Reports", "Reviews pipeline, quotations, discounts, order fulfilment and customer risk. Supports commercial approvals and delivery prioritisation.", "View plus approve / reject sales documents within policy."),
                         ("Dispatch / Fulfilment Coordinator", "Sales fulfilment; Inventory availability; UID deployment; Documents", "Plans pick/pack, dispatch and delivery evidence; coordinates with stores and service. Confirms only permitted logistics milestones.", "Create/edit fulfilment records; no billing or stock adjustment approval."),
                         ("Customer Service / Account Manager", "Customers; Sales order status; Service; Warranty / installed assets", "Answers customer enquiries, tracks commitments, initiates service requests and maintains customer-facing records.", "Read and create service / customer records; no financial approval."),
                     ])

    add_page_break(doc)
    add_role_section(doc, "5. Procurement and supplier management",
                     "A governed supply-to-pay process: request, approval, purchase order, receipt / service confirmation, invoice and payment.",
                     [
                         ("Procurement Requester", "Purchase Requisitions; catalogue / item reference; Documents", "Raises purchase requisitions with business purpose, quantity, required date, supplier context and supporting evidence.", "Create/edit/submit own PRs. Cannot approve own PRs or POs."),
                         ("Procurement Buyer", "Purchase Overview; Vendors; PRs; Purchase Orders; Spend Intelligence; Contracts", "Sources suppliers, compares offers, prepares POs from approved PRs, maintains vendor / commercial evidence and monitors commitments.", "Create/edit/submit POs. No self-approval; no GRN posting or payment."),
                         ("Procurement Approver", "PRs; POs; Spend Intelligence; Documents", "Approves or rejects requisitions and orders after reviewing necessity, quotation / supplier evidence, budget and terms.", "Approve / reject within delegated limits; view-only on downstream stock / finance unless separately assigned."),
                         ("Supplier / Contract Manager", "Vendors; Contracts; Strategic Sourcing; Spend Intelligence", "Maintains supplier profile, contracts, rate history, performance evidence and sourcing plans.", "Maintain master / contract data in assigned scope; no payment approval."),
                     ])

    add_page_break(doc)
    add_role_section(doc, "6. Stores, inventory and traceability",
                     "Roles that safeguard physical stock, warehouses, receipts, issues, returns and UID traceability.",
                     [
                         ("Storekeeper", "Inventory Overview; Stock Master; GRN; SIV; SRV; Store Vouchers; UID tracking", "Receives material against authorised documents, captures batches / UIDs, executes issues and returns, and records warehouse movements with evidence.", "Create/edit controlled inventory records. Cannot approve own QC, PO or payment."),
                         ("Warehouse Manager", "Warehouse Control; Warehouse Optimisation; Stock Adjustments; Reports", "Controls warehouse operations, reviews stock exceptions, authorises permitted stock adjustments and monitors cycle-count / availability exposure.", "View plus designated adjustment approval / posting authority."),
                         ("Inventory Planner", "Stock Master; Low Stock Planning; MRP; Working Capital & SLOB", "Monitors replenishment, reorder exposure, slow-moving stock and material availability for demand and production.", "Read / plan / recommend. Does not execute stock movement without stores authority."),
                         ("UID / Traceability Coordinator", "UID Overview; Deployment; Traceability; Documents", "Maintains item-level identity, deployment history, warranty evidence and asset traceability.", "Create/edit UID and deployment records in assigned scope; no financial posting."),
                     ])

    add_page_break(doc)
    add_role_section(doc, "7. Production, planning and subcontracting",
                     "Roles that plan, execute and govern the plan-to-produce process, including outside processing.",
                     [
                         ("Production Planner", "Production Overview; MRP; APS / capacity; Demand & S&OP; Job Orders; BOM / Routing", "Turns demand into material and capacity plans, prepares job orders, reviews constraints and releases approved work to operations.", "Create/edit/submit plans and job orders; no self-approval where policy requires separation."),
                         ("Production Supervisor", "Job Orders; Shop Floor; Work Stations; OEE; Engineering Changes", "Runs shop-floor execution, records progress / output / exceptions, assigns work centres and escalates material or quality constraints.", "Execute assigned production transactions; no financial approval."),
                         ("Subcontracting Coordinator", "Subcontracting; Vendors; BOM / Routing; Documents", "Creates a subcontract PR, attaches scope and commercial evidence, submits it for approval, issues material only after approval, records receipt and follows the order trail.", "Create/edit/submit. Cannot approve own subcontract PR / PO."),
                         ("Subcontracting Approver", "Subcontracting; PR / PO evidence; Reports", "Reviews outside-processing need, vendor, rate, quantity, delivery date and material exposure. Approves or rejects the subcontract PR / PO.", "Approve / reject within delegated authority; no material issue unless separately authorised."),
                         ("Subcontract Stores / Receipt User", "Subcontracting; Inventory; GRN / QC status; UID", "Issues approved material outward, records subcontract receipt / return, captures accepted quantity and traceability details.", "Execute inventory steps after the subcontract PO is approved."),
                     ])
    add_callout(doc, "Subcontracting control point", "New subcontract work starts as a draft PR. It is submitted and approved before the PO is released / printed and before material can be issued. Supplier financial exposure remains in normal Accounts Payable, with its subcontract source visible there.", PALE_GOLD)

    add_page_break(doc)
    add_role_section(doc, "8. Quality, engineering and EHS",
                     "Roles responsible for product conformity, non-conformance control, corrective action and compliant operations.",
                     [
                         ("Quality Inspector", "Quality; Inspection Plans; GRN QC; CAPA; Cost of Quality", "Performs inspections, records measurements and evidence, accepts / rejects / holds material or output, and raises NCR / CAPA where needed.", "Create/edit QC records; make designated acceptance / rejection decisions."),
                         ("Quality Manager", "Quality; CAPA & Supplier Recovery; Cost of Quality; Reports", "Owns inspection standards, monitors recurring defects, approves corrective action closure and supplier recovery decisions.", "View plus delegated quality approval / closure."),
                         ("Engineering / Process Owner", "BOM; Routing; Engineering Changes; Work Stations; Projects", "Maintains approved technical definitions, routings and change-control evidence; supports production and subcontract scope definition.", "Maintain assigned engineering masters; formal change approval where designated."),
                         ("EHS / Sustainability Officer", "EHS & Sustainability; Documents; Reports", "Records safety / environmental observations, corrective actions, supporting evidence and compliance performance.", "Create/edit/submit EHS records; view operational data as needed."),
                     ])

    add_page_break(doc)
    add_role_section(doc, "9. Finance, accounts and commercial control",
                     "Roles that manage supplier invoices, advances, receivables, cash, bank controls, statutory activity and accounting evidence.",
                     [
                         ("Accounts Payable Executive", "Supplier Invoices; Accounts Payable; Supplier advances; Subcontract source view; Documents", "Verifies supplier invoices against authorised PO / GRN or service-entry evidence, tracks outstanding and open advances, and prepares payment proposals.", "Create/edit/submit invoices and payment proposals. No self-payment approval."),
                         ("Accounts Receivable / Collections Executive", "Collections; Sales billing view; Margin-to-Cash; Reports", "Tracks customer invoices, receipts, collection commitments and ageing; follows up on overdue receivables.", "Create/edit/submit collections records; no journal / payment approval unless assigned."),
                         ("Finance Manager / Controller", "Accounts Control Centre; AP; Treasury; Cash Forecast; Budgets; Cost & Margin; Statutory Returns", "Reviews payables, overdue exposure, advances, cash position, forecast, budget and control exceptions; approves finance workflows within policy.", "Approve / reject within limits; view cross-functional commercial data."),
                         ("Treasury / Payment Officer", "Payment Runs; Bank Reconciliation; Cash Forecast; FX; Banking", "Prepares approved payments, records bank activity, reconciles bank statements and manages cash execution controls.", "Prepare / execute payments only after independent approval; no self-approval."),
                         ("Financial Accountant", "Accounting; Fixed Assets; Opening Balances; Cost Centres; Compliance / statutory modules", "Maintains accounting entries, ledgers, reconciliations, fixed assets, period controls and statutory evidence.", "Post accounting entries according to close policy; restricted administrative access."),
                     ])

    add_page_break(doc)
    add_role_section(doc, "10. Service, people, documents and platform administration",
                     "Roles supporting installed assets, post-sale service, workforce operations and a secure platform.",
                     [
                         ("Service Coordinator", "Service; Installed Assets; Service Contracts / Entitlements; Warranty; Documents", "Registers and maintains service context, assigns tickets, verifies entitlement / warranty status, schedules service and communicates progress.", "Create/edit/submit service records; no financial posting."),
                         ("Service Engineer", "Service Tickets; Installed Assets; UID / traceability; Documents", "Executes assigned field work, records diagnosis, labour / parts usage, service evidence, installation or commissioning outcome and customer acknowledgement.", "Update assigned tickets and evidence; no contract or invoice approval."),
                         ("HR Executive", "HR employee self-service / attendance; HR Overview; workforce skills", "Maintains employee information, attendance exceptions, leave / people records and workforce information needed by managers.", "Create/edit employee records in policy scope; no payroll approval unless separately authorised."),
                         ("HR / Payroll Manager", "HR management / payroll; workforce planning; Reports", "Reviews attendance controls, payroll inputs, skills / capacity risk and people performance data.", "Approve payroll and HR actions within policy; confidential access only."),
                         ("Document Controller", "Documents; Import Files; Audit evidence", "Controls document indexing, linking, retention, version evidence and source-document completeness.", "Upload / maintain documents; view related records; no transaction approval."),
                         ("IT / System Administrator", "Settings; Organisation; Users / roles; Integrations; Automation; Audit Trails", "Creates users, assigns approved role bundles, configures organisation / integrations and investigates technical issues with audit evidence.", "Admin only. No business transaction approval, posting or payment unless separately granted."),
                     ])

    add_page_break(doc)
    heading(doc, "11. Recommended role bundles", "Use these as starter bundles. Add only the module permissions that the job genuinely needs.")
    add_matrix(doc,
               ["Bundle", "Typical assignee", "Primary scope", "Do not combine with"],
               [
                   ("MIZ-OWNER", "Business owner", "All dashboards, settings, governance, reporting", "Routine AP / payment preparation where possible."),
                   ("MIZ-SALES", "Sales Executive", "Customers, quotations, sales orders, documents", "Final sales approval or customer receipt posting."),
                   ("MIZ-BUYER", "Procurement Buyer", "Vendors, PR / PO preparation, sourcing", "Own PO approval and supplier payment approval."),
                   ("MIZ-STORES", "Storekeeper", "GRN, issue / return, warehouse, UID", "QC acceptance and stock-adjustment approval."),
                   ("MIZ-PLANNER", "Production Planner", "MRP, capacity, job orders, BOM / routing view", "Final production / finance approvals."),
                   ("MIZ-SUB-APPROVER", "Subcontracting Approver", "Subcontract PR / PO review and approval", "Preparation of the same subcontract order."),
                   ("MIZ-QUALITY", "Quality Inspector", "Inspections, GRN QC, NCR / CAPA", "Supplier invoice / payment approval."),
                   ("MIZ-AP", "AP Executive", "Supplier invoices, AP, advances, payment proposal", "Final payment approval and bank release."),
                   ("MIZ-TREASURY", "Payment Officer", "Approved payment runs, bank reconciliation", "Payment proposal preparation for the same payment."),
                   ("MIZ-SERVICE", "Service Engineer / Coordinator", "Tickets, installed assets, warranty, contracts", "Commercial terms and invoice approval."),
                   ("MIZ-HR", "HR Executive", "Attendance, employee and workforce records", "Payroll approval for own input."),
                   ("MIZ-AUDIT-VIEW", "Auditor", "Read-only reports, documents and audit trails", "Any create/edit/approve/post capability."),
               ], [1.35, 1.45, 2.85, 2.55])

    add_page_break(doc)
    heading(doc, "12. Role assignment and sign-off worksheet", "Complete one row for each user or a group of users with the same responsibility.")
    add_callout(doc, "Before enabling a role", "Confirm the department owner, manager, approval limit, period of access and whether the user may create, approve, post or pay. Access should be reviewed at least quarterly and immediately on role change.")
    add_matrix(doc,
               ["User / group", "Department", "Role bundle(s)", "Modules", "Create / submit", "Approve / post / pay", "Owner sign-off", "Review date"],
               [("", "", "", "", "", "", "", "") for _ in range(12)],
               [1.00, 0.90, 1.05, 1.10, 1.05, 1.20, 1.05, 0.85])
    heading(doc, "13. Final go-live checklist")
    for item in [
        "Every named user has exactly the department role bundle(s) needed for their job.",
        "Approver limits and substitutes are documented and tested using a sample transaction.",
        "No person can prepare and finally approve / pay the same sensitive transaction without documented exception approval.",
        "Users have been trained on their module workflow, required evidence and escalation route.",
        "Access has been tested from the user’s own login, including menus, prompt workspace, mobile views and approval notifications.",
        "The business owner has signed off on the final access matrix and review schedule.",
    ]:
        bullet(doc, item)

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
