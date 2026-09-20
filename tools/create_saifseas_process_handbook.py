from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('output/docs/ERP_SaifSeas_Process_Handbook.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = '283B4D'; GOLD = '9A7B45'; LIGHT_GOLD = 'F7F1E5'; BLUE = '2E74B5'
PALE_BLUE = 'E8EEF5'; PALE_GRAY = 'F2F4F7'; INK = '1F2937'; WHITE = 'FFFFFF'

modules = [
    ('Executive Cockpit & Approvals', 'A single management view of priorities, approvals, risk and operational work.',
     ['Operations dashboard and KPI cards', 'Manager approvals and action reminders', 'Reports, audit trail and executive brief'],
     [('Review dashboard', 'Executive / manager', 'See pending approvals, exceptions and operating priorities.'), ('Open work item', 'Assigned owner', 'Navigate to the governing module and supporting document.'), ('Approve or return', 'Authorized approver', 'Decision is recorded with audit evidence.'), ('Monitor outcome', 'Management', 'Refresh dashboard, reports and follow-up actions.')],
     ['Approvals remain within the source workflow.', 'Dashboard metrics are operational indicators, not accounting postings.'],
     'All module owners start here each day; use Manager Approvals before creating follow-up transactions.'),
    ('Procurement', 'Controls supplier sourcing and the procure-to-pay document chain.',
     ['Vendor master and supplier comparison', 'Purchase requisitions and approval workflow', 'Purchase orders, contracts, service entries and debit notes', 'Spend intelligence and strategic sourcing'],
     [('Raise purchase requisition', 'Requestor', 'Need, item/service, quantity and required date are captured.'), ('Review and approve PR', 'Manager / procurement', 'Approved demand becomes eligible for sourcing.'), ('Source and create PO', 'Buyer', 'Supplier, price, tax, delivery and terms are frozen in the PO.'), ('Receive through GRN or service entry', 'Stores / requester', 'Receipt evidence is matched to the order.'), ('Send to quality and AP', 'QC / accounts', 'Accepted receipt supports invoice and payment processing.')],
     ['Only approved PRs should progress to PO.', 'PO rates and taxes are controlled source values; GRN quantity cannot exceed governed tolerance.', 'Service-only orders use Service Entry Sheets, not material GRN.'],
     'Outputs: approved PO, GRN/service entry, supplier invoice readiness, payable exposure.'),
    ('Inventory & Warehouse Control', 'Maintains accurate stock, reservation, warehouse movement and replenishment visibility.',
     ['Item master and warehouse stock master', 'GRN, stock issue voucher (SIV), stock return voucher (SRV)', 'Stock adjustments, alerts, low-stock planning and SLOB analysis', 'Warehouse controls and optimisation'],
     [('Create / maintain item', 'Master-data owner', 'Item, UOM, reorder policy and traceability rules are established.'), ('Receive accepted material', 'Stores / QC', 'GRN acceptance creates governed available stock.'), ('Reserve or issue material', 'Production / stores', 'Reservations and SIV protect and consume stock.'), ('Return or adjust', 'Stores / authorized approver', 'SRV or approved adjustment preserves stock audit trail.'), ('Plan replenishment', 'Planner / buyer', 'Low-stock and working-capital reports trigger demand action.')],
     ['No stock posting from an empty or unaccepted receipt.', 'On-hand, reserved and available quantities are shown separately.', 'Adjustments require reason and authority.'],
     'Outputs: real-time availability, traceability, movement history and replenishment signals.'),
    ('Production, Planning & MRP', 'Converts sales demand and forecasts into controlled production and material plans.',
     ['BOM, routing, job orders and shop-floor execution', 'MRP, APS configuration, demand/S&OP and capacity planning', 'OEE/loss control, engineering changes and maintenance'],
     [('Load demand', 'Planner', 'Sales order, forecast or manual demand defines product, quantity and target date.'), ('Run planning', 'Planner', 'BOM explosion, stock, open supply, lead times and capacity are evaluated.'), ('Review recommendations', 'Planner / manager', 'MRP recommendations are reviewed before release.'), ('Create and release job order', 'Production control', 'Route, work stations and material requirements are governed.'), ('Execute and record completion', 'Shop floor / QC', 'Operations, consumption, output and quality results are captured.')],
     ['MRP recommendations do not directly post stock, purchase or accounting.', 'BOM/routing changes must be controlled before release.', 'Shop-floor completion is traceable to the job order and operation.'],
     'Outputs: feasible job orders, material requirements, capacity exceptions and production performance.'),
    ('Subcontracting', 'Manages outside processing while preserving customer, material, quality and finance traceability.',
     ['Subcontract service orders', 'Material outward challans', 'Subcontract receipt, GRN QC and scrap/rework disposition', 'Subcontract payables and document flow'],
     [('Create subcontract order', 'Production / procurement', 'Operation, vendor, input material and expected output are defined.'), ('Send material outward', 'Stores', 'Challan records material issued to the subcontractor.'), ('Receive finished output', 'Stores', 'Subcontract GRN records received output against the service order.'), ('Inspect every line', 'QC', 'Accepted, rejected, rework or scrap quantities are explicitly recorded.'), ('Post finance outcome', 'Accounts', 'Approved service invoice/payable follows verified receipt and QC.')],
     ['Accepted + rejected must equal the received quantity.', 'Rejected balance requires rework or scrap disposition; scrap item is controlled.', 'Finance is visible only after governed receipt/QC evidence.'],
     'Outputs: subcontract trail from outward material to accepted output, scrap/rework and payable.'),
    ('Sales, Fulfilment & Warranty', 'Controls the order-to-cash flow from customer enquiry to collection and warranty.',
     ['Customer master, quotations and approvals', 'Sales orders, fulfilment, dispatch, billing and collections', 'Returns, logistics control and warranty records'],
     [('Create customer and quotation', 'Sales', 'Commercial scope, product, quantity, price and validity are captured.'), ('Approve and convert to sales order', 'Sales manager', 'Commercial approval creates a controlled order commitment.'), ('Plan fulfilment / ATP', 'Sales + production', 'Availability and production requirement are checked.'), ('Pick, pack and dispatch', 'Warehouse / logistics', 'UID/serial and delivery evidence travel with dispatch.'), ('Invoice and collect', 'Accounts receivable', 'Billing, receipt allocation and ageing are monitored.'), ('Activate warranty', 'Sales / service', 'Warranty record supports later service entitlement.')],
     ['Do not dispatch without governed availability and required UID/serial capture.', 'Billing follows the approved fulfilment/dispatch path.', 'Returns and warranty claims stay linked to the original sale.'],
     'Outputs: customer order visibility, dispatch proof, receivables and warranty traceability.'),
    ('Quality Management', 'Assures incoming, in-process and final output quality with corrective action traceability.',
     ['Inspection plans and GRN inspections', 'NCR, CAPA and supplier quality', 'Quality dashboard, cost of quality and EHS/sustainability'],
     [('Define inspection plan', 'Quality engineer', 'Criteria, sampling and acceptance rules are set.'), ('Perform inspection', 'Inspector', 'Actual accepted/rejected quantities and notes are recorded.'), ('Decide disposition', 'QC owner', 'Accept, rework, return or scrap is controlled.'), ('Raise NCR / CAPA', 'Quality', 'Non-conformance receives owner, root cause and action.'), ('Verify closure', 'Quality manager', 'Effectiveness and closure evidence are retained.')],
     ['QC disposition must reconcile to inspected quantity.', 'Rejected material cannot silently become available stock.', 'CAPA closure requires evidence, not only a status update.'],
     'Outputs: accepted stock/output, non-conformance record, corrective action and supplier feedback.'),
    ('Accounts & Financial Control', 'Provides governed payable, receivable, cash, cost, compliance and financial-control workflows.',
     ['Accounts payable, supplier invoices, payment runs and advances', 'Collections, cash forecast, bank reconciliation and treasury/FX', 'Costing, margin-to-cash, budgets, fixed assets and statutory returns', 'IFRS/UAE compliance, consolidation and financial reporting'],
     [('Receive verified commercial evidence', 'Accounts', 'GRN/QC, service entry, sales invoice or expense evidence is checked.'), ('Create or review accounting document', 'Accounts', 'Invoice, payable, receivable, expense or journal is governed.'), ('Approve payment / collection action', 'Finance approver', 'Authority, due date and available advance are validated.'), ('Post settlement', 'Accounts / treasury', 'Payment, receipt, bank movement and allocation are recorded.'), ('Reconcile and report', 'Finance controller', 'Bank, ageing, margin, tax and forecast reports are reviewed.')],
     ['Payables use verified supplier documents and available advance separately.', 'Overdue is governed from document due date/payment terms.', 'Bank reconciliation and statutory returns are controlled closing activities.'],
     'Outputs: AP/AR ageing, cash position, profitability, compliance and management reporting.'),
    ('Service, Installed Base & Contracts', 'Delivers traceable post-sale support, warranty, maintenance and service billing.',
     ['Service tickets and dispatch board', 'Installed base, service contracts and warranty check', 'Preventive maintenance, technicians, billing and service reports'],
     [('Confirm installation / commissioning', 'Service / commissioning team', 'Installed asset is created from customer UID deployment.'), ('Create warranty entitlement', 'System', 'Active warranty coverage is linked to the installed asset when a valid expiry exists.'), ('Open service ticket', 'Support desk', 'Customer, asset/UID, location and issue are captured.'), ('Validate entitlement and assign', 'System / dispatcher', 'Warranty or contract coverage, SLA and technician are applied.'), ('Complete service and bill if chargeable', 'Technician / accounts', 'Work evidence, parts/labour and customer billing are completed.')],
     ['An installed asset must belong to the selected customer.', 'Warranty dates are never guessed; invalid customer/UID mapping blocks commissioning.', 'Ticket SLA and entitlement are frozen at ticket creation.'],
     'Outputs: installed-base history, service performance, warranty/AMC coverage and service revenue.'),
    ('HR, Attendance & Payroll', 'Manages employee attendance, leave, workforce capacity and payroll controls.',
     ['Employee self-service check-in/check-out and leave', 'HR attendance review and payroll', 'Workforce skills and capacity-risk analysis'],
     [('Employee checks in', 'Employee', 'Attendance date/time and status are recorded.'), ('Manager reviews exceptions', 'Manager / HR', 'Missing, duplicate or exception attendance is corrected through control.'), ('Approve leave and attendance', 'Manager / HR', 'Attendance and leave records are finalized for payroll.'), ('Run payroll', 'HR / finance', 'Approved attendance and payroll inputs form the payroll run.'), ('Review workforce capacity', 'Management', 'Skills and availability feed planning decisions.')],
     ['Attendance uses the company operating date/time zone.', 'Payroll is based on approved attendance/leave data.', 'Managers must resolve exception records before closure.'],
     'Outputs: employee attendance, approved leave, payroll inputs and workforce capacity insight.'),
    ('UID Tracking & Traceability', 'Tracks serialized or UID-controlled products across manufacturing, sales, deployment and service.',
     ['UID management, trace UID and deployment history', 'Customer/location deployment, commissioning and public warranty lookup'],
     [('Generate or register UID', 'Production / stores', 'UID is linked to the governed product or item.'), ('Trace movement', 'Authorized user', 'Production, inventory, dispatch and service links are visible.'), ('Dispatch to customer', 'Logistics', 'UID follows the customer delivery record.'), ('Record deployment', 'Service / commissioning', 'Customer, location, deployment date and contact are captured.'), ('Confirm commissioning', 'Commissioning authority', 'Installed asset and valid warranty entitlement are created.')],
     ['A UID cannot be commissioned to the wrong customer.', 'Asset transfers/decommissioning use controlled service processes.', 'UID traceability ties operational evidence together.'],
     'Outputs: product genealogy, customer installed base and warranty/service linkage.'),
    ('Documents, Compliance & Audit', 'Stores controlled documents and demonstrates who did what, when and under which authority.',
     ['Document repository and attachments', 'Audit trails, continuous controls and approvals evidence', 'Letterheads, certificates and document templates'],
     [('Upload supporting document', 'Business user', 'Source file is attached to its governed transaction.'), ('Review evidence', 'Approver / auditor', 'Commercial, quality or compliance evidence is inspected.'), ('Approve / post controlled document', 'Authorized role', 'Decision and audit trail are retained.'), ('Retrieve document or trail', 'Authorized user', 'Trace source, version and actions for review.'), ('Monitor controls', 'Management / audit', 'Continuous-control exceptions are followed up.')],
     ['Documents support, but do not bypass, approval controls.', 'Audit records are retained against actions and approvals.', 'Use configured templates for external certificates and letters.'],
     'Outputs: searchable evidence, auditability and consistent client-facing documents.'),
    ('Administration, Access & Feature Control', 'Configures the tenant without exposing every feature to every user or customer.',
     ['Users, roles and permissions', 'Company settings, organization, header and notifications', 'Master Feature Access, segregation of duties and integrations', 'WhatsApp and automation controls'],
     [('Configure organization', 'Tenant owner', 'Company identity, document header and operating settings are defined.'), ('Create user and role', 'Administrator', 'Role permissions limit actions to business need.'), ('Enable feature access', 'Tenant owner', 'Only selected modules/features are exposed to the client tenant.'), ('Apply segregation checks', 'Administrator / audit', 'Conflicting duties are identified and resolved.'), ('Review audit and integrations', 'Owner / IT', 'Automation, communications and integration access are governed.')],
     ['Feature access is tenant and role controlled.', 'Do not grant approval, posting or payment rights broadly.', 'Review access after employee, customer or role changes.'],
     'Outputs: controlled user experience, safer feature rollout and audit-ready access governance.'),
    ('Mizantra Intelligence & Business Transformation', 'Provides a safe prompt workspace, analytics and cross-functional decision support.',
     ['Active Planner prompts, conversation history and controlled drafts', 'Executive AI brief, business transformation dashboard and command center', 'Read-only analytics, workflow handoff and approval boundary'],
     [('Ask a business question or request', 'Authorized user', 'Prompt is interpreted in tenant context.'), ('Resolve facts from ERP', 'System', 'Query uses governed inventory, sales, finance or operations data.'), ('Show checked answer or draft', 'System', 'Result identifies module, source workflow and missing information.'), ('Open native workflow', 'User', 'User reviews draft under normal controls.'), ('Approve/post in source module', 'Authorized approver', 'ERP approval and posting rules remain in force.')],
     ['The planner may prepare a draft but cannot approve, pay, post, release or message externally.', 'Use feedback to improve interpretation; it does not auto-post transactions.', 'Sensitive data stays tenant-scoped and permission checked.'],
     'Outputs: faster analysis and guided action without removing ERP governance.'),
]

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tcPr.append(shd)

def set_cell_width(cell, width):
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(width)); tcW.set(qn('w:type'), 'dxa')

def set_cell_text(cell, text, bold=False, color=INK, size=9.3):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color); r.font.name='Calibri'; r._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_table_geometry(table, widths):
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr; tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None: tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    for row in table.rows:
        for cell, width in zip(row.cells, widths): set_cell_width(cell, width)

def set_font(run, size=None, color=None, bold=None, name='Calibri'):
    run.font.name=name; run._element.rPr.rFonts.set(qn('w:ascii'),name); run._element.rPr.rFonts.set(qn('w:hAnsi'),name)
    if size: run.font.size=Pt(size)
    if color: run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold

def add_text(doc, text='', style=None, size=None, color=INK, bold=None, after=6, before=0, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.10
    if align is not None: p.alignment=align
    r=p.add_run(text); set_font(r,size,color,bold)
    return p

def add_bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.10
        r=p.add_run(item); set_font(r,10,INK)

def add_flow(doc, steps):
    for idx, (activity, owner, result) in enumerate(steps, 1):
        table = doc.add_table(rows=1, cols=2); set_table_geometry(table,[1500,7860])
        left,right=table.rows[0].cells; shade(left,GOLD); shade(right, LIGHT_GOLD if idx % 2 else PALE_GRAY)
        set_cell_text(left, f'{idx}', True, WHITE, 14)
        p=right.paragraphs[0]; p.paragraph_format.space_after=Pt(1)
        r=p.add_run(activity); set_font(r,10.2,NAVY,True)
        r=p.add_run(f'  |  Owner: {owner}'); set_font(r,9.2,GOLD,False)
        p2=right.add_paragraph(); p2.paragraph_format.space_after=Pt(2); p2.paragraph_format.space_before=Pt(0)
        r=p2.add_run(result); set_font(r,9.2,INK)
        if idx < len(steps):
            arrow=doc.add_paragraph(); arrow.alignment=WD_ALIGN_PARAGRAPH.CENTER; arrow.paragraph_format.space_before=Pt(0); arrow.paragraph_format.space_after=Pt(0)
            r=arrow.add_run('NEXT STEP'); set_font(r,8,GOLD,True)

def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run('Page '); set_font(run,8,'6B7280')
    fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin')
    instrText=OxmlElement('w:instrText'); instrText.set(qn('xml:space'),'preserve'); instrText.text=' PAGE '
    fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)

doc=Document()
section=doc.sections[0]
section.top_margin=Inches(0.85); section.bottom_margin=Inches(0.75); section.left_margin=Inches(0.8); section.right_margin=Inches(0.8)
section.header_distance=Inches(0.35); section.footer_distance=Inches(0.35)

styles=doc.styles
styles['Normal'].font.name='Calibri'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); styles['Normal'].font.size=Pt(10.3)
for name,size,color,before,after in [('Heading 1',16,NAVY,16,7),('Heading 2',13,BLUE,12,5),('Heading 3',11.5,NAVY,8,4)]:
    st=styles[name]; st.font.name='Calibri'; st._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True; st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)

header=section.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.LEFT; header.paragraph_format.space_after=Pt(0)
r=header.add_run('ERP SAIFSEAS  |  PROCESS HANDBOOK'); set_font(r,8,GOLD,True)
footer=section.footer.paragraphs[0]; add_page_number(footer)

# Cover
for _ in range(6): add_text(doc,'',after=0)
add_text(doc,'SAIF AUTOMATION SERVICES LLP',size=11,color=GOLD,bold=True,after=14,align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(doc,'ERP SaifSeas',size=31,color=NAVY,bold=True,after=6,align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(doc,'Business Process Handbook & Module Flowcharts',size=16,color='4B5563',after=28,align=WD_ALIGN_PARAGRAPH.CENTER)
cover=doc.add_table(rows=4,cols=2); set_table_geometry(cover,[2800,6560])
for (label,value),row in zip([('Prepared for','Client Operations & Implementation Team'),('System','erp.saifseas.com'),('Purpose','Operational training, role clarity and controlled process adoption'),('Version','Client handbook - September 2026')],cover.rows):
    shade(row.cells[0],NAVY); shade(row.cells[1],PALE_BLUE); set_cell_text(row.cells[0],label,True,WHITE,10); set_cell_text(row.cells[1],value,False,INK,10)
add_text(doc,'',after=40)
add_text(doc,'How to use this handbook',style='Heading 2')
add_text(doc,'Each module follows the same reading pattern: what the module controls, key features, a role-based flowchart, control points and the operational outputs that feed the next process.',size=10.5,after=8)
add_text(doc,'Important: configured feature access and role permissions determine which screens an individual user can see or act on. This handbook describes the standard controlled flow; it does not replace the company approval matrix.',size=10.5,color='7A5A00',after=10)
doc.add_page_break()

add_text(doc,'Process Map at a Glance',style='Heading 1')
add_text(doc,'The ERP connects commercial demand, supply, stock, production, quality, finance and after-sales service in one controlled operating model.',size=10.5)
overview=[('Demand to Cash','Customer -> Quotation -> Sales Order -> Fulfilment -> Dispatch -> Invoice -> Collection'),('Supply to Pay','PR -> Approval -> PO -> GRN / Service Entry -> QC -> Supplier Invoice -> Payment'),('Plan to Produce','Demand -> MRP / Capacity -> Job Order -> Material Issue -> Shop Floor -> QC -> Finished Output'),('Subcontract Cycle','Subcontract Order -> Material Outward -> Receipt -> QC -> Rework/Scrap -> Service Payable'),('Install to Serve','UID Dispatch -> Deployment -> Commissioning -> Installed Asset -> Warranty/Contract -> Service Ticket')]
table=doc.add_table(rows=1,cols=2); set_table_geometry(table,[2600,6760]); shade(table.rows[0].cells[0],NAVY); shade(table.rows[0].cells[1],NAVY); set_cell_text(table.rows[0].cells[0],'VALUE STREAM',True,WHITE,9.5); set_cell_text(table.rows[0].cells[1],'CONTROLLED FLOW',True,WHITE,9.5)
for name,flow in overview:
    cells=table.add_row().cells; shade(cells[0],PALE_BLUE); shade(cells[1],PALE_GRAY); set_cell_text(cells[0],name,True,NAVY,10); set_cell_text(cells[1],flow,False,INK,10)
add_text(doc,'Control principles',style='Heading 2')
add_bullets(doc,['Create work in the source module; approve, post and pay only through the prescribed authority.', 'Keep commercial, quality, stock and finance evidence attached to the underlying document.', 'Use the Active Planner for analysis and draft preparation, then complete the controlled action in the native workflow.', 'Escalate a mismatch rather than correcting a quantity, customer, rate, UID or tax value outside its governed source.'])
doc.add_page_break()

add_text(doc,'Module Directory',style='Heading 1')
directory=doc.add_table(rows=1,cols=2); set_table_geometry(directory,[700,8660]); shade(directory.rows[0].cells[0],NAVY); shade(directory.rows[0].cells[1],NAVY); set_cell_text(directory.rows[0].cells[0],'#',True,WHITE,9.5); set_cell_text(directory.rows[0].cells[1],'MODULE',True,WHITE,9.5)
for i,(name,_,_,_,_,_) in enumerate(modules,1):
    cells=directory.add_row().cells; shade(cells[0],LIGHT_GOLD); set_cell_text(cells[0],str(i),True,GOLD,9.5); set_cell_text(cells[1],name,False,INK,10)
add_text(doc,'Training suggestion',style='Heading 2')
add_text(doc,'Train staff by value stream first (Demand to Cash, Supply to Pay, Plan to Produce, Install to Serve), then use the module chapters as role-specific desk references.',size=10.5)

for number,(name,purpose,features,steps,controls,outputs) in enumerate(modules,1):
    doc.add_page_break()
    add_text(doc,f'{number}. {name}',style='Heading 1')
    add_text(doc,purpose,size=11,color='4B5563',after=9)
    add_text(doc,'Key features',style='Heading 2')
    add_bullets(doc,features)
    add_text(doc,'Process flowchart',style='Heading 2')
    add_flow(doc,steps)
    add_text(doc,'Controls that must be followed',style='Heading 2')
    add_bullets(doc,controls)
    add_text(doc,'Operational result',style='Heading 2')
    add_text(doc,outputs,size=10.3,after=6)

doc.add_page_break()
add_text(doc,'Implementation & Training Checklist',style='Heading 1')
checklist=[('1','Confirm tenant feature access','Owner enables only the modules/features contracted for this client.'),('2','Set up master data','Company, users, roles, customers, vendors, items, warehouses, BOMs and approval limits.'),('3','Train by role','Requestor, buyer, storekeeper, QC, planner, shop floor, sales, accounts, service and manager.'),('4','Run controlled test transactions','One complete transaction for each relevant value stream before operating live.'),('5','Review opening reports','Stock, open PR/PO/SO, AP/AR, work orders, service tickets and attendance.'),('6','Start daily control routine','Dashboard review, approvals, exceptions, reconciliation and end-of-day follow-up.')]
tab=doc.add_table(rows=1,cols=3); set_table_geometry(tab,[650,3100,4610])
for cell,text in zip(tab.rows[0].cells,['#','CHECKPOINT','SUCCESS CRITERION']): shade(cell,NAVY); set_cell_text(cell,text,True,WHITE,9)
for n,check,criterion in checklist:
    cells=tab.add_row().cells; shade(cells[0],LIGHT_GOLD); set_cell_text(cells[0],n,True,GOLD,9.5); set_cell_text(cells[1],check,True,NAVY,9.4); set_cell_text(cells[2],criterion,False,INK,9.4)
add_text(doc,'Support & governance note',style='Heading 2')
add_text(doc,'If a process needs a different policy for this client (for example approval value, warranty start rule, scrap disposition, document format or role access), record it as a controlled configuration decision before enabling it for users.',size=10.3)

doc.core_properties.title='ERP SaifSeas - Business Process Handbook & Module Flowcharts'
doc.core_properties.subject='Client training and operational flow handbook'
doc.core_properties.author='SAIF Automation Services LLP'
doc.save(OUT)
print(OUT)
