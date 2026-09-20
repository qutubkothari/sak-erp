from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Users\QK\Downloads\Second Round — Challenging the Positioning.docx")
OUT = Path(r"C:\Users\QK\Documents\GitHub\sak-erp\deliverables\SAK_ERP_Second_Round_Positioning_Response.docx")

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def add_answer(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.6)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, 'F6EFE2'); set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run('SAK RESPONSE  ')
    r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(139,111,71)
    r = p.add_run(text)
    r.font.size = Pt(10); r.font.color.rgb = RGBColor(63,45,32)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)

source = Document(str(SOURCE))
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.72); sec.left_margin = Inches(.82); sec.right_margin = Inches(.82)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5); styles['Normal'].font.color.rgb = RGBColor(63,45,32)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size, color in [('Title', 24, RGBColor(63,45,32)), ('Heading 1', 15, RGBColor(92,71,56)), ('Heading 2', 11.5, RGBColor(139,111,71))]:
    styles[name].font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
    styles[name].font.size = Pt(size); styles[name].font.color.rgb = color
    styles[name].font.bold = True

p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run('SAK ERP — Second-Round Positioning Response')
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(14)
r = p.add_run('Completed response to “Second Round — Challenging the Positioning”'); r.italic = True; r.font.color.rgb = RGBColor(122,103,86)

lead = doc.add_table(rows=1, cols=1); lead.alignment = WD_TABLE_ALIGNMENT.LEFT; lead.autofit = False
cell = lead.cell(0,0); shade(cell,'3F2D20'); set_cell_margins(cell,180,180,180,180)
p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(4)
r=p.add_run('Recommended positioning'); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor(232,220,196)
p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(0)
r=p.add_run('Complete traceability from material to money — built for manufacturers whose real processes do not fit a generic template.'); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=RGBColor(255,255,255)
doc.add_paragraph('The responses below are written for UAE manufacturing prospects while remaining compatible with SAK ERP’s India deployment. Market-specific tax, currency, numbering and date presentation are separated by company profile; operating records are not mixed between markets.')

sections = [
('1. Why SAK, not Odoo?', 'SAK should not compete by claiming that it has more generic features. The strongest answer is operational depth in the manufacturing situations that create the most leakage: subcontracting, vendor-held work-in-progress, material accountability, QC, service execution and the commercial documents that follow. SAK connects those steps in one controlled trail, with the practical screens and workflows already shaped around manufacturing teams. Odoo can be configured for many scenarios; SAK starts closer to the customer’s difficult scenario, reducing the distance between implementation and usable control.', None),
('2. Where is SAK exceptionally strong?', 'The two lead differentiators are outside processing / subcontract manufacturing and material-to-money traceability. They are reinforced by production and inventory integration.', ['One route or service order can define multiple output products and quantities.', 'A single vendor order issues the raw material once, tracks vendor WIP, receives multiple finished products, accounts for unused material, scrap and loss, and carries service pricing into payables.', 'The same document trail links order, issue, receipt, QC, stock movement, invoice and payment status.', 'The system can be adapted to unusual length, weight, count and multi-output manufacturing rules without losing auditability.']),
('3. What does “one accountable operating system” mean?', 'It means management can follow a controlled chain from commercial demand to cash, not just view isolated transactions. SAK records the document number, status, quantity, UOM, user action and next step across the chain: customer order or requirement → procurement → receipt → issue → production or subcontracting → QC → finished stock → dispatch → invoice → payment. For subcontracting, the trail also shows vendor, raw material issued, vendor WIP, received output, backflush consumption, unused return, scrap, deductions, payable value and QC outcome. The depth is transaction-level: the user can open the source document and its linked documents rather than relying on a manually maintained spreadsheet.', None),
('4. How important is outside processing?', 'It should be a primary positioning pillar for target manufacturers that send material or semi-finished goods to outside processors. The exact promise is: “Send part of your manufacturing process outside and retain complete visibility of material, vendor WIP, receipts, QC and cost.” This is stronger and more credible than a generic “subcontracting module” claim because it describes the control problem SAK solves. It is especially relevant to machining, fabrication, anodizing, plating, heat treatment, assembly, job work and other multi-step external operations.', None),
('5. What are the three biggest financial consequences?', 'The three most important consequences to lead with are:', ['Material leakage and uncontrolled subcontracting: raw material issued to vendors is reconciled against finished goods, unused return, scrap and approved loss.', 'Working-capital blockage and stock shortages: procurement, inventory, production WIP and demand are visible together so money is not trapped in excess stock while production waits for a missing item.', 'Margin leakage and delayed cash: actual output, service charges, deductions, tax, invoice, receivable and payment status are connected to the originating order, improving commercial accuracy and follow-up.']),
('6. Should digital transformation be part of the proposition?', 'Yes, but it should support the product promise rather than replace it. SAK is not only a software licence; implementation includes process discovery, document design, role controls, data migration, workflow configuration, user adoption and measurable operating controls. The message should be: “We do not simply install ERP. We help manufacturers turn fragmented factory, commercial and finance processes into one traceable operating system.” This keeps SAK ERP as the product while making the implementation capability a credible differentiator.', None),
('7. What happens after implementation?', 'The proposition should be ERP software + implementation + ongoing business/process improvement. Manufacturing changes after go-live: new products, new vendors, new service models, acquisitions, new tax requirements and revised approval policies. SAK should support controlled changes through configuration, regional profiles, route/BOM updates, workflow adjustments, training, data-quality review and enhancement releases. The promise is continuity and improvement, not a one-time go-live followed by a generic support queue.', None),
('8. Which positioning statement is closest?', 'Option C is closest: “SAK ERP — Complete Traceability from Material to Money.” It is specific, outcome-led and supported by the product’s document trails, stock movements, subcontracting controls, QC, service, invoicing and payment workflows. For a more differentiated market-facing version, use: “SAK ERP — Complete traceability from material to money for manufacturers with complex production and outside processing.”', None),
('9. Is the core problem that complexity grows faster than control?', 'Partly, and this is the right strategic frame. The wording should be sharpened to avoid implying that every customer has the same problem: “As manufacturers grow, production complexity, outside processing and commercial commitments can grow faster than control. SAK ERP reconnects the factory, inventory, service, commercial operations and finance in one traceable system.” This is fundamentally what SAK ERP is designed to do. The claim should be supported in sales conversations with a process map and a live demonstration of one order’s end-to-end trail.', None),
]

for heading, answer, bullets in sections:
    doc.add_heading(heading, level=1)
    add_answer(doc, answer)
    if bullets:
        add_bullets(doc, bullets)

doc.add_heading('Recommended UAE market message', level=1)
doc.add_paragraph('SAK ERP is a manufacturing operating system for UAE companies that need more control than a disconnected accounting, inventory and production setup can provide. It connects material, production, subcontractors, quality, service, commercial documents and finance in one auditable flow, while adapting to the company’s actual processes instead of forcing every operation into a generic template.')
doc.add_heading('Proof points to demonstrate in a sales meeting', level=1)
add_bullets(doc, ['Create a route/BOM with multiple outputs and UOMs.', 'Create one vendor order and issue the total raw-material quantity once.', 'Show vendor WIP and receive output through GRN/QC.', 'Open the single document trail from order through stock, invoice and payable.', 'Show service assignment, technician, customer contact, parts, evidence and warranty treatment.', 'Switch company profile between India and UAE to demonstrate separated tax, currency, date and registration labels.'])
doc.add_heading('Positioning guardrails', level=1)
doc.add_paragraph('Do not position SAK as a universal replacement for every enterprise suite or claim that configuration alone solves every process. Lead with the manufacturing control problems SAK demonstrably addresses, show the evidence in the product, and describe integrations or advanced regulatory capabilities as implementation scope where appropriate.')

doc.core_properties.title = 'SAK ERP — Second-Round Positioning Response'
doc.core_properties.subject = 'Completed positioning questionnaire for UAE manufacturing ERP market'
doc.core_properties.author = 'SAK ERP'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(OUT)
