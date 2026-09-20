from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docx" / "Mizantra_Intelligent_CRM_UAT_Test_Guide.docx"

INK = "2F241B"
BROWN = "3E2A1F"
GOLD = "B08D57"
PALE = "F7F3EA"
LIGHT = "FBF9F5"
BORDER = "D8C8AE"
MUTED = "6F5A49"
GREEN = "146C43"
AMBER = "8A5A00"
RED = "A61B1B"
WHITE = "FFFFFF"
BLUEGRAY = "E8EEF5"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run(run, size=11, bold=False, color=INK, italic=False, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_header(table.rows[0])


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run(run, size=9, color=MUTED)


def configure_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(existing_abs or [0]) + 1
    num_id = max(existing_num or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abs_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_numbered(doc, text, num_id):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)
    set_run(p.add_run(text), size=11)
    return p


def restart_numbering(doc, base_num_id):
    numbering = doc.part.numbering_part.element
    existing_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    new_id = max(existing_ids or [0]) + 1
    base = next(x for x in numbering.findall(qn("w:num")) if int(x.get(qn("w:numId"))) == base_num_id)
    abstract_id = base.find(qn("w:abstractNumId")).get(qn("w:val"))
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.extend([abstract, override])
    numbering.append(num)
    return new_id


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text), size=11)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]))
    else:
        set_run(p.add_run(text))
    return p


def add_callout(doc, title, body, tone="info"):
    colors = {
        "info": (PALE, BROWN),
        "success": ("EAF7EF", GREEN),
        "warning": ("FFF4D6", AMBER),
        "risk": ("FDECEC", RED),
    }
    fill, accent = colors[tone]
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(title), size=11, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    set_run(p2.add_run(body), size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_result_box(doc):
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [3120, 3120, 3120])
    labels = ["Result", "Evidence reference", "Tester / date"]
    values = ["PASS / FAIL / BLOCKED", "Screenshot or ticket no.: __________", "________________________"]
    for index, label in enumerate(labels):
        shade(table.cell(0, index), PALE)
        p = table.cell(0, index).paragraphs[0]
        set_run(p.add_run(label), size=9, bold=True, color=BROWN)
        p2 = table.cell(1, index).paragraphs[0]
        set_run(p2.add_run(values[index]), size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_test_case(doc, code, title, purpose, steps, expected, num_id, caution=None):
    add_heading(doc, f"{code} - {title}", 2)
    add_body(doc, f"Purpose: {purpose}", "Purpose:")
    if caution:
        add_callout(doc, "Control", caution, "warning")
    case_num_id = restart_numbering(doc, num_id)
    for step in steps:
        add_numbered(doc, step, case_num_id)
    add_body(doc, f"Expected result: {expected}", "Expected result:")
    add_result_box(doc)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    heading_tokens = {
        1: (16, BROWN, 18, 10),
        2: (13, BROWN, 14, 7),
        3: (12, GOLD, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(hp.add_run("MIZANTRA | INTELLIGENT CRM UAT"), size=8.5, bold=True, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(fp.add_run("Internal testing guide  |  Page "), size=9, color=MUTED)
    add_field(fp, "PAGE")

    num_id = configure_numbering(doc)

    # Customer-pack opening block.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run("TEAM ENABLEMENT AND ACCEPTANCE TESTING"), size=10, bold=True, color=GOLD)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    set_run(title.add_run("Mizantra Intelligent CRM"), size=30, bold=True, color=BROWN)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    set_run(subtitle.add_run("Step-by-Step User Acceptance Test Guide"), size=15, color=MUTED)

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    set_table_geometry(meta, [4680, 4680])
    meta_rows = [
        ("Environment", "https://mizantra.saksolution.com"),
        ("Release scope", "Lead capture, routing, pipeline, follow-up, prompt actions, conversion and Customer 360"),
        ("Document version", "1.0 | 02 September 2026"),
        ("Testing status", "To be completed by the assigned UAT team"),
    ]
    for row, (label, value) in zip(meta.rows, meta_rows):
        shade(row.cells[0], PALE)
        set_run(row.cells[0].paragraphs[0].add_run(label), size=10, bold=True, color=BROWN)
        set_run(row.cells[1].paragraphs[0].add_run(value), size=10)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_callout(
        doc,
        "Release rule",
        "Do not mark CRM accepted unless every mandatory case is PASS. Record screenshots for failures and do not use real customer secrets, API tokens, or personal data during testing.",
        "warning",
    )

    doc.add_page_break()
    add_heading(doc, "1. How to use this guide", 1)
    add_body(doc, "Run the tests in order. Earlier cases establish configuration and data required by later cases. Each tester must record the result and evidence reference directly below every case.")
    for item in [
        "Use a Super Admin for configuration cases and a normal Sales user for operational cases.",
        "Use the sample names in this guide, adding your initials if another tester is working simultaneously.",
        "Capture a screenshot whenever the actual result differs from the expected result.",
        "Log each defect with environment, user, time, steps, expected result, actual result and screenshot.",
        "After fixing a defect, repeat the failed case and every related downstream case.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Test accounts and responsibilities", 2)
    roles = doc.add_table(rows=1, cols=3)
    roles.style = "Table Grid"
    set_table_geometry(roles, [2100, 3600, 3660])
    for cell, text in zip(roles.rows[0].cells, ["Role", "Used for", "Assigned tester"]):
        shade(cell, BROWN)
        set_run(cell.paragraphs[0].add_run(text), size=9.5, bold=True, color=WHITE)
    for values in [
        ("Super Admin", "Readiness, routing rules and inbound channel setup", "________________________"),
        ("Sales User A", "Lead ownership, activities and pipeline execution", "________________________"),
        ("Sales User B", "Assignment distribution and access separation", "________________________"),
        ("Approver / Manager", "Forecast review and final UAT sign-off", "________________________"),
    ]:
        cells = roles.add_row().cells
        for cell, value in zip(cells, values):
            set_run(cell.paragraphs[0].add_run(value), size=9.5)
    set_repeat_header(roles.rows[0])

    add_heading(doc, "Test data", 2)
    for item in [
        "Company A: UAT Bluewater Systems <tester initials>",
        "Company B: UAT Coastal Controls <tester initials>",
        "Contact: Sameer Test | Phone: 9000012345 | Email: use a non-deliverable test address",
        "Expected value: INR 250,000 | Priority: High | Product interest: Marine antenna",
        "Inbound channel name: UAT Website <date and initials>",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2. Test sequence at a glance", 1)
    matrix = doc.add_table(rows=1, cols=4)
    matrix.style = "Table Grid"
    set_table_geometry(matrix, [1100, 3500, 1760, 3000])
    for cell, text in zip(matrix.rows[0].cells, ["ID", "Test area", "Tester", "Mandatory outcome"]):
        shade(cell, BROWN)
        set_run(cell.paragraphs[0].add_run(text), size=9, bold=True, color=WHITE)
    rows = [
        ("CRM-01", "Access and permissions", "Admin + Sales", "Correct access"),
        ("CRM-02", "Readiness and routing", "Admin", "Two owners configured"),
        ("CRM-03", "Manual lead and duplicate control", "Sales", "Create + reuse"),
        ("CRM-04", "CSV import", "Admin/Sales", "Valid rows imported"),
        ("CRM-05", "Secure inbound capture", "Admin", "Token-gated lead"),
        ("CRM-06", "WhatsApp capture", "Admin", "Opt-in behavior"),
        ("CRM-07", "Pipeline and stages", "Sales", "History retained"),
        ("CRM-08", "Activities and reminders", "Sales", "Due work visible"),
        ("CRM-09", "Ask Mizantra CRM prompts", "Sales", "Correct interpretation"),
        ("CRM-10", "Conversion and Customer 360", "Sales", "Connected customer"),
        ("CRM-11", "Mobile usability", "Sales", "No clipping/overflow"),
        ("CRM-12", "Security and recovery", "Admin", "Controls enforced"),
    ]
    for values in rows:
        cells = matrix.add_row().cells
        for cell, value in zip(cells, values):
            set_run(cell.paragraphs[0].add_run(value), size=9)
    set_repeat_header(matrix.rows[0])

    doc.add_page_break()
    add_heading(doc, "3. Detailed test cases", 1)
    add_test_case(
        doc, "CRM-01", "Access and permissions",
        "Confirm that authorized users can reach CRM and that normal users do not receive administration capabilities.",
        [
            "Sign in to Mizantra as Super Admin and open Sales > CRM.",
            "Confirm the Intelligent CRM heading, KPI cards, Pipeline, All leads, Follow-ups and Assignment rules are visible.",
            "Sign out and sign in as Sales User A. Open the same CRM workspace.",
            "Confirm the operational lead screens are visible and administrative actions follow the user's assigned permissions.",
            "Attempt access with a user who has no CRM or Sales permission, if such a test user is available.",
        ],
        "Authorized users can open CRM; restricted actions and screens are not exposed to unauthorized users.", num_id,
    )

    add_test_case(
        doc, "CRM-02", "Readiness and automatic assignment",
        "Configure dependable lead distribution before testing automatic intake.",
        [
            "As Super Admin, open Assignment rules.",
            "Review the CRM readiness checklist. Confirm it reports the current number of owners and channels.",
            "Create a rule named UAT General Lead Distribution.",
            "Choose Load balanced or Round robin, leave filters blank, search for Sales User A and Sales User B, and select both.",
            "Activate the rule and refresh CRM.",
            "Confirm the readiness warning for single-owner dependency is cleared.",
            "Create two test leads and confirm ownership is distributed according to the selected strategy.",
        ],
        "An active rule contains at least two valid users; new leads are assigned automatically and ownership is auditable.", num_id,
        "Choose actual authorized sales users. Do not select every active ERP user merely to clear the warning.",
    )

    add_test_case(
        doc, "CRM-03", "Manual lead creation and duplicate control",
        "Validate intelligent capture, assignment, scoring and duplicate prevention.",
        [
            "Open CRM and select New lead.",
            "Enter Company A, contact, phone, product interest, expected value and priority from the Test data section.",
            "Save the lead and note its lead number and assigned owner.",
            "Confirm the lead appears in All leads and the correct pipeline stage.",
            "Create the same company again with the same phone or email.",
            "Confirm Mizantra reuses or flags the existing lead instead of creating an uncontrolled duplicate.",
        ],
        "One governed lead exists with a lead number, stage, score and owner; duplicate details do not create an uncontrolled second record.", num_id,
    )

    add_test_case(
        doc, "CRM-04", "CSV lead import",
        "Confirm bulk lead capture handles valid, duplicate and invalid rows safely.",
        [
            "Prepare a CSV with headers for company_name, contact_person, email, phone, source and product_interest.",
            "Include one new valid lead, one row matching Company A and one row without a company name.",
            "Select Import CSV and upload the file.",
            "Record the created, reused and rejected counts shown by Mizantra.",
            "Search All leads and confirm only the valid new company was added.",
        ],
        "The new row is created, the duplicate is reused, the invalid row is rejected and the summary counts are accurate.", num_id,
    )

    doc.add_page_break()
    add_test_case(
        doc, "CRM-05", "Secure website/API inbound capture",
        "Validate token-protected, tenant-scoped external lead intake and replay protection.",
        [
            "As Super Admin, open Assignment rules and locate Website, email and API lead channels.",
            "Create a WEBSITE channel using the UAT channel name from Test data.",
            "Copy the one-time token into an approved secure test location; do not place it in screenshots or chat.",
            "Ask the technical tester to POST a sample lead to /api/v1/crm/inbound/<channel-id> using the token header.",
            "Repeat the identical event using the same external event ID.",
            "Confirm the first call creates or links a lead and the repeated call returns the existing result without duplication.",
            "Send one request with an invalid token and confirm it is rejected.",
            "Rotate the token and confirm the previous token no longer works.",
        ],
        "Valid intake is accepted once, duplicates are idempotent, invalid/old tokens are rejected and the lead is assigned automatically.", num_id,
        "The token is a secret and is displayed only once. Rotate it immediately if accidentally shared.",
    )

    add_test_case(
        doc, "CRM-06", "WhatsApp lead capture",
        "Confirm WhatsApp becomes a CRM source only when an administrator explicitly enables it.",
        [
            "Open Settings > WhatsApp Business and confirm the connection is healthy before proceeding.",
            "With CRM capture disabled, send a test inbound message from a number not already used in the UAT data.",
            "Confirm no CRM lead is created solely by that message.",
            "Enable CRM lead capture and select an appropriate owner if the screen requests one.",
            "Send a new inbound message from another test number.",
            "Confirm a CRM lead is created and its source/activity shows WhatsApp context.",
            "Send a second message from the same number and confirm it appends activity rather than creating another lead.",
        ],
        "WhatsApp capture is opt-in, tenant-scoped and duplicate-safe; repeated messages retain conversation history on one lead.", num_id,
        "Run this case only when approved WhatsApp test credentials and numbers are available. Otherwise mark BLOCKED, not FAIL.",
    )

    add_test_case(
        doc, "CRM-07", "Pipeline stages and ownership",
        "Confirm lead progression is clear, controlled and historically traceable.",
        [
            "Open Company A from the Pipeline view.",
            "Move it through Qualified, Proposal and Negotiation stages one stage at a time.",
            "Refresh after each change and confirm the card appears in the correct pipeline column.",
            "Change the owner to Sales User B and confirm the owner filter reflects the change.",
            "Open stage history and confirm previous stages, timestamps and users were retained.",
            "For a separate disposable lead, select Lost and confirm a loss reason is required.",
        ],
        "Stage and owner changes persist without refresh issues; history identifies who changed what and when; Lost requires a reason.", num_id,
    )

    add_test_case(
        doc, "CRM-08", "Activities, follow-ups and reminders",
        "Validate that customer work is scheduled, visible and completed with an outcome.",
        [
            "Open Company A and add a FOLLOW_UP activity scheduled for today with subject UAT proposal discussion.",
            "Add a second activity scheduled for a future date.",
            "Open Follow-ups and confirm the due activity is visible to the correct owner.",
            "Complete today's activity and enter a meaningful outcome.",
            "Refresh and confirm the activity is completed while the future activity remains open.",
            "Confirm the CRM KPI/reminder count responds correctly to the completion.",
        ],
        "Due work, future work and completed outcomes are distinguishable; reminders and KPIs use the saved activity state.", num_id,
    )

    doc.add_page_break()
    add_test_case(
        doc, "CRM-09", "Ask Mizantra CRM intelligence",
        "Confirm natural-language CRM requests are understood without relying on a single rigid sentence.",
        [
            "Open Ask Mizantra from the left navigation.",
            "Enter: Create a new prospect called UAT Harbor Electronics, contact Arif, phone 9000098765.",
            "Confirm the preview identifies a CRM lead action and shows the extracted company, contact and phone before any draft is created.",
            "Enter: Remind me next Friday to discuss the proposal with Company A.",
            "Confirm the preview identifies a follow-up and resolves Company A from CRM context.",
            "Enter: Move Company A to negotiation.",
            "Confirm the preview identifies a stage change and asks only for genuinely missing information.",
            "Repeat one request with spelling errors and changed word order, for example: nxt friday folowup Company A proposal.",
            "Confirm the interpretation remains correct; mark FAIL if it silently performs an unrelated action.",
        ],
        "The AI returns the correct CRM action and entities, preserves approval/control boundaries and does not substitute unrelated report results.", num_id,
        "Review the system-checked preview before confirming any write action. The prompt must never bypass normal approvals or permissions.",
    )

    add_test_case(
        doc, "CRM-10", "Lead conversion and Customer 360",
        "Validate the controlled transition from prospect to connected customer operations.",
        [
            "Open Company A and verify its contact information is sufficient for conversion.",
            "Select Convert to customer and confirm the action.",
            "Record the generated customer code.",
            "Open Customer 360 from the converted lead.",
            "Confirm the customer master and original CRM lead/activity history are linked.",
            "Confirm Sales, Finance, Service and Installed Assets sections load even when some sections contain no transactions yet.",
            "Start the quotation hand-off and confirm the converted customer is available for selection.",
        ],
        "One customer master is created, CRM history remains linked, Customer 360 loads safely and the customer is available to downstream Sales.", num_id,
        "Conversion is a controlled master-data action. Use the designated UAT company and do not convert junk or duplicate records.",
    )

    add_test_case(
        doc, "CRM-11", "Mobile usability",
        "Confirm the CRM is practical on a phone without hidden controls or horizontal page overflow.",
        [
            "Open Mizantra on a phone or browser device mode at approximately 390 px width.",
            "Open CRM and verify New lead, Pipeline, All leads and Follow-ups are reachable.",
            "Open Assignment rules and verify the readiness checklist, owner search and secure intake section are readable.",
            "Create and open a disposable test lead.",
            "Confirm the modal, forms, buttons and activity history fit the screen and can be scrolled vertically.",
            "Rotate the device or test a second narrow width and confirm the page never requires horizontal body scrolling.",
        ],
        "No content is cut off, primary actions remain reachable, and the browser shows no blank screen or client-side exception.", num_id,
    )

    add_test_case(
        doc, "CRM-12", "Security, error handling and recovery",
        "Confirm failures are contained and users receive safe, understandable feedback.",
        [
            "Attempt to create a fixed-owner rule with zero or more than one owner.",
            "Attempt to create a load-balanced rule with no eligible owners.",
            "Confirm the server rejects each invalid configuration with a readable message.",
            "Temporarily disconnect the network and attempt a non-destructive CRM read or draft operation.",
            "Restore the connection, refresh and confirm existing CRM data remains intact.",
            "Sign out, revisit the CRM URL and confirm authentication is required.",
            "Review browser console/network logs and confirm no secret token or another tenant's data is exposed.",
        ],
        "Invalid rules are rejected, network failures do not corrupt CRM records, authentication is enforced and tenant/security data is not leaked.", num_id,
    )

    doc.add_page_break()
    add_heading(doc, "4. Defect reporting standard", 1)
    add_body(doc, "Use one defect for each distinct problem. Avoid messages such as 'not working' without evidence.")
    defect = doc.add_table(rows=1, cols=2)
    defect.style = "Table Grid"
    set_table_geometry(defect, [2700, 6660])
    for cell, text in zip(defect.rows[0].cells, ["Field", "What to record"]):
        shade(cell, BROWN)
        set_run(cell.paragraphs[0].add_run(text), size=9.5, bold=True, color=WHITE)
    defect_rows = [
        ("Title", "Module + action + observed failure"),
        ("Environment", "Mizantra URL, browser/device and date/time"),
        ("User", "Role and username; never include the password"),
        ("Steps", "Exact sequence that reproduces the issue"),
        ("Expected", "What this guide or approved requirement says should occur"),
        ("Actual", "What occurred, including exact error text"),
        ("Evidence", "Screenshot, screen recording, request ID or lead number"),
        ("Severity", "Critical, High, Medium or Low using the rules below"),
    ]
    for label, value in defect_rows:
        cells = defect.add_row().cells
        shade(cells[0], PALE)
        set_run(cells[0].paragraphs[0].add_run(label), size=9.5, bold=True, color=BROWN)
        set_run(cells[1].paragraphs[0].add_run(value), size=9.5)
    set_repeat_header(defect.rows[0])

    add_heading(doc, "Severity rules", 2)
    for item in [
        "Critical: security exposure, cross-tenant data, uncontrolled transaction, data corruption or complete CRM outage.",
        "High: mandatory workflow cannot be completed and no safe workaround exists.",
        "Medium: workflow works with a safe workaround, or an important result is confusing or wrong.",
        "Low: visual or wording issue that does not alter business data or block the workflow.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5. Acceptance and sign-off", 1)
    add_callout(
        doc,
        "Go-live acceptance gate",
        "CRM is accepted only when CRM-01 through CRM-05 and CRM-07 through CRM-12 are PASS. CRM-06 may be BLOCKED only when WhatsApp credentials are intentionally out of scope for this release.",
        "success",
    )
    signoff = doc.add_table(rows=1, cols=4)
    signoff.style = "Table Grid"
    set_table_geometry(signoff, [2400, 2160, 2400, 2400])
    for cell, text in zip(signoff.rows[0].cells, ["Approval role", "Decision", "Name / signature", "Date"]):
        shade(cell, BROWN)
        set_run(cell.paragraphs[0].add_run(text), size=9, bold=True, color=WHITE)
    for role in ["CRM Process Owner", "Sales Manager", "System Administrator", "Final Business Approver"]:
        cells = signoff.add_row().cells
        values = [role, "ACCEPT / REJECT", "________________", "____________"]
        for cell, value in zip(cells, values):
            set_run(cell.paragraphs[0].add_run(value), size=9.5)
    set_repeat_header(signoff.rows[0])

    add_heading(doc, "Final release notes", 2)
    for item in [
        "Confirmed CRM assignment owners: ______________________________________________",
        "Confirmed live intake channel: _________________________________________________",
        "Open accepted exceptions and owners: __________________________________________",
        "Planned production activation date: ____________________________________________",
    ]:
        add_body(doc, item)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Mizantra Intelligent CRM - UAT Test Guide"
    doc.core_properties.subject = "Step-by-step team testing and sign-off"
    doc.core_properties.author = "SAK Solutions"
    doc.core_properties.keywords = "Mizantra, CRM, UAT, testing, acceptance"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
